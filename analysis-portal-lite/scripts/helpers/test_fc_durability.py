"""FC durability, milestones M1-M3: names, ordering, holds, clock, the OCV
analysis and death rule, characterizations, the crossover rule and the
verdict, outputs.

Builds synthetic runs in the station format -- one CSV per step, the real
column layouts and file-name grammar, Excel-edited extra columns, H2/N2 steps
and the 30 s every step drops from the end of its log -- in a temp dir, so
the harness needs no lab data. The repository is public; the real runs stay
off it. The OCV cases are the plan's verification cases 2 (death mid-hold 4),
4 (a transient at a hold's start) and 5 (a hold split across two files);
the characterization cases are 1 (crossover death), 3 (alive), 6 (BOL
failure), 7 (short-dominated) and 8 (missing crossover); each comes with
the guards its part of the rule exists for. Characterization files are
synthetic CVs, GEIS sweeps and stepped pol curves with known answers.

Run on its own:  PYTHONPATH=. python scripts/helpers/test_fc_durability.py
"""
import io
import json
import os
import shutil
import sys
import tempfile
import contextlib
from dataclasses import replace as dc_replace
from pathlib import Path

os.environ.setdefault('MPLBACKEND', 'Agg')
import numpy as np

passed = failed = 0


def check(name, got, want):
    global passed, failed
    if got == want:
        passed += 1
    else:
        failed += 1
        print(f'  FAIL {name}\n       got  {got!r}\n       want {want!r}')


def close(name, got, want, tol):
    global passed, failed
    if got is not None and abs(got - want) <= tol:
        passed += 1
    else:
        failed += 1
        print(f'  FAIL {name}: got {got!r}, want {want!r} +/- {tol}')


def true(name, cond, detail=''):
    global passed, failed
    if cond:
        passed += 1
    else:
        failed += 1
        print(f'  FAIL {name} {detail}')


# Importing the package prints which registered scripts are absent; that is
# noise here.
with contextlib.redirect_stdout(io.StringIO()):
    from scripts.fc_durability_analysis import (
        parse_step_name, natural_key, build_timeline, order_folders, run,
        analyze_ocv, load_hold_trace, below_runs, rolling_decay, OcvSettings,
        PLOT_BINS_CAP, ocv_plot_bin, analyze_characterizations, decide,
        CharSettings, _pick_curves, analyze_crossover_file, _series_labels, BAND_LOW, BAND_HIGH,
        rh_cycle_map, rh_cycles_absent)
    from scripts.ecsa_analysis import compute_ecsa_hupd, generate_synthetic_cv
    from scripts.eis_analysis import extract_hfr, MODELS


# ── station-format writers ──────────────────────────────────────────

FOUR = 'Time Stamp,Counter Electrode Voltage,Working Electrode Voltage,Current'
POL = 's,mV,mA,mW'
GEIS = ('Time Stamp,DC Electrode Voltage,DC Current,Frequency,Abs Impedance,'
        'Phase Angle,Real Impedance,Imaginary Impedance,Total harmonic distortion')


def stamps(programmed_s, dt, t0=None):
    """Log stamps as the stations write them: from one interval in, ending
    30 s before the step does."""
    start = dt if t0 is None else t0
    return np.arange(start, programmed_s - 30 + 1e-9, dt)


def write_four(path, programmed_s, dt, v=0.9, i=0.0, extra=None, t0=None):
    t = stamps(programmed_s, dt, t0)
    vs = v(t) if callable(v) else np.full(len(t), v)
    head = FOUR + (',' + ','.join(extra) if extra else '')
    pad = ',' * len(extra) if extra else ''
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    with open(path, 'w') as fh:
        fh.write(head + '\n')
        for tt, vv in zip(t, vs):
            fh.write(f'{tt:.6f},{1e-5:.6g},{vv:.6f},{i:.6g}{pad}\n')
    return len(t), float(t[0]), float(t[-1])


def write_pol(path, programmed_s, dt):
    t = stamps(programmed_s, dt)
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    with open(path, 'w') as fh:
        fh.write(POL + '\n')
        for k, tt in enumerate(t):
            mv = 900 - (k % 40) * 15
            ma = (k % 40) * 180.0
            fh.write(f'{tt:.3f},{mv},{ma:.1f},{mv * ma / 1000:.1f}\n')


def write_geis(path):
    f = np.logspace(5, -1, 61)
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    with open(path, 'w') as fh:
        fh.write(GEIS + '\n')
        for k, ff in enumerate(f):
            fh.write(f'{2.0 * (k + 1):.3f},0.10,-0.501,{ff:.4f},0.012,-10.0,0.010,'
                     f'{-0.003 * np.sin(k / 9):.6f},0.01\n')


def touch(path, data=b''):
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    Path(path).write_bytes(data)


def scribner_run(root):
    """One folder, the Gore-A shape: break-in and BOL, an 8 h H2/N2 step that
    the old 50-minute rule would have called a hold, six back-to-back 10 h
    holds (c5..c10, so string order would separate c10 from c9), EOL with
    repeated crossovers, and the clutter a real folder carries."""
    d = Path(root) / '231127_Synth-A'
    n = 'Synth-A'
    write_four(d / f'a0_{n}_OCV_80.3Cell-100RH_0.05WH2-0.05WN2.csv', 5400, 10, v=0.09)
    write_four(d / f'a4_{n}_H2X_80.3Cell-100RH_0.4WH2-0.4WN2.csv', 720, 0.5, v=0.4, i=0.018)
    write_pol(d / f'b1_{n}_Breakin_80.3Cell-100RH_0.2WH2-0.2WA.csv', 32400, 60)
    write_four(d / f'b2_{n}_OCV_80.3Cell-95RH_0.1WH2-0.1WA.csv', 3600, 10, v=0.88)
    write_pol(d / f'b3_{n}_PolarizationCurve_80.3Cell-95RH_0.2WH2-0.2WA.csv', 4900, 5)
    touch(d / f'b3_{n}_PolarizationCurve_80.3Cell-95RH_0.2WH2-0.2WA_FILTERDATA.csv',
          b'mV,mA/cm2,mW/cm2\n878,0.0,0.0\n')
    touch(d / f'b3_{n}_PolarizationCurve_80.3Cell-95RH_0.2WH2-0.2WA_IVCurve.tiff', b'II*\x00')
    write_four(d / f'b6_{n}_OCV_80.3Cell-100RH_0.05WN2-0.5DN2B3.csv', 600, 10, v=0.7)
    write_four(d / f'b14_{n}_OCV_90.3Cell-100RH_0.05WH2-0.05WN2.csv', 28800, 30, v=0.1)
    write_four(d / f'b20_{n}_H2X_80.3Cell-100RH_0.4WH2-0.4WN2.csv', 720, 0.5, v=0.4,
               i=0.0185, extra=['cd', '', ''])
    write_four(d / f'b24_{n}_ECSA_80.3Cell-100RH_0.05WH2-0.05WN2.csv', 140, 0.5, v=0.5, i=0.01)
    write_geis(d / f'b26_{n}_GEIS_80.3Cell-100RH_0.05WH2-0.05WN2.csv')
    write_four(d / f'c2_{n}_OCV_95Cell-30RH_0.05WH2-0.05WN2.csv', 7200, 10, v=0.09)
    write_four(d / f'c3_{n}_OCV_95Cell-30RH_0.05WH2-0.05WA.csv', 60, 1, v=0.87)
    write_four(d / f'c4_{n}_OCV_95Cell-30RH_0.1WH2-0.2WA.csv', 60, 1, v=0.92)
    for k in range(5, 11):
        write_four(d / f'c{k}_{n}_OCV_95Cell-30RH_0.1WH2-0.2WA.csv', 36000, 30,
                   v=lambda t, k=k: 0.92 - 0.01 * (k - 5) - 1e-7 * t)
    write_four(d / f'c20_{n}_OCV_80.3Cell-95RH_0.05WH2-0.05WA.csv', 3600, 10, v=0.79)
    for k in (32, 40, 41):
        write_four(d / f'c{k}_{n}_H2X_80.3Cell-100RH_0.4WH2-0.4WN2.csv', 720, 0.5, v=0.4, i=0.02)
    touch(d / 'Synth-A_OCV-Notes.txt', b'notes\n')
    touch(d / 'Synth-A_Full-OCV.xlsx', b'PK\x03\x04')
    touch(d / 'readme_export.csv', b'a,b\n1,2\n')
    return d


def fcts_folder(d, sample='01M12umSynth_2'):
    """One 40 h FCTS folder: two RH cycles' worth of the real protocol
    collapsed to one -- 10 h at 30% RH, then 10.425 h at 0% RH."""
    n = sample
    write_four(d / f'c1_{n}_OCV_90Cell-30RH_0.1WH2-0.1WN2.csv', 3600, 10, v=0.11)
    write_four(d / f'c2_{n}_OCV_90Cell-30RH_0.1WH2-0.1WA.csv', 60, 1, v=0.95)
    write_four(d / f'c3_{n}_OCV_90Cell-30RH_0.1WH2-0.1WA.csv', 36000, 30, v=0.92)
    write_four(d / f'c4_{n}_OCV_90Cell-100RH_0.1WH2-0.1WN2.csv', 3600, 10, v=0.12)
    write_four(d / f'c5_{n}_OCV_90Cell-0RH_0.1WH2-0.1DA.csv', 60, 1, v=0.92)
    write_four(d / f'c6_{n}_OCV_90Cell-0RH_0.1WH2-0.1DA.csv', 37530, 30, v=0.90)
    write_four(d / f'c9_{n}_OCV_90Cell-30RH_0.1WH2-0.1WN2.csv', 600, 1,
               v=lambda t: 0.84 - 0.41 * t / t[-1])
    write_four(d / f'e11_{n}_H2X_80.3Cell-100RH_0.4WH2-0.4WN2.csv', 720, 0.5, v=0.4, i=0.009)


def fcts_run(root):
    """A BOL folder named so it sorts last alphabetically, and three 40 h
    folders whose step letters restart -- 40H_10 must follow 40H_2."""
    base = Path(root) / 'FCTS_Synth'
    bol = base / 'Conditioning_BOL'
    n = '01M12umSynth_2'
    write_four(bol / f'a0_{n}_OCV_80.3Cell-100RH_0.05WH2-0.05WN2.csv', 3600, 10, v=0.09)
    write_pol(bol / f'b1_{n}_Breakin_80.3Cell-100RH_0.2WH2-0.2WA.csv', 7200, 60)
    write_four(bol / f'b20_{n}_H2X_80.3Cell-100RH_0.4WH2-0.4WN2.csv', 720, 0.5, v=0.4, i=0.0085)
    for name in ('40H_10', '40H_2', '40H_1'):
        fcts_folder(base / name)
    return base


tmp = Path(tempfile.mkdtemp(prefix='fcdur_m1_'))

# ── 1. file names ────────────────────────────────────────────────────
print('1. file names')
r = parse_step_name('c5_Synth-A_OCV_95Cell-30RH_0.1WH2-0.2WA.csv')
check('step', (r['step'], r['letter'], r['number']), ('c5', 'c', 5))
check('sample', r['sample'], 'Synth-A')
check('kind', (r['kind'], r['kind_token']), ('ocv', 'OCV'))
check('T and RH', (r['T_C'], r['RH_pct']), (95.0, 30.0))
check('anode', (r['anode_gas'], r['anode_slpm'], r['anode_humid']), ('H2', 0.1, 'wet'))
check('cathode', (r['cathode_gas'], r['cathode_slpm'], r['cathode_humid']), ('Air', 0.2, 'wet'))
check('no suffix', r['suffix'], '')

r = parse_step_name('e5_01M12umSynth_2_OCV_80.3Cell-100RH_0.05WN2-0.5DN2B3.csv')
check('sample with underscore', r['sample'], '01M12umSynth_2')
check('N2/N2 with B3 tail', (r['anode_gas'], r['cathode_gas'], r['cathode_slpm'],
                             r['cathode_humid']), ('N2', 'N2', 0.5, 'dry'))
check('dry air', parse_step_name('c6_S_OCV_90Cell-0RH_0.1WH2-0.1DA.csv')['cathode_humid'], 'dry')
check('EClean50x', parse_step_name('a2_S_EClean50x_80.3Cell-100RH_0.2WH2-0.2WN2.csv')['kind'], 'eclean')
check('GEIS', parse_step_name('b26_S_GEIS_80.3Cell-100RH_0.05WH2-0.05WN2.csv')['kind'], 'eis')
check('Breakin', parse_step_name('b1_S_Breakin_80.3Cell-100RH_0.2WH2-0.2WA.csv')['kind'], 'breakin')
r = parse_step_name('b5_S_PolarizationCurve_80.3Cell-95RH_0.75WH2-1WA.csv')
check('pol curve flows', (r['kind'], r['anode_slpm'], r['cathode_slpm']), ('polcurve', 0.75, 1.0))
check('FILTERDATA kept as suffix',
      parse_step_name('b3_S_PolarizationCurve_80.3Cell-95RH_0.2WH2-0.2WA_FILTERDATA.csv')['suffix'],
      'FILTERDATA')
check('O2 cathode, upper-case step', (lambda q: (q['step'], q['cathode_gas']))(
    parse_step_name('C12_S_OCV_80Cell-50RH_0.4WH2-0.4WO2.csv')), ('c12', 'O2'))
check('Air spelled out', parse_step_name('c3_S_OCV_80Cell-50RH_0.4WH2-0.4WAir.csv')['cathode_gas'], 'Air')
check('kind word inside the sample', (lambda q: (q['sample'], q['kind']))(
    parse_step_name('c3_OCV-batch_OCV_80Cell-50RH_0.4WH2-0.4WA.csv')), ('OCV-batch', 'ocv'))
check('unknown kind', parse_step_name('c3_S_Hold_80Cell-50RH_0.4WH2-0.4WA.csv')['kind'], 'other')
check('not a step file', parse_step_name('readme_export.csv'), None)
check('no conditions token', parse_step_name('c3_S_OCV_notes.csv'), None)
check('no sample', parse_step_name('c3_OCV_80Cell-50RH_0.4WH2-0.4WA.csv'), None)
check('natural sort', sorted(['40H_10', '40H_2', '40H_1'], key=natural_key),
      ['40H_1', '40H_2', '40H_10'])
# A letter after the step number (f3b, a step added after f3) is part of the
# step, as the index reads it (19 Oct: f3b was skipped as no step file).
r = parse_step_name('f3b_Synth-F_OCV_80.3Cell-100RH_0.05WH2-0.05WN2.csv')
check('a letter after the step number', (r['step'], r['letter'], r['number'], r['trail'], r['sample']),
      ('f3b', 'f', 3, 'b', 'Synth-F'))
check('... upper case and a leading zero as for any step; two letters after it is no step',
      [(lambda r: r and (r['step'], r['trail']))(parse_step_name(n)) for n in (
          'B17B_S_OCV_80.3Cell-100RH_0.4WH2-0.4WO2.csv', 'f03b_S_OCV_80.3Cell-100RH_0.4WH2-0.4WO2.csv',
          'f3bc_S_OCV_80.3Cell-100RH_0.4WH2-0.4WO2.csv')], [('b17b', 'b'), ('f3b', 'b'), None])
for folder, names in (('trail', ('f2', 'f3', 'f3b', 'f4')), ('trail_up', ('F3B', 'F2', 'F3', 'f3a')),
                      ('trail_lone', ('f2', 'f3b', 'f4'))):
    for k, step in enumerate(names):
        write_four(tmp / folder / '261019_Synth-F' / f'{step}_Synth-F_OCV_80.3Cell-100RH_0.05WH2-0.05WN2.csv',
                   600, 10, v=0.1)
tl = build_timeline(str(tmp / 'trail'))
check('a run with f3 and f3b: both read, f3b after f3 and before f4, nothing skipped, '
      'no step said to repeat or be missing',
      ([f.step for f in tl.files], tl.skipped, [n for n in tl.notes if 'twice' in n or 'no file for' in n]),
      (['f2', 'f3', 'f3b', 'f4'], [], []))
check('... ordered by the step, not the name, whatever the case: F3 before f3a before F3B',
      [f.step for f in build_timeline(str(tmp / 'trail_up')).files], ['f2', 'f3', 'f3a', 'f3b'])
tl = build_timeline(str(tmp / 'trail_lone'))
check('a run with f3b and no f3: f3b stands for step 3, so no step is said to be missing',
      ([f.step for f in tl.files], [n for n in tl.notes if 'no file for' in n]), (['f2', 'f3b', 'f4'], []))

# ── 2. single-folder run (Scribner shape) ────────────────────────────
print('2. single-folder run')
sc_dir = scribner_run(tmp / 'scribner')
tl = build_timeline(str(tmp / 'scribner'))
roles = {}
for f in tl.files:
    roles[f.role] = roles.get(f.role, 0) + 1
check('roles', roles, {'h2n2': 3, 'crossover': 5, 'breakin': 1, 'air-step': 4,
                       'polcurve': 1, 'n2n2': 1, 'ecsa': 1, 'eis': 1, 'hold': 6})
check('step files', len(tl.files), 23)
check('holds are c5-c10 in order', [f.step for f in tl.holds],
      ['c5', 'c6', 'c7', 'c8', 'c9', 'c10'])
check('one hold segment', len(tl.segments), 1)
close('hours tested = 6 x (35,970 + 30) s', tl.hours_tested, 60.0, 1e-9)
b14 = next(f for f in tl.files if f.step == 'b14')
check('8 h H2/N2 step is not a hold', (b14.role, b14.length_s), ('h2n2', 28800.0))
check('1 h air OCV is a step', next(f for f in tl.files if f.step == 'b2').role, 'air-step')
check('steps in order', [f.step for f in tl.files][:6], ['a0', 'a4', 'b1', 'b2', 'b3', 'b6'])
check('c10 follows c9', [f.step for f in tl.files].index('c10') -
      [f.step for f in tl.files].index('c9'), 1)
b20 = next(f for f in tl.files if f.step == 'b20')
check('BOL crossover stamped at 0 h', (b20.interval, b20.H_start_h), (0, 0.0))
check('Excel-edited file reads', b20.n_rows, 1380)
c41 = next(f for f in tl.files if f.step == 'c41')
check('EOL crossover stamped after the hold', (c41.interval, round(c41.H_start_h, 6)), (1, 60.0))
c7 = next(f for f in tl.files if f.step == 'c7')
check('hold file clock', (c7.segment, round(c7.H_start_h, 6), round(c7.H_end_h, 6), c7.clock_s),
      (1, 20.0, 30.0, 36000.0))
check('ignored files', tl.ignored, {'.tiff': 1, '.txt': 1, '.xlsx': 1})
check('skipped files', sorted(reason for _, reason in tl.skipped),
      ['name does not follow the step-file pattern', 'station pol-curve summary (FILTERDATA)'])
check('gap note', [n for n in tl.notes if 'no file for' in n][:1],
      ['231127_Synth-A: no file for a1, a2, a3'])
check('conditions text', tl.segments[0].conditions, '95 °C, 30% RH, H2/Air 0.1/0.2 slpm')

tl2 = build_timeline(str(tmp / 'scribner'), hold_min_h=0.5)
check('hold_min_h 0.5 makes the 1 h air steps holds', [f.step for f in tl2.holds],
      ['b2', 'c5', 'c6', 'c7', 'c8', 'c9', 'c10', 'c20'])
check('... in three segments', len(tl2.segments), 3)
close('... adding 2 h', tl2.hours_tested, 62.0, 1e-9)
close('step_tail_s 0 counts logged time only',
      build_timeline(str(tmp / 'scribner'), step_tail_s=0).hours_tested, 6 * 35970 / 3600, 1e-9)

# ── 3. multi-folder run (FCTS shape) ─────────────────────────────────
print('3. multi-folder run')
fcts_run(tmp / 'fcts')
tl = build_timeline(str(tmp / 'fcts'))
check('folder order', tl.folders, ['FCTS_Synth/Conditioning_BOL', 'FCTS_Synth/40H_1',
                                   'FCTS_Synth/40H_2', 'FCTS_Synth/40H_10'])
check('holds', [(f.folder.split('/')[-1], f.step) for f in tl.holds],
      [('40H_1', 'c3'), ('40H_1', 'c6'), ('40H_2', 'c3'), ('40H_2', 'c6'),
       ('40H_10', 'c3'), ('40H_10', 'c6')])
check('six segments', len(tl.segments), 6)
close('hours tested = 3 x (10 + 10.425) h', tl.hours_tested, 61.275, 1e-9)
check('segment ends', [round(s.H_end_h, 6) for s in tl.segments],
      [10.0, 20.425, 30.425, 40.85, 50.85, 61.275])
e11 = [(f.folder.split('/')[-1], f.interval, round(f.H_start_h, 6))
       for f in tl.files if f.step == 'e11']
check('characterization stamped per folder', e11,
      [('40H_1', 2, 20.425), ('40H_2', 4, 40.85), ('40H_10', 6, 61.275)])
c9 = next(f for f in tl.files if f.step == 'c9')
check('purge falling through 0.7 V is a gas step', c9.role, 'h2n2')
check('missing c7, c8 noted per folder',
      sum(1 for n in tl.notes if n.endswith('no file for c7, c8')), 3)
check('dry hold conditions', tl.segments[1].conditions, '90 °C, 0% RH, H2/dry Air 0.1/0.1 slpm')

# Folder names and segment hours label the timeline from the figure, so the
# stored copy keeps only the readout: every text of an axes is stored as
# readout, and a re-render listed the folders and hours there.
from scripts.fc_durability_analysis import plot_timeline
_fig = plot_timeline(tl)
check('timeline: folders and hours labelled from the figure, the axes holding only the readout; on '
      'hold hours a folder of no hold time shares its boundary with the next',
      (len(_fig.axes[0].texts), sorted(t.get_text().strip() for t in _fig.texts)),
      (1, sorted(['Conditioning_BOL → 40H_1', '40H_2', '40H_10']
                 + [f'{s.H_end_h:.2f} h' for s in tl.segments])))
check('... each hold segment after the first parted from the one before by a line the figure holds',
      [round(float(ln.get_xdata()[0]), 6) for ln in _fig.artists], [10.0, 20.425, 30.425, 40.85, 50.85])
# On the hold clock every step that is not a hold takes no time: it is a tick
# at the hold hour it ran at, so the hold lane runs end to end and the
# crossover lane ticks at each characterization.
lanes_ = {ln.get_label(): ln.get_xdata() for ln in _fig.axes[0].get_lines() if len(ln.get_xdata()) > 2}
pulse_mid = lambda xs: (np.asarray(xs)[1:-1:4] + np.asarray(xs)[3:-1:4]) / 2.0     # each pulse's middle
check('... holds fill their lane over the hold hours, the line dipping between segments; a '
      'characterization is a tick at the hold hour it ran at; the axis is the hold clock, a hair wider',
      (list(np.round(lanes_['OCV hold'], 3)), list(np.round(pulse_mid(lanes_['H₂ crossover']), 4)),
       [round(x, 3) for x in _fig.axes[0].get_xlim()]),
      ([-0.368] + [x for a_, b_ in zip([0.0, 10.0, 20.425, 30.425, 40.85, 50.85],
                                        [10.0, 20.425, 30.425, 40.85, 50.85, 61.275])
                   for x in (a_, a_, b_, b_)] + [61.643],
       [0.0, 20.425, 40.85, 61.275], [-0.368, 61.643]))
gas_h = {}
for f_ in tl.files:
    if f_.role == 'h2n2':
        gas_h[round(f_.H_start_h, 4)] = gas_h.get(round(f_.H_start_h, 4), 0) + 1
check('... steps at one hold hour share one tick (several H2/N2 steps before each hold)',
      (max(gas_h.values()) > 1, list(np.round(pulse_mid(lanes_['H₂/N₂ step']), 4))),
      (True, sorted(gas_h)))
import matplotlib.pyplot as _plt
_plt.close(_fig)

# Explicit order: exact paths win, the rest follow canonically.
tl = build_timeline(str(tmp / 'fcts'),
                    folder_order=['FCTS_Synth/40H_10', 'FCTS_Synth/Conditioning_BOL'])
check('explicit folder order', tl.folders, ['FCTS_Synth/40H_10', 'FCTS_Synth/Conditioning_BOL',
                                            'FCTS_Synth/40H_1', 'FCTS_Synth/40H_2'])
true('note on folders left out', any('placed after it' in n for n in tl.notes), tl.notes)

# ── 4. ordering and segment edge cases ───────────────────────────────
print('4. edge cases')
check('exact folder matching, no substrings',
      order_folders({'Trap/A_hold1': 'c', 'Trap/hold1': 'c'}, ['Trap/hold1', 'Trap/A_hold1'])[0],
      ['Trap/hold1', 'Trap/A_hold1'])
order, notes = order_folders({'Trap/A_hold1': 'c', 'Trap/hold1': 'c'}, ['hold1'])
check('partial path matches nothing', order, ['Trap/A_hold1', 'Trap/hold1'])
true('... and says so', any('matches no folder exactly' in n for n in notes), notes)

cond = tmp / 'cond'
write_four(cond / 'c3_S_OCV_90Cell-30RH_0.1WH2-0.1WA.csv', 36000, 30)
write_four(cond / 'c4_S_OCV_90Cell-0RH_0.1WH2-0.1DA.csv', 36000, 30)
write_four(cond / 'c5_S_OCV_90Cell-0RH_0.1WH2-0.1DA.csv', 36000, 30)
tl = build_timeline(str(cond))
check('a change of conditions starts a new hold', [[f.step for f in s.files] for s in tl.segments],
      [['c3'], ['c4', 'c5']])

edge = tmp / 'edge'
write_four(edge / 'c3_S_OCV_90Cell-30RH_0.1WH2-0.1WA.csv', 7200, 30)
write_four(edge / 'c5_S_OCV_90Cell-30RH_0.1WH2-0.1WN2.csv', 600, 10)
write_four(edge / 'c6_S_OCV_90Cell-30RH_0.1WH2-0.1WA.csv', 7170, 30)
tl = build_timeline(str(edge))
check('a 2 h step logged every 30 s is a hold; 1.99 h is not',
      [(f.step, f.role) for f in tl.files if f.cathode_gas == 'Air'],
      [('c3', 'hold'), ('c6', 'air-step')])

split = tmp / 'split'
write_four(split / 'part_1' / 'c3_S_OCV_90Cell-30RH_0.1WH2-0.1WA.csv', 36000, 30)
write_four(split / 'part_2' / 'c3_S_OCV_90Cell-30RH_0.1WH2-0.1WA.csv', 36000, 30)
tl = build_timeline(str(split))
check('a hold never spans folders', [len(s.files) for s in tl.segments], [1, 1])
close('... clock still adds up', tl.hours_tested, 20.0, 1e-9)

late = tmp / 'late'
write_four(late / 'c3_S_OCV_90Cell-30RH_0.1WH2-0.1WA.csv', 41030, 30, t0=5000)
tl = build_timeline(str(late))
check('late-starting log counts its span', tl.holds[0].clock_s, 36030.0)
true('... with a note', 'log starts at 5000 s' in tl.holds[0].note, tl.holds[0].note)

empty = tmp / 'empty'
touch(empty / 'only.tiff', b'II*\x00')
tl = build_timeline(str(empty))
check('no step files', (len(tl.files), tl.hours_tested), (0, 0.0))

nohold = tmp / 'nohold'
write_four(nohold / 'b20_S_H2X_80.3Cell-100RH_0.4WH2-0.4WN2.csv', 720, 0.5)
tl = build_timeline(str(nohold))
true('no-hold note', any(n.startswith('no OCV holds found') for n in tl.notes), tl.notes)

# ── 5. run() end to end ──────────────────────────────────────────────
print('5. run()')
out = tmp / 'out'
with contextlib.redirect_stdout(io.StringIO()):
    res = run(str(tmp / 'scribner'), str(out), {'image_format': 'png', 'loading': '0.2'})
check('status', res['status'], 'success')
row = res['summary'][0]
check('summary identity', (row['Analysis'], row['Label']), ('fc_durability', 'fcdur'))
close('summary hours', row['Hours tested'], 60.0, 1e-9)
check('summary counts', (row['Hold files'], row['Hold segments'], row['Step files']), (6, 1, 23))
check('loading and area parameters', (row['loading'], row['geo_area']), (0.2, 5.0))
true('timeline png', (out / 'fcdur_timeline.png').exists())
sidecar_path = out / '_plot_data' / 'fcdur_timeline.json'
true('timeline sidecar', sidecar_path.exists())


def _no_nan(text):
    raise ValueError(f'non-finite value {text} in sidecar')


sidecar = json.loads(sidecar_path.read_text(), parse_constant=_no_nan)
check('sidecar plot type', sidecar['plot_type'], 'fcdur_timeline')
ax0 = sidecar['data']['axes'][0]
check('one square-wave line per lane', len(ax0['lines']), 9)
true('lanes kept as lines, not axhlines',
     all(len(ln['x']) >= 6 for ln in ax0['lines']) and not ax0['axhlines'])
from scripts.helpers.record import extract_values
vals = extract_values(sidecar)
close('readout reaches the record layer', (vals.get('Hours tested') or {}).get('value'), 60.0, 1e-9)

import openpyxl
wb = openpyxl.load_workbook(out / 'fc_durability.xlsx', read_only=True)
check('workbook sheets', wb.sheetnames, ['Summary', 'Cycles', 'Holds', 'OCV by condition',
                                         'OCV events', 'Decay rate', 'OCV trace', 'Hold files',
                                         'Characterizations', 'Manifest', 'Notes'])
check('manifest rows: 23 step files + 2 skipped', wb['Manifest'].max_row - 1, 25)
check('holds sheet rows', wb['Holds'].max_row - 1, 1)
summary_items = {r[0]: r[1] for r in wb['Summary'].iter_rows(min_row=2, values_only=True)}
check('nothing left out at upload: said so, and no summary field',
      (summary_items.get('Folders left out at upload'), 'Folders left out' in row), ('none', False))
check('no RH cycles from station CSVs logged every 30 s, too coarsely to read them off, and the Summary '
      'says so; none in the stored row',
      (summary_items.get('RH cycles tested'), 'RH cycles tested' in row),
      ('none counted: c5, c6, c7, c8, c9, c10: unreadable (see the notes)', False))
wb.close()

# Folders skipped in the upload page's folder list never arrive; the page
# names them, and every output says the analysis did not see them.
out_left = tmp / 'out_left'
with contextlib.redirect_stdout(io.StringIO()):
    res_left = run(str(tmp / 'scribner'), str(out_left),
                   {'image_format': 'none', 'skipped_folders': ['Run/40H_3/', '  ', 'Run/extra']})
wb = openpyxl.load_workbook(out_left / 'fc_durability.xlsx', read_only=True)
summary_items = {r[0]: r[1] for r in wb['Summary'].iter_rows(min_row=2, values_only=True)}
notes_left = [r[0] for r in wb['Notes'].iter_rows(min_row=2, values_only=True)]
wb.close()
check('skipped folders: named in the workbook summary, blanks dropped',
      summary_items.get('Folders left out at upload'), 'Run/40H_3, Run/extra')
true('... noted as not in the analysis',
     any('left out at upload' in str(n) and 'Run/40H_3, Run/extra' in str(n) for n in notes_left),
     notes_left)
check('... and in the stored summary row', res_left['summary'][0].get('Folders left out'),
      'Run/40H_3, Run/extra')
with contextlib.redirect_stdout(io.StringIO()):
    res_bad = run(str(tmp / 'scribner'), str(tmp / 'out_left_bad'),
                  {'image_format': 'none', 'skipped_folders': 'Run/40H_3'})
check('skipped_folders that is not a list is ignored',
      'Folders left out' in res_bad['summary'][0], False)
# Names as the page sends them: a control character cannot go into a workbook
# cell, a non-string is not a name, and a folder that was read is not left out.
read_folder = build_timeline(str(tmp / 'scribner')).folders[0]
try:
    with contextlib.redirect_stdout(io.StringIO()):
        res_odd = run(str(tmp / 'scribner'), str(tmp / 'out_left_odd'),
                      {'image_format': 'none',
                       'skipped_folders': ['Run/40H\x01_9', None, 7, read_folder + '/']})
    odd = res_odd['summary'][0].get('Folders left out')
except Exception as e:                               # noqa: BLE001
    odd = f'run failed: {type(e).__name__}: {e}'
check('left-out names cleaned: control characters dropped, non-names and read folders ignored',
      odd, 'Run/40H_9')

out2 = tmp / 'out_noimg'
with contextlib.redirect_stdout(io.StringIO()):
    run(str(tmp / 'scribner'), str(out2), {'image_format': 'none'})
check('no images: workbook only', sorted(p.name for p in out2.rglob('*') if p.is_file()),
      ['fc_durability.xlsx'])

try:
    with contextlib.redirect_stdout(io.StringIO()):
        run(str(empty), str(tmp / 'out_empty'), {})
    true('empty upload raises', False)
except RuntimeError:
    true('empty upload raises', True)

# ── 6. OCV analysis ──────────────────────────────────────────────────
print('6. OCV analysis')

HOLD = '{}_S_OCV_95Cell-30RH_0.1WH2-0.2WA.csv'
DRY = '{}_S_OCV_95Cell-0RH_0.1WH2-0.2DA.csv'
rng = np.random.default_rng(7)


def write_ocv(path, programmed_s, dt, H0_s, v, i=None, drop=None, header=FOUR, extra=None,
              no_current=False, dup=False):
    """An OCV step whose voltage, and optionally current, are functions of
    hold-clock seconds, for a file that sits at H0_s on the clock. `drop`
    removes samples (a gap in the log) where it is true; a NaN current is
    written as a blank cell; `no_current` leaves the Current column out and
    `dup` writes every row twice. Returns the clock times and voltages as
    written, to work out expected values independently."""
    t = stamps(programmed_s, dt)
    H = H0_s + t
    keep = ~drop(H) if drop else np.ones(len(t), bool)
    vs = np.round(np.asarray(v(H), float), 6)
    cs = np.asarray(i(H), float) if i else np.zeros(len(t))
    if no_current:
        header = header.rsplit(',', 1)[0]
    head = header + (',' + ','.join(extra) if extra else '')
    pad = ',' * len(extra) if extra else ''
    rows = []
    for tt, vv, cc in zip(t[keep], vs[keep], cs[keep]):
        cur = '' if no_current else ',' + ('' if np.isnan(cc) else f'{cc:.6g}')
        rows += [f'{tt:.6f},1e-05,{vv:.6f}{cur}{pad}\n'] * (2 if dup else 1)
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    with open(path, 'w') as fh:
        fh.write(head + '\n')
        fh.writelines(rows)
    return H[keep], vs[keep]


def gas_step(d, step):
    write_four(Path(d) / f'{step}_S_OCV_95Cell-30RH_0.1WH2-0.1WN2.csv', 600, 10, v=0.1)


def xover(d, step):
    write_four(Path(d) / f'{step}_S_H2X_80.3Cell-100RH_0.4WH2-0.4WN2.csv', 720, 0.5,
               v=0.4, i=0.0185)


def flat(level):
    return lambda H: np.full(len(H), level)


def quiet(fn, *a, **k):
    with contextlib.redirect_stdout(io.StringIO()):
        return fn(*a, **k)


# The stretch finder and the rate fit on their own.
t1 = np.arange(0.0, 1800.0, 1.0)
check('1,800 samples at 1 s last 30 min', below_runs(t1, np.full(1800, 0.6), 0.7, 600.0),
      [(0, 1799, 1800.0)])
gap_t = np.r_[np.arange(0.0, 600.0), np.arange(1201.0, 1800.0)]
check('a gap longer than the window ends a stretch',
      [(a, b) for a, b, _ in below_runs(gap_t, np.full(len(gap_t), 0.6), 0.7, 600.0)],
      [(0, 599), (600, 1198)])
roll_t = np.r_[np.arange(0.0, 600.0), np.arange(631.0, 1800.0)]
check('a file rollover (31 s) does not', len(below_runs(roll_t, np.full(len(roll_t), 0.6), 0.7, 600.0)), 1)
check('isolated samples last one interval each',
      [s for _, _, s in below_runs(np.arange(10.0) * 10, np.where(np.arange(10) % 2, 0.6, 0.8),
                                   0.7, 600.0)], [10.0] * 5)
check('nothing under the threshold, no stretches', below_runs(t1, np.full(1800, 0.8), 0.7, 600.0), [])
check('no samples, no stretches', below_runs(np.empty(0), np.empty(0), 0.7, 600.0), [])
bins_h = (np.arange(240) + 0.5) / 60.0
rates = rolling_decay(bins_h, 0.9 - 0.001 * bins_h, 0.0, 4.0)
check('a hold shorter than the window gets one window over all of it',
      [(r['start'], r['end']) for r in rates], [(0.0, 4.0)])
close('... with the exact slope, positive while falling', rates[0]['decay'], 1000.0, 1e-6)
rates = rolling_decay((np.arange(720) + 0.5) / 60.0, np.full(720, 0.9), 0.0, 12.0)
check('windows end every hour from H0 + 10 h', [r['end'] for r in rates], [10.0, 11.0, 12.0])
sparse = rolling_decay((np.arange(0, 240, 2) + 0.5) / 60.0, np.full(120, 0.9), 0.0, 4.0)
true('data in half its minutes: no rate', np.isnan(sparse[0]['decay'])
     and abs(sparse[0]['coverage'] - 0.5) < 1e-9, sparse)

# Case 2: OCV drops below 0.7 V mid-hold 4 and stays.
death = tmp / 'death'
xover(death, 'b20')
T_DEATH = 12 * 3600 + 9000.0                     # 2.5 h into hold 4: 14.5 h
for j, step in enumerate(['c2', 'c4', 'c6', 'c8', 'c10']):
    if j:
        gas_step(death, f'c{2 * j + 1}')
    if j < 3:
        prof = lambda H: 0.90 - 1e-7 * H         # noqa: E731
    elif j == 3:
        prof = lambda H: np.where(H < T_DEATH, 0.90, 0.60)  # noqa: E731
    else:
        prof = flat(0.58)
    write_ocv(death / HOLD.format(step), 14400, 10, j * 14400.0, prof)
xover(death, 'c11')
res = analyze_ocv(build_timeline(str(death)))
check('case 2: five holds', len(res.holds), 5)
close('case 2: death within one bin of the built-in time', res.death_h, T_DEATH / 3600, 60 / 3600)
check('... in hold 4', (res.death.segment, res.death.step), (4, 'c8'))
check('... and later stretches never move it', [r.verdict for r in res.runs], ['death', 'after death'])
check('... no recovery', res.recovered_h, None)
check('decay rate from the latest window ending by the death, not the dead trace',
      res.decay_window_h, (8.0, 12.0))
close('... at the 360 μV/h the holds before it fell', res.decay_uV_h, 360.0, 0.5)
true('death note', any(n.startswith('OCV death at 14.50 h in hold 4 (c8)') for n in res.notes),
     res.notes)

# Case 4 and the guards: a transient at a hold's start, a 20-min dip, a
# noisy dry hold, a load pulse, and a logging gap between two dips.
dips = tmp / 'dips'
write_ocv(dips / HOLD.format('c2'), 14400, 10, 0.0, lambda H: np.where(H < 300, 0.30, 0.92))
gas_step(dips, 'c3')
write_ocv(dips / HOLD.format('c4'), 14400, 10, 14400.0,
          lambda H: np.where((H >= 14400 + 5400) & (H < 14400 + 6600), 0.60, 0.91))
gas_step(dips, 'c5')
write_ocv(dips / DRY.format('c6'), 14400, 10, 28800.0,
          lambda H: 0.76 + rng.uniform(-0.08, 0.08, len(H)))
gas_step(dips, 'c7')


def on_load(H):
    return (H >= 43200 + 3600) & (H < 43200 + 6000)


write_ocv(dips / HOLD.format('c8'), 14400, 10, 43200.0,
          lambda H: np.where(on_load(H), 0.55, 0.91), i=lambda H: np.where(on_load(H), 0.5, 0.0))
gas_step(dips, 'c9')


def in_dip(H):
    x = H - 57600
    return ((x >= 3600) & (x < 4800)) | ((x >= 5700) & (x < 6900))


write_ocv(dips / HOLD.format('c10'), 14400, 10, 57600.0,
          lambda H: np.where(in_dip(H), 0.60, 0.91),
          drop=lambda H: (H - 57600 >= 4800) & (H - 57600 < 5700))
res = analyze_ocv(build_timeline(str(dips)))
check('case 4: five holds, no death', (len(res.holds), res.death), (5, None))
h1, h2, h3, h4, h5 = res.holds
check('case 4: 5 min under 0.7 V at a hold start is never judged', (h1.runs, h1.settling), ([], 59))
close('... and the settled start lies past it', h1.start_v, 0.92, 1e-9)
check('a 20-min dip is a dip, not a death', [(r.verdict, round(r.minutes)) for r in h2.runs],
      [('dip', 20)])
close('... dated at its start', h2.runs[0].start_h, (14400 + 5400) / 3600, 60 / 3600)
check('noisy dry hold: the median never trips', h3.runs, [])
true('... though raw samples fall under 0.7 V', h3.raw_below > 100, h3.raw_below)
check('a load pulse is not judged', (h4.on_load, h4.runs), (240, []))
check('a logging gap longer than the window splits a stretch',
      [(r.verdict, round(r.minutes)) for r in h5.runs], [('dip', 20), ('dip', 20)])
true('dip note', any(n.startswith('hold 2 (c4): 10-min median under 0.7 V 1 time(s)')
                     for n in res.notes), res.notes)
true('load note', any(n.startswith('hold 4 (c8): 240 samples drew more than 1 mA')
                      for n in res.notes), res.notes)
check('two conditions, the first hold the reference', [r['Reference'] for r in res.by_condition],
      ['yes', ''])

# Case 5: a hold split across two files, dying across the rollover.
roll = tmp / 'rollover'
write_ocv(roll / HOLD.format('c2'), 14400, 10, 0.0, flat(0.90))
gas_step(roll, 'c3')
T_DIE = 14400 + 13170.0                          # 20 min before c4's log ends
write_ocv(roll / HOLD.format('c4'), 14400, 10, 14400.0, lambda H: np.where(H < T_DIE, 0.90, 0.60))
write_ocv(roll / HOLD.format('c5'), 14400, 10, 28800.0,
          lambda H: np.where(H < 28800 + 1200, 0.60, 0.90))
gas_step(roll, 'c6')
write_ocv(roll / HOLD.format('c7'), 7200, 10, 43200.0, flat(0.90))
tl = build_timeline(str(roll))
res = analyze_ocv(tl)
check('case 5: a hold split across two files is one hold',
      [[f.step for f in s.files] for s in tl.segments], [['c2'], ['c4', 'c5'], ['c7']])
close('... its durations summed', tl.segments[1].H_end_h - tl.segments[1].H_start_h, 8.0, 1e-9)
check('... settling timed once, from the start of the hold', res.holds[1].settling, 59)
close('case 5: death across the rollover, dated at its start', res.death_h, T_DIE / 3600, 60 / 3600)
close('... a 40 min stretch', res.death.minutes, 40.4, 0.3)
close('recovery flagged when the median comes back', res.recovered_h, (28800 + 1200) / 3600, 0.1)
true('... and the death stands', any('Flagged; the death at 7.66 h stands' in n for n in res.notes),
     res.notes)
comeback = tmp / 'comeback'
write_ocv(comeback / HOLD.format('c2'), 14400, 10, 0.0, lambda H: np.where(H < 7200, 0.90, 0.60))
gas_step(comeback, 'c3')
write_ocv(comeback / HOLD.format('c4'), 14400, 10, 14400.0, flat(0.90))
res = analyze_ocv(build_timeline(str(comeback)))
close('a death that holds to the end of its hold', res.death_h, 2.0, 60 / 3600)
close('... is flagged as recovered at the next hold\'s first judged sample', res.recovered_h,
      (14400 + 600) / 3600, 1e-9)
tr = load_hold_trace(tl.segments[1])
check('stitched trace: one rollover, the unlogged tail its only gap',
      (int(np.sum(np.diff(tr.file_idx) != 0)), round(float(np.max(np.diff(tr.t_s))), 6)), (1, 40.0))
tr = load_hold_trace(build_timeline(str(late)).segments[0])
check('a log that started late is placed by its own span', (tr.t_s[0], tr.t_s[-1]), (0.0, 36000.0))
res = analyze_ocv(build_timeline(str(dips)), OcvSettings(settle_min=0))
check('settle_min 0 judges the start transient, as a dip', [r.verdict for r in res.holds[0].runs],
      ['dip'])

# The median, not single samples: noise around 0.67 V still trips it.
noisy = tmp / 'noisy'
write_ocv(noisy / HOLD.format('c2'), 14400, 10, 0.0,
          lambda H: np.where(H < 7200, 0.90, 0.67 + rng.uniform(-0.06, 0.06, len(H))))
res = analyze_ocv(build_timeline(str(noisy)))
true('noisy death: the median trips though raw samples cross back over 0.7 V',
     res.death is not None and 0 <= res.death_h - 2.0 <= 5 / 60, res.death_h)

# Exact values on a linear decline over three files, with a start overshoot
# the settle window must skip and Excel columns in the middle file.
lin = tmp / 'linear'
prof = lambda H: np.where(H < 300, 0.98, 0.95 - 0.001 * H / 3600)  # noqa: E731
written = [write_ocv(lin / HOLD.format(f'c{5 + j}'), 36000, 30, j * 36000.0, prof,
                     extra=['cd', '', ''] if j == 1 else None) for j in range(3)]
H = np.concatenate([w[0] for w in written])
V = np.concatenate([w[1] for w in written])
res = analyze_ocv(build_timeline(str(lin)))
h = res.holds[0]
Hj, Vj = H[H >= 600], V[H >= 600]
start, end = np.median(Vj[Hj < Hj[0] + 1800]), np.median(Vj[Hj >= Hj[-1] - 1800])
close('settled start: median of the first 30 min after the settle window', h.start_v, start, 1e-9)
close('end: median of the last 30 min', h.end_v, end, 1e-9)
close('loss', res.loss_pct, (start - end) / start * 100, 1e-6)
check('settling counted once over three files', (h.settling, h.judged),
      (int((H < 600).sum()), int((H >= 600).sum())))
close('lowest 10-min median: the centred window at the last sample', res.lowest_v,
      np.median(Vj[Hj > Hj[-1] - 300]), 1e-9)
check('windows end every hour from 10 h to 30 h', [round(r['end'], 6) for r in h.rates],
      [float(e) for e in range(10, 31)])
true('every window fits the 1,000 μV/h slope', all(abs(r['decay'] - 1000) < 0.5 for r in h.rates),
     [round(r['decay'], 3) for r in h.rates])
close('OCV decay rate is the last window', res.decay_uV_h, 1000.0, 0.5)
check('... which spans 20-30 h', tuple(round(x, 6) for x in res.decay_window_h), (20.0, 30.0))

# Loss and rate follow the first hold's conditions, even when the run ends
# on another (the FCTS cycle ends on its dry hold).
multi = tmp / 'multi'
for j, (pat, level) in enumerate([(HOLD, 0.92), (DRY, 0.89), (HOLD, 0.91), (DRY, 0.88)]):
    if j:
        gas_step(multi, f'c{2 * j + 1}')
    write_ocv(multi / pat.format(f'c{2 * j + 2}'), 14400, 10, j * 14400.0, flat(level))
res = analyze_ocv(build_timeline(str(multi)))
check('reference: the first hold\'s conditions', res.reference, '95 °C, 30% RH, H2/Air 0.1/0.2 slpm')
check('initial and final at the reference conditions', (res.initial_v, res.final_v), (0.92, 0.91))
close('reference loss', res.loss_pct, 1 / 0.92, 1e-9)
check('every condition in its own row',
      [(r['Conditions'], r['Holds'], r['First settled start (V)'], r['Last end (V)'])
       for r in res.by_condition],
      [('95 °C, 30% RH, H2/Air 0.1/0.2 slpm', 2, 0.92, 0.91),
       ('95 °C, 0% RH, H2/dry Air 0.1/0.2 slpm', 2, 0.89, 0.88)])
close('flat 4 h holds decay at 0 μV/h', res.decay_uV_h, 0.0, 1e-6)
true('reference note', any("first hold's conditions" in n for n in res.notes), res.notes)
# The headline rate is the reference conditions' too: this run ends on a dry
# hold falling 2 mV/h while its reference hold is flat.
mr = tmp / 'multirate'
write_ocv(mr / HOLD.format('c2'), 14400, 10, 0.0, flat(0.92))
gas_step(mr, 'c3')
write_ocv(mr / DRY.format('c4'), 14400, 10, 14400.0, lambda H: 0.89 - 2e-3 * (H - 14400.0) / 3600)
res = analyze_ocv(build_timeline(str(mr)))
check("the headline rate is the reference conditions', not the last hold's",
      (round(res.decay_uV_h, 1),
       [(r['Reference'], r['Latest decay rate (μV/h)']) for r in res.by_condition]),
      (0.0, [('yes', 0.0), ('', 2000.0)]))

bad = tmp / 'badcol'
gas_step(bad, 'c1')
write_ocv(bad / HOLD.format('c2'), 14400, 10, 0.0, flat(0.9),
          header='Time Stamp,Counter Electrode Voltage,Cell Voltage,Current')
res = analyze_ocv(build_timeline(str(bad)))
check('a hold file without a voltage column still counts its hours',
      (len(res.holds), res.holds[0].judged, res.death), (1, 0, None))
true('... and says it could not be read',
     any("could not be read (ValueError: no 'Working Electrode Voltage' column)" in n
         for n in res.notes), res.notes)
check('... and that the rule could not be applied', (res.judged, res.notes[-1]),
      (False, 'no hold has OCV to judge, so the OCV death rule could not be applied'))

# Samples whose current is unknown are not judged; a file with no current
# column at all is taken as at open circuit, and says so.
blank = tmp / 'blank'
write_ocv(blank / HOLD.format('c2'), 14400, 10, 0.0,
          lambda H: np.where((H >= 3600) & (H < 6000), 0.30, 0.92),
          i=lambda H: np.where((H >= 3600) & (H < 6000), np.nan, 0.0))
gas_step(blank, 'c3')
write_ocv(blank / HOLD.format('c4'), 14400, 10, 14400.0, flat(0.92), no_current=True)
res = analyze_ocv(build_timeline(str(blank)))
check('blank current readings are not judged', (res.death, res.holds[0].not_finite), (None, 240))
true('... with a note', any(n.startswith('hold 1 (c2): 240 samples lack a time, voltage or current')
                            for n in res.notes), res.notes)
check('a file without a current column is judged as at open circuit',
      (res.holds[1].judged, res.holds[1].on_load), (len(stamps(14400, 10)) - 59, 0))
true('... with a note', any('has no Current column' in n for n in res.notes), res.notes)

# Duplicated rows do not shorten a stretch: 181 samples 10 s apart under
# 0.7 V make a 30 min median stretch whether or not each row is doubled.
for doubled in (False, True):
    d = tmp / f'dup_{doubled}'
    write_ocv(d / HOLD.format('c2'), 14400, 10, 0.0,
              lambda H: np.where((H >= 3600) & (H <= 5400), 0.60, 0.90), dup=doubled)
    res = analyze_ocv(build_timeline(str(d)))
    check(f'a 30 min stretch is a death, rows doubled {doubled}',
          (res.death is not None and round(res.death.minutes, 6), res.death_h), (30.0, 1.0))

# OCV initial falls back to a later reference hold with data, and says so.
fb = tmp / 'fallback'
write_ocv(fb / HOLD.format('c2'), 14400, 10, 0.0, flat(0.9),
          header='Time Stamp,Counter Electrode Voltage,Cell Voltage,Current')
gas_step(fb, 'c3')
write_ocv(fb / DRY.format('c4'), 14400, 10, 14400.0, flat(0.89))
gas_step(fb, 'c5')
write_ocv(fb / HOLD.format('c6'), 14400, 10, 28800.0, flat(0.91))
res = analyze_ocv(build_timeline(str(fb)))
check('initial and final from the reference hold that has data',
      (res.initial_v, res.initial_from.segment.number, res.final_from.segment.number), (0.91, 3, 3))
true('... with a note', any(n.startswith('OCV initial is the settled start of hold 3 (c6)')
                            for n in res.notes), res.notes)

# run(): plot, readout, workbook and summary on the death and the dip runs.
out = tmp / 'out_death'
res = quiet(run, str(death), str(out), {'image_format': 'png'})
row = res['summary'][0]
close('summary: OCV death', row.get('OCV death'), T_DEATH / 3600, 60 / 3600)
close('run() returns it', res['ocv_death_h'], T_DEATH / 3600, 60 / 3600)
true('summary row is JSON-safe', json.dumps(row, allow_nan=False) is not None)
true('OCV plot', (out / 'fcdur_ocv_vs_time.png').exists())
sc = json.loads((out / '_plot_data' / 'fcdur_ocv_vs_time.json').read_text(), parse_constant=_no_nan)
check('OCV sidecar plot type', sc['plot_type'], 'fcdur_ocv_vs_time')
vals = extract_values(sc)
close('readout: OCV death', (vals.get('OCV death') or {}).get('value'), T_DEATH / 3600, 0.006)
close('readout: OCV initial as in the summary', (vals.get('OCV initial') or {}).get('value'),
      row['OCV initial'], 5e-5)
top = sc['data']['axes'][0]
true('threshold and death lines', any(abs(a['y'] - 0.7) < 1e-12 for a in top['axhlines'])
     and any(abs(a['x'] - res['ocv_death_h']) < 1e-3 for a in top['axvlines']), top['axvlines'])
true('characterization ticks at 0 h and 20 h', {round(a['x'], 4) for a in top['axvlines']} >= {0.0, 20.0},
     top['axvlines'])
wb = openpyxl.load_workbook(out / 'fc_durability.xlsx', read_only=True)
check('events sheet', [(r[0], r[6]) for r in wb['OCV events'].iter_rows(min_row=2, values_only=True)],
      [(4, 'death'), (5, 'after death')])
holds_sheet = list(wb['Holds'].iter_rows(values_only=True))
col = holds_sheet[0].index('OCV death (h)')
check('holds sheet marks the death', [r[col] for r in holds_sheet[1:]],
      [None, None, None, round(res['ocv_death_h'], 4), None])
wb.close()

# The OCV plot's bin width: automatic, or set on the form (ocv_plot_bin_min).
# Only the plot and its stored copy follow it; every rule works on the samples.
check('plot bin, automatic: the finest the 3,000-bin cap allows', ocv_plot_bin(150.0),
      (180.0, 'automatic: the finest that keeps about 3,000 bins over the test, at least 1 min'))
check('... never under 1 min', ocv_plot_bin(5.0)[0], 60.0)
check('... and a NaN or negative width is automatic too',
      [ocv_plot_bin(150.0, x)[0] for x in (float('nan'), -5.0)], [180.0, 180.0])
check('... in whole RH cycles when the holds cycle RH: one 75 s cycle over 41.69 h',
      ocv_plot_bin(41.689, 0.0, 75.0),
      (75.0, 'automatic: 1 whole RH cycle of 75 s, at least 1 min and no more than about 3,000 '
             'bins over the test'))
check('... the fewest whole cycles at least the cap\'s width, or 1 min',
      [ocv_plot_bin(h, 0.0, period)[0] for h, period in ((500.0, 75.0), (20.0, 40.0), (20.0, 75.0))],
      [600.0, 80.0, 75.0])
check('... while a cycle is no more than two of the finest bins; a slower one is cut into equal '
      'bins at least that fine, and said so',
      ([ocv_plot_bin(h, 0.0, period)[0] for h, period in ((20.0, 120.0), (20.0, 121.0), (20.0, 179.0),
                                                           (150.0, 7200.0))],
       ocv_plot_bin(150.0, 0.0, 7200.0)[1]),
      ([120.0, 60.5, 89.5, 180.0],
       'automatic: 40 bins to each RH cycle of 7200 s, at least 1 min and no more than about '
       '3,000 bins over the test'))
check('... while a width set on the form is used as set, cycles or not',
      ocv_plot_bin(41.689, 5.0, 75.0), (300.0, 'as set'))
check('plot bin set: used as set', ocv_plot_bin(150.0, 5.0), (300.0, 'as set'))
bin_s, why = ocv_plot_bin(150.0, 1.0)
check(f'... raised to keep about {PLOT_BINS_CAP:,} bins at most, and says so',
      (bin_s, why), (150 * 3600.0 / PLOT_BINS_CAP,
                     'raised from the 1 min set: bins are at least 1 min, and no more than about '
                     '3,000 are plotted and stored'))
check('... raised to 1 min at least', (ocv_plot_bin(20.0, 0.5)[0], ocv_plot_bin(20.0, 0.5)[1][:7]),
      (60.0, 'raised '))
check('... and a width exactly at the cap is as set', ocv_plot_bin(50.0, 1.0), (60.0, 'as set'))
# 2 min on a 100.02 h test is a hair too fine for the cap: raised at the four
# figures the outputs print, so the note never reads "2 min, raised from the
# 2 min set", and the width it prints, typed back, is taken as set.
shown = []
for hours, asked in ((100.02, 2.0), (51.2, 1.0), (163.4, 1.0)):
    bin_s, why = ocv_plot_bin(hours, asked)
    shown.append((f'{bin_s / 60:.4g}', why[:6], ocv_plot_bin(hours, float(f'{bin_s / 60:.4g}'))))
check('... a raise past the width set shows in the minutes printed, which typed back hold',
      [(s, w, round(b2, 6), why2) for s, w, (b2, why2) in shown],
      [('2.001', 'raised', 120.06, 'as set'), ('1.024', 'raised', 61.44, 'as set'),
       ('3.268', 'raised', 196.08, 'as set')])
check('... nor is a width typed with more figures than it prints, a hair under the cap',
      ocv_plot_bin(51.2, 1.02399999999)[1], 'as set')
bin_s, why = ocv_plot_bin(51.67, 1.033999)
check('... and one raised by a little says the width set with all its figures',
      (f'{bin_s / 60:.4g}', why.split(':')[0]), ('1.034', 'raised from the 1.033999 min set'))


def ocv_plot_points(out_dir):
    """Median points in the stored OCV plot, and the bin its readout names."""
    sc = json.loads((Path(out_dir) / '_plot_data' / 'fcdur_ocv_vs_time.json').read_text(),
                    parse_constant=_no_nan)
    n = sum(len(ln['x']) for ln in sc['data']['axes'][0]['lines']
            if ln.get('label') and ln['label'] not in (BAND_LOW, BAND_HIGH))
    return n, extract_values(sc).get('OCV plot bin')


def workbook_sheets(out_dir):
    book = openpyxl.load_workbook(Path(out_dir) / 'fc_durability.xlsx', read_only=True)
    got = {name: list(book[name].iter_rows(values_only=True)) for name in book.sheetnames}
    book.close()
    return got


def plot_bin_row(sheets):
    return next((r[1] for r in sheets['Summary'] if r[0] == 'OCV plot bin (min)'), '')


auto_n, auto_bin = ocv_plot_points(out)
out1 = tmp / 'out_death_bin1'
res1 = quiet(run, str(death), str(out1), {'image_format': 'png', 'ocv_plot_bin_min': '2'})
n1, bin1 = ocv_plot_points(out1)
check('plot bin: automatic over the 20 h death run is 1 min, five holds of 230 bins',
      (auto_bin, auto_n), ({'value': 1.0, 'unit': 'min'}, 1150))
check('... set to 2 min on the form: half the points, and the readout says 2 min',
      (bin1, 2 * n1), ({'value': 2.0, 'unit': 'min'}, auto_n))
book_auto, book1 = workbook_sheets(out), workbook_sheets(out1)
check('... the workbook says which bin, and why',
      [plot_bin_row(b).split(';')[0] for b in (book_auto, book1)],
      ['1 (automatic: the finest that keeps about 3,000 bins over the test, at least 1 min)',
       '2 (as set)'])
for b in (book_auto, book1):
    b['Summary'] = [r for r in b['Summary'] if r[0] != 'OCV plot bin (min)']
check('... and nothing else moves: the summary row, the notes and every sheet are the same',
      (res1['summary'], res1['notes'], book1), (res['summary'], res['notes'], book_auto))
res_r = quiet(run, str(death), str(tmp / 'out_death_bin_r'),
              {'image_format': 'none', 'ocv_plot_bin_min': '0.5'})
check('a bin under 1 min is raised to 1 min, with a note and in the workbook',
      ([n for n in res_r['notes'] if n.startswith('OCV plot bin')],
       plot_bin_row(workbook_sheets(tmp / 'out_death_bin_r')).split(';')[0]),
      (['OCV plot bin 1 min, raised from the 0.5 min set: bins are at least 1 min, and no more '
        'than about 3,000 are plotted and stored'],
       '1 (raised from the 0.5 min set: bins are at least 1 min, and no more than about 3,000 '
       'are plotted and stored)'))
# The widest bin the form allows, a day, leaves each 4 h hold one bin.
out_day = tmp / 'out_death_bin_day'
quiet(run, str(death), str(out_day), {'image_format': 'png', 'ocv_plot_bin_min': '1440'})
check('the widest bin the form allows reads as a plain number',
      plot_bin_row(workbook_sheets(out_day)).split(';')[0], '1440 (as set)')
day_lines = json.loads((out_day / '_plot_data' / 'fcdur_ocv_vs_time.json').read_text(),
                       parse_constant=_no_nan)['data']['axes'][0]['lines']
check('... a hold left one bin is drawn as a marker, so the plot is not empty',
      sorted({(len(ln['x']), ln.get('marker')) for ln in day_lines}), [(1, 'o')])

# The high–low band: each plot bin's lowest and highest judged sample, shaded
# around the median. One 10-s sample 50 mV high, one 40 mV low, in a flat hold:
# the band's edges keep both whole, the median does not move, and neither
# reaches the rule (shorter than the 10-min median window).
import pandas as pd
from scripts.fc_durability_analysis import plot_ocv
spike = tmp / 'spike'
write_ocv(spike / HOLD.format('c2'), 14400, 10, 0.0,
          lambda H: np.where(H == 7200, 0.95, np.where(H == 10800, 0.86, 0.90)))
gas_step(spike, 'c3')
write_ocv(spike / HOLD.format('c4'), 14400, 10, 14400.0, flat(0.89))
tl_s = build_timeline(str(spike))
res_s = analyze_ocv(tl_s)
p1 = res_s.holds[0].plot
at = lambda h: int(np.argmin(np.abs(p1['t_h'].to_numpy() - h)))  # noqa: E731
check('band: a one-sample spike and dip are the bin\'s high and low, the median flat',
      [(round(p1['max'][at(h)], 5), round(p1['median'][at(h)], 5), round(p1['min'][at(h)], 5))
       for h in (2.0, 3.0)],
      [(0.95, 0.9, 0.9), (0.9, 0.9, 0.86)])
check('... every other bin is flat, and the rule sees neither', (
    sorted({round(v, 5) for v in p1['max'].drop(at(2.0))}),
    sorted({round(v, 5) for v in p1['min'].drop(at(3.0))}), res_s.death, res_s.runs),
    ([0.9], [0.9], None, []))
# Against the samples themselves, bin by bin: the death run's dying hold.
tr4 = load_hold_trace(build_timeline(str(death)).segments[3], OcvSettings())
res_d = analyze_ocv(build_timeline(str(death)))
k = np.floor((tr4.t_s[tr4.valid] - tr4.segment.H_start_h * 3600.0) / res_d.plot_bin_s)
want = pd.Series(tr4.v[tr4.valid]).groupby(k).agg(['min', 'max'])
got = res_d.holds[3].plot
check('band: each bin\'s low and high are its judged samples\' min and max',
      (np.allclose(got['min'].to_numpy(), want['min'].to_numpy()),
       np.allclose(got['max'].to_numpy(), want['max'].to_numpy()), len(got) == len(want)),
      (True, True, True))
fig_s = plot_ocv(tl_s, res_s)
ax_s = fig_s.axes[0]
from matplotlib.collections import PolyCollection
legend_s = [t.get_text() for t in ax_s.get_legend().get_texts()]
check('... drawn as one shaded band per hold, with one key and no edge entries in the legend',
      (sum(isinstance(c, PolyCollection) for c in ax_s.collections),
       [t for t in legend_s if t.startswith('high–low')],
       [t for t in legend_s if t in (BAND_LOW, BAND_HIGH, 'bin minimum')]),
      (2, ['high–low of the samples in each 1-min bin'], []))
ylo, yhi = ax_s.get_ylim()
true('... the y range takes the whole band in', ylo < 0.86 and yhi > 0.95, (ylo, yhi))
_plt.close(fig_s)
# The stored copy keeps the band as its two edges, labelled, on every hold's
# x, so the portal can shade between them.
out_s = tmp / 'out_spike'
quiet(run, str(spike), str(out_s), {'image_format': 'png'})
lines_s = json.loads((out_s / '_plot_data' / 'fcdur_ocv_vs_time.json').read_text(),
                     parse_constant=_no_nan)['data']['axes'][0]['lines']
meds = [ln for ln in lines_s if ln.get('label') and ln['label'] not in (BAND_LOW, BAND_HIGH)]
lows = [ln for ln in lines_s if ln.get('label') == BAND_LOW]
highs = [ln for ln in lines_s if ln.get('label') == BAND_HIGH]
check('stored: a low and a high edge per hold, on its median\'s x, around it',
      (len(meds), len(lows), len(highs),
       all(lo['x'] == m['x'] == hi['x'] for lo, m, hi in zip(lows, meds, highs)),
       all(a <= b <= c for lo, m, hi in zip(lows, meds, highs)
           for a, b, c in zip(lo['y'], m['y'], hi['y'])),
       max(max(hi['y']) for hi in highs), min(min(lo['y']) for lo in lows)),
      (2, 2, 2, True, True, 0.95, 0.86))

# The portal's redraw (Recall's run detail, and the comparison Export renders)
# pairs each low edge with the high on the same x and shades between them.
from scripts.helpers.plot_compare import band_pairs, render_overlay_comparison
from scripts.compare_polcurves import export_comparison_excel
pairs, loose = band_pairs(lines_s)
check('band_pairs: each hold\'s low with its high, none left over',
      (len(pairs), [lo is a and hi is b for (lo, hi), a, b in zip(pairs, lows, highs)], loose),
      (2, [True, True], []))
lo1 = {'label': BAND_LOW, 'x': [0, 1], 'y': [1, 1]}
lo2 = {'label': BAND_LOW, 'x': [0, 1], 'y': [0, 0]}
hi_off = {'label': BAND_HIGH, 'x': [0, 1.5], 'y': [2, 2]}
hi_short = {'label': BAND_HIGH, 'x': [0, 1], 'y': [2]}
hi1 = {'label': BAND_HIGH, 'x': [0, 1], 'y': [2, 2]}
pairs, loose = band_pairs([lo1, hi_off, hi_short, hi1, lo2])
check('... a high on other x or short of points pairs with nothing, and a second low waits '
      'for a high of its own',
      ([(lo is lo1, hi is hi1) for lo, hi in pairs], [id(t) for t in loose]),
      ([(True, True)], [id(lo2), id(hi_off), id(hi_short)]))
sc_s = json.loads((out_s / '_plot_data' / 'fcdur_ocv_vs_time.json').read_text(), parse_constant=_no_nan)


def redrawn(band):
    """What the redraw of the stored spike plot draws on its OCV axes."""
    item = {'label': 'spike', 'sidecar': sc_s}
    if band != 'unsaid':
        item['band'] = band
    with contextlib.redirect_stdout(io.StringIO()):
        fig = render_overlay_comparison([item])
    ax = fig.axes[0]
    fills = [c for c in ax.collections if isinstance(c, PolyCollection)]
    ys = [v for f in fills for p in f.get_paths() for v in p.vertices[:, 1]]
    edges = [ln for ln in ax.get_lines() if ln.get_alpha() == 0.3]
    medians = [ln for ln in ax.get_lines() if abs(ln.get_linewidth() - 1.4 * 0.9) < 1e-9]
    got = (len(fills), (round(min(ys), 5), round(max(ys), 5)) if ys else None,
           len(edges), all(any(ln.get_ydata().tolist() == e['y'] for e in lows + highs)
                           for ln in edges),
           len(medians), sorted({(ln.get_linestyle(), ln.get_alpha()) for ln in medians}),
           [t.get_text() for t in fig.legends[0].get_texts()])
    _plt.close(fig)
    return got


check('redraw: each hold\'s band shaded from its stored edges, the edges faint over it, the '
      'medians solid, one key for the band',
      redrawn('unsaid'), (2, (0.86, 0.95), 4, True, 2, [('-', 1.0)], ['spike', 'high–low of each bin']))
check('... the same with the band asked for', redrawn(True), redrawn('unsaid'))
check('... and with it turned off, no shading, no edges and no key, the medians as before',
      redrawn(False), (0, None, 0, True, 2, [('-', 1.0)], ['spike']))
# An edge left without a partner (here the first hold's high taken out) has
# no shape to shade: it is drawn faint like any edge, never as a series, and
# not at all with the band off; its hold's line is styled as any line.
loose_sc = json.loads(json.dumps(sc_s))
loose_ax = loose_sc['data']['axes'][0]
loose_ax['lines'].remove(next(ln for ln in loose_ax['lines'] if ln.get('label') == BAND_HIGH))
loose_y = next(ln for ln in loose_ax['lines'] if ln.get('label') == BAND_LOW)['y']


def loose_drawn(band):
    with contextlib.redirect_stdout(io.StringIO()):
        fig = render_overlay_comparison([{'label': 'loose', 'sidecar': loose_sc, 'band': band}])
    ax = fig.axes[0]
    got = (sum(isinstance(c, PolyCollection) for c in ax.collections),
           sorted(ln.get_alpha() for ln in ax.get_lines()
                  if np.asarray(ln.get_ydata(), float).tolist() == loose_y))
    _plt.close(fig)
    return got


check('redraw: an edge without a partner drawn faint, once, and not at all with the band off',
      (loose_drawn(True), loose_drawn(False)), ((1, [0.3]), (0, [])))
# The shading goes under the markers it surrounds (the settled start and end
# values are scatters, which the redraw draws at zorder 1).
with contextlib.redirect_stdout(io.StringIO()):
    fig3 = render_overlay_comparison([{'label': 'spike', 'sidecar': sc_s}])
fills3 = [c.get_zorder() for c in fig3.axes[0].collections if isinstance(c, PolyCollection)]
marks3 = [c.get_zorder() for c in fig3.axes[0].collections if type(c).__name__ == 'PathCollection']
true('redraw: the shading drawn under the markers', fills3 and marks3 and max(fills3) < min(marks3),
     (fills3, marks3))
_plt.close(fig3)
# Two durability tests compared stack both readouts under the plot, about
# twenty lines: the box must clear the legend below it.
with contextlib.redirect_stdout(io.StringIO()):
    fig2 = render_overlay_comparison([{'label': 'one', 'sidecar': sc_s},
                                      {'label': 'two', 'sidecar': sc_s}], title='OCV')
fig2.canvas.draw()
rdr = fig2.canvas.get_renderer()
boxes = [t.get_bbox_patch().get_window_extent(rdr) for t in fig2.texts if t.get_bbox_patch()]
leg = fig2.legends[0].get_window_extent(rdr)
true('comparison: two tests\' readouts clear the legend under them',
     len(boxes) == 1 and boxes[0].y0 > leg.y1, (boxes and boxes[0].y0, leg.y1))
_plt.close(fig2)


def trace_labels(path):
    """The comparison workbook's first trace sheet: {sample: [trace labels]}."""
    wb = openpyxl.load_workbook(path)
    ws = next(w for w in wb.worksheets if w.title != 'Metrics')
    got, cur = {}, None
    for c in range(1, ws.max_column + 1):
        if ws.cell(1, c).value is not None:
            cur = ws.cell(1, c).value
            got[cur] = []
        if ws.cell(2, c).value is not None and cur is not None:
            got[cur].append(ws.cell(2, c).value)
    wb.close()
    return got


with contextlib.redirect_stdout(io.StringIO()):
    export_comparison_excel([{'label': 'on', 'sidecar': sc_s, 'band': True},
                             {'label': 'off', 'sidecar': sc_s, 'band': False}],
                            'fcdur_ocv_vs_time', str(tmp / 'band_cmp.xlsx'))
wl = trace_labels(tmp / 'band_cmp.xlsx')
edge = lambda labs: [t for t in labs if t in (BAND_LOW, BAND_HIGH)]  # noqa: E731
check('comparison workbook: a sample with its band keeps the edges as traces, one without '
      'leaves them out, the other traces alike',
      (sorted(wl), edge(wl.get('on', [])), edge(wl.get('off', [])),
       [t for t in wl.get('on', []) if t not in (BAND_LOW, BAND_HIGH)] == wl.get('off')),
      (['off', 'on'], [BAND_LOW, BAND_HIGH] * 2, [], True))

out = tmp / 'out_dips'
res = quiet(run, str(dips), str(out), {'image_format': 'png'})
row = res['summary'][0]
true('alive: no OCV death in the summary', 'OCV death' not in row and res['ocv_death_h'] is None, row)
check('summary carries the OCV numbers', sorted(k for k in row if k.startswith('OCV')),
      ['OCV decay rate', 'OCV final', 'OCV initial', 'OCV loss', 'OCV min'])
sc = json.loads((out / '_plot_data' / 'fcdur_ocv_vs_time.json').read_text(), parse_constant=_no_nan)
check('readout says no death', extract_values(sc).get('OCV death'), 'none')
pts = sum(len(ln['x']) for ln in sc['data']['axes'][0]['lines'])
# Three lines a hold (median, and the band's low and high edges), each of
# about PLOT_BINS_CAP bins over the test at most, plus one bin a hold.
true('plot within its point budget', pts <= 3 * (PLOT_BINS_CAP + 5) + 50, pts)
true('dip notes reach run()', any('not a death' in n for n in res['notes']), res['notes'])
res = quiet(run, str(dips), str(tmp / 'out_dips_p'), {'image_format': 'none', 'ocv_persist_min': '15'})
close('ocv_persist_min reaches the rule: the 20-min dip is now a death', res['ocv_death_h'],
      (14400 + 5400) / 3600, 60 / 3600)
res = quiet(run, str(dips), str(tmp / 'out_dips_q'),
            {'image_format': 'none', 'ocv_persist_min': 'abc', 'ocv_window_min': '-5',
             'ocv_threshold': '700', 'geo_area': '-5', 'rate_window_h': '1e300',
             'ocv_plot_bin_min': '-1'})
check('unreadable or out-of-range parameters fall back to the defaults',
      (res['ocv_death_h'], res['summary'][0]['geo_area']), (None, 5.0))
check('... each with a note', [n.split(' = ')[0] for n in res['notes'] if n.startswith('parameter ')],
      ['parameter geo_area', 'parameter ocv_threshold', 'parameter ocv_window_min',
       'parameter ocv_persist_min', 'parameter rate_window_h', 'parameter ocv_plot_bin_min'])

# Pairs of markers or rates never come back from the sidecar as reference
# lines: flat holds, and 10.425 h holds with two equal rate windows.
for name, src in (('multi', multi), ('fcts', tmp / 'fcts')):
    out = tmp / f'out_{name}'
    quiet(run, str(src), str(out), {'image_format': 'png'})
    sc = json.loads((out / '_plot_data' / 'fcdur_ocv_vs_time.json').read_text(),
                    parse_constant=_no_nan)
    check(f'{name}: only the threshold and zero lines are horizontal',
          [len(a['axhlines']) for a in sc['data']['axes']], [1, 1])
true('settled markers are kept as points', len(sc['data']['axes'][0]['scatters']) >= 1)
# Each hold's median line names its conditions, so a stored copy can be read
# per condition (Comparisons): six FCTS holds, two conditions alternating.
named_lines = [ln.get('label') for ln in sc['data']['axes'][0]['lines']
               if ln.get('label') and ln.get('label') not in (BAND_LOW, BAND_HIGH)]
check('fcts: every hold\'s OCV line labelled with its conditions',
      (len(named_lines), sorted(set(named_lines))),
      (6, ['90 °C, 0% RH, H2/dry Air 0.1/0.1 slpm', '90 °C, 30% RH, H2/Air 0.1/0.1 slpm']))
# ... while the legend still names each condition once.
from scripts.fc_durability_analysis import plot_ocv
_tl = build_timeline(str(tmp / 'fcts'))
_fig = plot_ocv(_tl, analyze_ocv(_tl))
_legend = [t.get_text() for t in _fig.axes[0].get_legend().get_texts()]
check('... and the legend names each condition once',
      sorted(t for t in _legend if ' °C' in t),
      ['90 °C, 0% RH, H2/dry Air 0.1/0.1 slpm', '90 °C, 30% RH, H2/Air 0.1/0.1 slpm'])
import matplotlib.pyplot as _plt
_plt.close(_fig)

out = tmp / 'out_bad'
res = quiet(run, str(bad), str(out), {'image_format': 'png'})
sc = json.loads((out / '_plot_data' / 'fcdur_ocv_vs_time.json').read_text(), parse_constant=_no_nan)
check('nothing judged: the readout says so rather than "none"',
      extract_values(sc).get('OCV death'), 'not judged')
check('... and the summary carries no OCV values',
      sorted(k for k in res['summary'][0] if k.startswith('OCV')), [])

out = tmp / 'out_nohold'
quiet(run, str(nohold), str(out), {'image_format': 'png'})
wb = openpyxl.load_workbook(out / 'fc_durability.xlsx', read_only=True)
check('no holds: no OCV sheets and no OCV plot',
      (wb.sheetnames, (out / 'fcdur_ocv_vs_time.png').exists()),
      (['Summary', 'Cycles', 'Holds', 'Characterizations', 'Manifest', 'Notes'], False))
wb.close()

# ── 7. Characterizations and verdict ─────────────────────────────────
print('7. characterizations and verdict')
import pandas as pd

AREA = 5.0
SETPOINTS = (0.90, 0.85, 0.80, 0.75, 0.70, 0.65, 0.60, 0.55, 0.50, 0.45, 0.40, 0.35)


def write_cv(path, j_x, R=None, rate=0.01, v_lo=0.12, v_hi=0.80, dt=0.5, cycles=2, c_dl=0.6,
             extra=None, sign=1.0):
    """An H2/N2 crossover CV in the station layout: j = j_x + V/R plus or minus
    double-layer charging (mA/cm², Ω·cm²), so the window mean is j_x plus
    the short's share and the intercept is j_x. sign=-1 logs oxidation as
    negative current."""
    n = int(round((v_hi - v_lo) / rate / dt))
    up = np.linspace(v_lo, v_hi, n + 1)
    V = np.concatenate([np.r_[up, up[-2::-1]][:-1] for _ in range(cycles)] + [up[:1]])
    sweep = np.r_[np.sign(np.diff(V)), -1.0]
    j = sign * (j_x + (V / R * 1000.0 if R else 0.0) + sweep * c_dl)
    t = dt * np.arange(1, len(V) + 1)
    head = FOUR + (',' + ','.join(extra) if extra else '')
    pad = ',' * len(extra) if extra else ''
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    with open(path, 'w') as fh:
        fh.write(head + '\n')
        fh.writelines(f'{a:.3f},1e-05,{v:.6f},{c * AREA / 1000:.8f}{pad}\n'
                      for a, v, c in zip(t, V, j))


def write_ecsa_cv(path, ecsa_cm2, rate=0.040, seed=1, cycles=3, t_scale=1.0):
    """Three cycles of the portal's synthetic H_UPD CV, logged at `rate` --
    40 mV/s, so a fixed 50 mV/s would read 25% high. t_scale=1000 stamps
    the log in milliseconds."""
    np.random.seed(seed)
    V1, I1 = generate_synthetic_cv(ecsa_cm2=ecsa_cm2, geo_area=AREA, scan_rate=rate)
    V, I = np.tile(V1, cycles), np.tile(I1, cycles)
    t = (0.5 + np.r_[0.0, np.cumsum(np.abs(np.diff(V)) / rate)]) * t_scale
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    with open(path, 'w') as fh:
        fh.write(FOUR + '\n')
        fh.writelines(f'{a:.6f},1e-05,{v:.6f},{c:.8g}\n' for a, v, c in zip(t, V, I))


def write_geis_sweep(path, r_ohm, seed=1, L=2e-8):
    """A GEIS sweep from 100 kHz to 0.1 Hz in the station layout: an R-RC
    cell (r_ohm in Ω·cm²) behind lead inductance, so Z'' crosses zero."""
    rng = np.random.default_rng(seed)
    freq = np.logspace(5, -1, 61)
    w = 2 * np.pi * freq
    Z = MODELS['R-RC']['func']([r_ohm, 0.15, 0.01, 0.88], w) / AREA + 1j * w * L
    Z = Z + 2e-5 * (rng.standard_normal(len(Z)) + 1j * rng.standard_normal(len(Z)))
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    with open(path, 'w') as fh:
        fh.write(GEIS + '\n')
        for k, (f, z) in enumerate(zip(freq, Z)):
            fh.write(f'{0.5 * (k + 1):.3f},0.038,-0.4990,{f:.5f},{abs(z):.6g},'
                     f'{np.degrees(np.angle(z)):.4f},{z.real:.8g},{z.imag:.8g},1.5\n')


def pol_model(j, loss=0.0):
    return 1.0 - 0.065 * np.log10(1 + j / 2e-4) - (0.12 + loss) * j + 0.03 * np.log(1 - j / 2.0)


def pol_j(v, loss=0.0):
    jj = np.linspace(1e-5, 1.99, 40000)
    return float(np.interp(v, pol_model(jj, loss)[::-1], jj[::-1]))


def write_stepped_pol(path, loss=0.0, dwell=25, dt=1.0, sweeps=3):
    """A potentiostatic pol curve (s, mV, mA, mW): a dwell at each setpoint,
    down and up three times, the current settling within each dwell."""
    order = list(SETPOINTS) + list(SETPOINTS[-2::-1])
    rows, t = [], 0.0
    for v in order * sweeps + [SETPOINTS[0]]:
        jv = pol_j(v, loss)
        for k in range(dwell):
            t += dt
            rows.append((t, round(v * 1000), jv * AREA * 1000 * (1 + 0.02 * np.exp(-k / 4.0))))
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    with open(path, 'w') as fh:
        fh.write(POL + '\n')
        fh.writelines(f'{a:.2f},{mv},{ma:.1f},{mv * ma / 1000:.1f}\n' for a, mv, ma in rows)


def xo(d, step, j_x, R=None, rh=100, flows='0.4WH2-0.4WN2', **kw):
    write_cv(Path(d) / f'{step}_S_H2X_80.3Cell-{rh}RH_{flows}.csv', j_x, R, **kw)


def ec(d, step, ecsa_cm2, **kw):
    write_ecsa_cv(Path(d) / f'{step}_S_ECSA_80.3Cell-100RH_0.05WH2-0.05WN2.csv', ecsa_cm2, **kw)


def eis(d, step, r_ohm, rh=100):
    write_geis_sweep(Path(d) / f'{step}_S_GEIS_80.3Cell-{rh}RH_0.05WH2-0.05WN2.csv', r_ohm)


def pol(d, step, loss=0.0, rh=95, T='80.3', flows='0.2WH2-0.2WA', **kw):
    write_stepped_pol(Path(d) / f'{step}_S_PolarizationCurve_{T}Cell-{rh}RH_{flows}.csv', loss,
                      **kw)


def ast_hold(d, step, k, prof=None):
    """The k-th 2 h hold, from hold hour 2(k - 1)."""
    write_ocv(Path(d) / HOLD.format(step), 7200, 30, 7200.0 * (k - 1), prof or flat(0.9))


def judge(d):
    tl = build_timeline(str(d))
    ca = analyze_characterizations(tl)
    ocv = analyze_ocv(tl)
    return tl, ca, ocv, decide(tl, ocv, ca)


# Case 1: crossover passes 15 mA/cm² between C5 and C6. BOL has a pre-break-in
# repeat and a 50% RH crossover above 15 that must not count (decision 13);
# C5's first repeat is above 15 but its last is not (decision 12); C1's
# crossover line falls with V, so it has no R_short.
xd_dir = tmp / 'xdeath'
xo(xd_dir, 'a1', 4.0)
xo(xd_dir, 'a2', 4.2)
xo(xd_dir, 'a3', 16.0, rh=50)
ec(xd_dir, 'a4', 300.0)
eis(xd_dir, 'a5', 0.045)
pol(xd_dir, 'a6', 0.0)
ast_hold(xd_dir, 'c1', 1)
xo(xd_dir, 'c2', 5.0, R=-5000.0, extra=['cd', '', ''])
ec(xd_dir, 'c3', 280.0)
eis(xd_dir, 'c4', 0.05)
pol(xd_dir, 'c5', 0.02)
for k, (hold_step, steps) in enumerate([('c6', [('c7', 7.0)]), ('c8', [('c9', 9.0)]),
                                        ('c10', [('c11', 11.0)]),
                                        ('c12', [('c13', 16.0), ('c14', 13.0)]),
                                        ('c15', [('c16', 11.0), ('c17', 17.0)])], start=2):
    ast_hold(xd_dir, hold_step, k)
    for step, j_x in steps:
        xo(xd_dir, step, j_x)
ec(xd_dir, 'c18', 250.0)
eis(xd_dir, 'c19', 0.06)
pol(xd_dir, 'c20', 0.05)
tl, ca, ocv, v = judge(xd_dir)
check('characterizations: BOL, then one after each hold', [iv.label for iv in ca.intervals],
      ['BOL', '2 h', '4 h', '6 h', '8 h', '10 h', '12 h'])
by = {x.file.step: x for x in ca.results}
close('crossover: window mean, no short', by['a2'].values['j_xover'], 4.2, 1e-3)
close('... intercept', by['a2'].values['j_intercept'], 4.2, 1e-3)
true('... no R_short from a flat line', np.isnan(by['a2'].values['R_short']), by['a2'].values)
true('... nor from a falling one', np.isnan(by['c2'].values['R_short']), by['c2'].values)
check('the death rule reads only the 100% RH crossover', ca.death_series, ['Crossover, 100% RH'])
check('repeats: the last represents each block',
      [x.file.step for x in ca.death_blocks()], ['a2', 'c2', 'c7', 'c9', 'c11', 'c14', 'c17'])
check('case 1: death at C6, bracket back to C5', (ca.xover_death.H, ca.xover_death.low), (12.0, 10.0))
check('... the verdict', (v.dead, v.hours, v.cause, v.bracket_text),
      (True, 12.0, 'crossover', '(10.00, 12.00] h'))
true('repeats more than 20% apart are flagged',
     any(n.startswith('10 h: Crossover, 100% RH repeats differ by 23%') for n in ca.notes), ca.notes)
check('... repeats 5% apart are not', [n for n in ca.notes if n.startswith('BOL: Crossover')], [])
check('DOE crossover breached at BOL, R_short never', (ca.doe_xover.label, ca.doe_r_short),
      ('BOL', None))

ecsa_csv = pd.read_csv(next(xd_dir.glob('a4_*.csv')))
with contextlib.redirect_stdout(io.StringIO()):
    oracle = compute_ecsa_hupd(ecsa_csv['Working Electrode Voltage'].to_numpy(),
                               ecsa_csv['Current'].to_numpy(), 0.040, AREA, v_low=0.08,
                               v_high=0.40, loading_mg_cm2=0.4, cycle=2)['roughness_factor']
close('ECSA: roughness factor as compute_ecsa_hupd at the logged 40 mV/s', by['a4'].values['RF'],
      oracle, 1e-3 * oracle)
close('... scan rate read from the log', by['a4'].values['scan_rate'], 40.0, 0.1)
close('... ECSA in m²/g from the 0.4 mg/cm² loading', by['a4'].values['ECSA'],
      by['a4'].values['RF'] / 4.0, 1e-9)
ca2 = analyze_characterizations(build_timeline(str(xd_dir)), CharSettings(loading=0.2))
close('... and from whatever loading is given', next(
    x for x in ca2.results if x.file.step == 'a4').values['ECSA'], by['a4'].values['RF'] / 2.0, 1e-9)
geis_csv = pd.read_csv(next(xd_dir.glob('a5_*.csv')))
o = np.argsort(geis_csv['Frequency'].to_numpy())[::-1]
oracle = extract_hfr(*(geis_csv[c].to_numpy()[o] for c in
                       ('Frequency', 'Real Impedance', 'Imaginary Impedance'))) * AREA * 1000
close('EIS: HFR as extract_hfr, in mΩ·cm²', by['a5'].values['HFR'], oracle, 1e-6)
close('... near the 45 mΩ·cm² built in', by['a5'].values['HFR'], 45.0, 2.5)
close('pol curve: V @ 1 A/cm² from the last up-sweep', by['a6'].values['V_1A'],
      float(np.interp(1.0, [pol_j(s) for s in SETPOINTS], SETPOINTS)), 2e-3)
close('... after decay', by['c20'].values['V_1A'],
      float(np.interp(1.0, [pol_j(s, 0.05) for s in SETPOINTS], SETPOINTS)), 2e-3)

# Case 8: C3 has no crossover file; death at C4 brackets back to C2. The
# third hold also loses more than 20% of its OCV without dying (DOE); a
# 10-min dip that deep in the first hold is too short to count.
gap = tmp / 'gap'
xo(gap, 'a1', 3.0)
ast_hold(gap, 'c1', 1, lambda H: np.where((H >= 3600) & (H < 4200), 0.74, 0.95))
xo(gap, 'c2', 5.0)
ast_hold(gap, 'c3', 2, flat(0.95))
xo(gap, 'c4', 8.0)
ast_hold(gap, 'c5', 3, lambda H: np.where(H < 5 * 3600, 0.95, 0.74))
ec(gap, 'c6', 280.0)
ast_hold(gap, 'c7', 4, flat(0.95))
xo(gap, 'c8', 16.0)
tl, ca, ocv, v = judge(gap)
true('case 8: a characterization without a crossover is skipped, and noted',
     '6 h: no 100% RH crossover, so this characterization is skipped for the crossover rule'
     in ca.notes, ca.notes)
check('... the bracket reaches back to the last crossover', (ca.xover_death.H, ca.xover_death.low),
      (8.0, 4.0))
close('DOE OCV loss: the first 20% loss that persists', ocv.loss_breach_h, 5.0, 1 / 60)
check('... with no OCV death', ocv.death, None)

# Case 6: BOL failure.
bolf = tmp / 'bolfail'
xo(bolf, 'a1', 20.0)
ast_hold(bolf, 'c1', 1)
xo(bolf, 'c2', 21.0)
tl, ca, ocv, v = judge(bolf)
check('case 6: a BOL crossover over 15 is death at 0 h', (v.dead, v.hours, v.cause, v.at_bol,
                                                          v.bracket_text, v.text),
      (True, 0.0, 'crossover', True, 'at BOL', 'dead at BOL (0 h), crossover'))

# Case 7: a low R_short puts the window mean over 15 and leaves the intercept under.
short = tmp / 'short'
xo(short, 'a1', 3.0)
ast_hold(short, 'c1', 1)
xo(short, 'c2', 12.0, R=100.0)
tl, ca, ocv, v = judge(short)
c2 = next(x for x in ca.results if x.file.step == 'c2')
close('case 7: the window mean carries the short', c2.values['j_xover'], 16.25, 0.05)
close('... the intercept removes it', c2.values['j_intercept'], 12.0, 0.01)
close('... R_short', c2.values['R_short'], 100.0, 0.5)
check('case 7: the verdict follows the window mean (decision 3)', (v.dead, v.hours, v.cause),
      (True, 2.0, 'crossover'))
true('... and flags the short-corrected value',
     any('short-corrected intercept 12.00 is not' in n for n in ca.notes), ca.notes)
check('DOE R_short breached there', ca.doe_r_short.label, '2 h')

# Both rules: the OCV trips inside the bracket of a later crossover trip;
# then an OCV trip in an earlier interval than the crossover's.
both = tmp / 'both'
xo(both, 'a1', 3.0)
ast_hold(both, 'c1', 1)
xo(both, 'c2', 5.0)
ast_hold(both, 'c3', 2, lambda H: np.where(H < 3 * 3600, 0.9, 0.6))
xo(both, 'c4', 18.0)
tl, ca, ocv, v = judge(both)
check('both rules in one interval: the OCV time, cause both', (v.cause, round(v.hours, 3)),
      ('both', 3.0))
check('... the OCV time is exact, so no bracket is shown', (v.bracket_text, v.text),
      ('', 'dead at 3.00 h, both'))
first = tmp / 'ocvfirst'
xo(first, 'a1', 3.0)
ast_hold(first, 'c1', 1, lambda H: np.where(H < 3600, 0.9, 0.6))
xo(first, 'c2', 4.0)
ast_hold(first, 'c3', 2, flat(0.6))
xo(first, 'c4', 18.0)
tl, ca, ocv, v = judge(first)
check('OCV first, crossover in a later interval: cause OCV', (v.cause, round(v.hours, 3), v.bracket),
      ('OCV', 1.0, None))
close('DOE OCV loss breach at the OCV drop', ocv.loss_breach_h, 1.0, 1 / 60)

check('evolution keeps the first, the last and evenly between', _pick_curves(10), [0, 3, 6, 9])
check('... and all of a few', _pick_curves(3), [0, 1, 2])

# run() end to end on case 1.
out = tmp / 'out_xdeath'
res = quiet(run, str(xd_dir), str(out), {'image_format': 'png'})
row = res['summary'][0]
check('summary verdict', (row['Verdict'], row['Hours to death'], row['Death cause'],
                          row['Death bracket']), ('dead', 12.0, 'crossover', '(10.00, 12.00] h'))
close('summary j_xover initial is BOL', row['j_xover initial'], 4.2, 1e-3)
close('summary j_xover final is the last block', row['j_xover final'], 17.0, 1e-3)
true('no R_short final from a flat line', 'R_short final' not in row, row)
true('summary row is JSON-safe', json.dumps(row, allow_nan=False) is not None)
plots = ['fcdur_timeline', 'fcdur_ocv_vs_time', 'fcdur_xover_vs_time', 'fcdur_hfr_vs_time',
         'fcdur_ecsa_vs_time', 'fcdur_perf_vs_time', 'fcdur_polcurve_evolution',
         'fcdur_xover_evolution', 'fcdur_nyquist_evolution', 'fcdur_ecsa_evolution']
check('every plot written', sorted(p.stem for p in out.glob('*.png')), sorted(plots))
side = {n: json.loads((out / '_plot_data' / f'{n}.json').read_text(), parse_constant=_no_nan)
        for n in plots}
check('sidecar plot types', [side[n]['plot_type'] for n in plots], plots)
vals = extract_values(side['fcdur_timeline'])
check('timeline readout: the verdict', (vals.get('Verdict'), (vals.get('Hours to death') or {}).get('value'),
                                        vals.get('Death cause'), vals.get('Death bracket')),
      ('dead', 12.0, 'crossover', '(10.00, 12.00] h'))
vals = extract_values(side['fcdur_xover_vs_time'])
close('crossover readout: death', (vals.get('Crossover death') or {}).get('value'), 12.0, 1e-9)
close('... DOE breach at BOL', (vals.get('DOE crossover breach') or {}).get('value'), 0.0, 1e-9)
EVO_CAP = {'fcdur_polcurve_evolution': 40, 'fcdur_xover_evolution': 50,
           'fcdur_nyquist_evolution': 61, 'fcdur_ecsa_evolution': 60}
evo = {n: max(len(ln['x']) for ax in side[n]['data']['axes'] for ln in ax['lines']) for n in EVO_CAP}
true('evolution curves thinned', all(evo[n] <= EVO_CAP[n] for n in EVO_CAP), evo)
size = sum((out / '_plot_data' / f'{n}.json').stat().st_size for n in plots)
true('all sidecars together under 250 KB', size < 250_000, size)
wb = openpyxl.load_workbook(out / 'fc_durability.xlsx', read_only=True)
check('workbook: the verdict heads the summary',
      list(wb['Summary'].iter_rows(max_row=2, values_only=True)),
      [('Item', 'Value'), ('Verdict', 'dead at 12.00 h, crossover, between 10.00 and 12.00 h')])
check('Cycles: one row per characterization', wb['Cycles'].max_row - 1, 7)
check('Characterizations: every file', wb['Characterizations'].max_row - 1, 20)
wb.close()

# Case 3, and the crossover parameters reaching the rule.
res = quiet(run, str(xd_dir), str(tmp / 'out_alive'), {'image_format': 'none', 'xover_threshold': '20'})
row = res['summary'][0]
check('case 3: alive at the hours tested, no Hours to death',
      (row['Verdict'], 'Hours to death' in row, row['Hours tested'], res['verdict']),
      ('alive', False, 12.0, 'alive at 12.00 h'))
res = quiet(run, str(short), str(tmp / 'out_win'),
            {'image_format': 'none', 'xover_window_lo': '0.15', 'xover_window_hi': '0.20'})
check('the crossover window reaches the rule', res['verdict'], 'alive at 2.00 h')
res = quiet(run, str(short), str(tmp / 'out_win2'),
            {'image_format': 'none', 'xover_window_lo': '0.5', 'xover_window_hi': '0.4'})
check('an empty window falls back to 0.35-0.50 V, with a note',
      (res['verdict'], any('is empty' in n for n in res['notes'])),
      ('dead at 2.00 h, crossover, between 0.00 and 2.00 h', True))
res = quiet(run, str(tmp / 'scribner'), str(tmp / 'out_placeholder'), {'image_format': 'png'})
check('placeholder characterization files: alive, crossover rule not judged',
      (res['verdict'], any('crossover rule could not be applied' in n for n in res['notes'])),
      ('alive at 60.00 h', True))
sc = json.loads((tmp / 'out_placeholder' / '_plot_data' / 'fcdur_timeline.json').read_text(),
                parse_constant=_no_nan)
vals = extract_values(sc)
check('timeline readout when alive: Verdict alive as in the summary row, no stray keys',
      (vals.get('Verdict'), sorted(vals)),
      ('alive', ['Hold files', 'Hold segments', 'Hours tested', 'Verdict']))
true('... and the rule that could not be applied is shown',
     any('crossover rule not applied (no data)' in t['text']
         for ax in sc['data']['axes'] for t in ax['texts']), vals)

# ── 8. M3 review guards ──────────────────────────────────────────────
print('8. review guards')

# Two 100% RH crossover series (two flow settings) are judged apart, not as
# repeats of one: the 0.2/0.2 slpm series trips at 4 h while the 0.4/0.4 one
# is under 15 there, and its bracket reaches back to its own last
# measurement (BOL), not the other series' (2 h).
two = tmp / 'twoflows'
xo(two, 'a1', 4.0)
xo(two, 'a2', 4.0, flows='0.2WH2-0.2WN2')
ast_hold(two, 'c1', 1)
xo(two, 'c2', 5.0)
ast_hold(two, 'c3', 2)
xo(two, 'c4', 15.5, flows='0.2WH2-0.2WN2')
xo(two, 'c5', 6.0)
ast_hold(two, 'c6', 3)
xo(two, 'c7', 16.0)
tl, ca, ocv, v = judge(two)
check('two 100% RH crossover series are each judged; the earliest trip counts',
      (len(ca.death_series), ca.xover_death.result.file.step, ca.xover_death.H, ca.xover_death.low),
      (2, 'c4', 4.0, 0.0))
check('... the headline follows the series measured most often',
      (ca.primary_death_series, [x.file.step for x in ca.death_blocks()]),
      ('Crossover, 100% RH, 80.3 °C, H2/N2 0.4/0.4 slpm', ['a1', 'c2', 'c5', 'c7']))
true('... and the notes say so', any(n.startswith('2 100% RH crossover series') for n in ca.notes),
     ca.notes)
f0 = next(f for f in tl.files if f.role == 'crossover')
twin = dc_replace(f0, anode_humid=f0.anode_humid + 'x')     # differs where no name shows it
check('series no name tells apart are counted', sorted(_series_labels([f0, twin]).values()),
      ['Crossover, 100% RH, 80.3 °C, H2/N2 0.4/0.4 slpm (1)',
       'Crossover, 100% RH, 80.3 °C, H2/N2 0.4/0.4 slpm (2)'])

# Died on crossover before any OCV trip: the headline OCV decay rate is the
# latest window ending by the death (5 mV/h), not the faster fall after it.
xr = tmp / 'xrate'
xo(xr, 'a1', 3.0)
ast_hold(xr, 'c1', 1, lambda H: 0.95 - 5e-3 * H / 3600)
xo(xr, 'c2', 5.0)
ast_hold(xr, 'c3', 2, lambda H: 0.95 - 5e-3 * H / 3600)
xo(xr, 'c4', 16.0)
ast_hold(xr, 'c5', 3, lambda H: 0.93 - 0.065 * (H - 14400) / 3600)
xo(xr, 'c6', 17.0)
res = quiet(run, str(xr), str(tmp / 'out_xrate'), {'image_format': 'none'})
close('crossover death first: the decay rate stops there too', res['summary'][0]['OCV decay rate'],
      5000.0, 5.0)
true('... and says so', 'OCV decay rate is from the latest window ending by the crossover death '
     'at 4.00 h (2.00–4.00 h)' in res['notes'], res['notes'])

# A window the CVs do not reach falls back to 0.35-0.50 V, rather than
# leaving the rule unjudged and calling a dead sample alive.
lowcv = tmp / 'lowcv'
xo(lowcv, 'a1', 3.0, v_hi=0.6)
ast_hold(lowcv, 'c1', 1)
xo(lowcv, 'c2', 16.0, v_hi=0.6)
res = quiet(run, str(lowcv), str(tmp / 'out_lowcv'),
            {'image_format': 'none', 'xover_window_lo': '0.7', 'xover_window_hi': '0.8'})
check('a window outside the CVs falls back to 0.35-0.50 V, with a note',
      (res['verdict'],
       'crossover window 0.7–0.8 V lies outside the CVs; used 0.35–0.5 V' in res['notes']),
      ('dead at 2.00 h, crossover, between 0.00 and 2.00 h', True))

# Pol curves at two RH and two temperatures on the same flows: separate,
# readable series with their own readout lines. The 60 °C curve after the
# hold no longer reaches 1 A/cm²: it still represents its block, with j @
# 0.65 V and peak power, and the readout says 'not reached' rather than
# repeating an earlier V @ 1 A/cm².
pols = tmp / 'pols'
pol(pols, 'a1')
pol(pols, 'a2', rh=50)
pol(pols, 'a3', T='60')
ast_hold(pols, 'c1', 1)
pol(pols, 'c2', 0.02)
pol(pols, 'c3', 0.02, rh=50)
pol(pols, 'c4', 0.45, T='60')
out = tmp / 'out_pols'
res = quiet(run, str(pols), str(out), {'image_format': 'png'})
tl = build_timeline(str(pols))
check('pol-curve series apart by RH and by temperature',
      sorted(set(_series_labels([f for f in tl.files if f.role == 'polcurve']).values())),
      ['Pol curve, 50% RH, H2/Air 0.2/0.2 slpm',
       'Pol curve, 95% RH, H2/Air 0.2/0.2 slpm, 60 °C',
       'Pol curve, 95% RH, H2/Air 0.2/0.2 slpm, 80.3 °C'])
vals = extract_values(json.loads((out / '_plot_data' / 'fcdur_perf_vs_time.json').read_text(),
                                 parse_constant=_no_nan))
check('... each with its own readout lines',
      sum(k.startswith(('V @ 1 A/cm² ', 'j @ 0.65 V ')) for k in vals), 12)
close('... 50% RH final from its own curve, j @ 0.65 V',
      (vals.get('j @ 0.65 V final, 50% RH, H2/Air 0.2/0.2 slpm') or {}).get('value'),
      pol_j(0.65, 0.02), 5e-3)
check('a last curve short of 1 A/cm² reads not reached',
      vals.get('V @ 1 A/cm² final, 95% RH, H2/Air 0.2/0.2 slpm, 60 °C'), 'not reached')
close('... while its j @ 0.65 V is read',
      (vals.get('j @ 0.65 V final, 95% RH, H2/Air 0.2/0.2 slpm, 60 °C') or {}).get('value'),
      pol_j(0.65, 0.45), 5e-3)
check('... and no j @ 0.6 V is reported any more', [k for k in vals if k.startswith('j @ 0.6 V')], [])
wb = openpyxl.load_workbook(out / 'fc_durability.xlsx', read_only=True)
rows = list(wb['Characterizations'].iter_rows(values_only=True))
wb.close()
c4 = dict(zip(rows[0], next(r for r in rows if r[3] == 'c4')))
check('... it represents its block, unflagged, and says how far it reached',
      (c4['Represents block'], c4['V @ 1 A/cm² (V)'], c4['Error'], c4['Note']),
      ('yes', None, None, f'reaches {pol_j(0.35, 0.45):.2f} A/cm², short of 1 A/cm²'))

# Oxidation logged as negative current: flipped, with a note, same numbers;
# and the crossover scales with the geometric area.
neg = tmp / 'neg.csv'
write_cv(neg, 12.0, R=100.0, sign=-1.0)
nv, _, note = analyze_crossover_file(str(neg), CharSettings())
check('negative oxidation current: flipped and noted, R_short kept',
      (round(nv['j_xover'], 3), round(nv['j_intercept'], 3), round(nv['R_short'], 1), note),
      (16.246, 12.0, 100.0, 'oxidation current logged as negative; sign flipped'))
close('crossover per geometric area', analyze_crossover_file(str(neg), CharSettings(geo_area=2.5))[0]
      ['j_xover'], 2 * nv['j_xover'], 1e-9)

# ECSA: a log stamped in milliseconds gives an implausible scan rate, so 50
# mV/s is assumed, and said; a single-cycle CV is read on its last cycle, and
# said.
ecd = tmp / 'ecsa'
ec(ecd, 'a1', 300.0)
ast_hold(ecd, 'c1', 1)
ec(ecd, 'c2', 300.0, t_scale=1000.0)
ast_hold(ecd, 'c3', 2)
ec(ecd, 'c4', 300.0, cycles=1)
tl, ca, ocv, v = judge(ecd)
by = {x.file.step: x for x in ca.results}
ecsa_csv = pd.read_csv(next(ecd.glob('c2_*.csv')))
with contextlib.redirect_stdout(io.StringIO()):
    oracle = compute_ecsa_hupd(ecsa_csv['Working Electrode Voltage'].to_numpy(),
                               ecsa_csv['Current'].to_numpy(), 0.050, AREA, v_low=0.08,
                               v_high=0.40, loading_mg_cm2=0.4, cycle=2)['roughness_factor']
close('ECSA: an implausible logged scan rate falls back to 50 mV/s', by['c2'].values['RF'], oracle,
      1e-3 * oracle)
check('... said, with the scan rate left blank',
      ('2 h: c2 (ECSA) scan rate from the log (0.04 mV/s) is implausible; assumed 50 mV/s'
       in ca.notes, np.isnan(by['c2'].values['scan_rate'])), (True, True))
check('a single-cycle CV: its last cycle, said',
      (by['c4'].values['cycle'],
       '4 h: c4 (ECSA) only 1 cycle(s); the last was analyzed, not cycle 2' in ca.notes),
      (1.0, True))
close('... the same roughness factor as cycle 2 of three', by['c4'].values['RF'],
      by['a1'].values['RF'], 1e-6 * by['a1'].values['RF'])

# A long run -- BOL and eight characterizations, crossover at two RH, EIS
# at three, pol curves at two flows -- keeps its sidecars small: at most two
# series per evolution plot, four curves per series, each thinned.
long_run = tmp / 'long'


def characterize(d, k, n):
    xo(d, f'{k}1', 3.0 + 0.5 * n)
    xo(d, f'{k}2', 2.0 + 0.4 * n, rh=50)
    ec(d, f'{k}3', 300.0 - 10 * n)
    for m, rh in enumerate((100, 50, 30)):
        eis(d, f'{k}{4 + m}', 0.045 + 0.002 * n, rh=rh)
    pol(d, f'{k}7', 0.01 * n, sweeps=1)
    pol(d, f'{k}8', 0.01 * n, flows='0.1WH2-0.1WA', sweeps=1)


characterize(long_run, 'a', 0)
for n in range(1, 9):
    ast_hold(long_run, f'c{n}0', n)
    characterize(long_run, f'c{n}', n)
out = tmp / 'out_long'
res = quiet(run, str(long_run), str(out), {'image_format': 'png'})
side = {n: json.loads((out / '_plot_data' / f'{n}.json').read_text(), parse_constant=_no_nan)
        for n in EVO_CAP}
check('long run: evolution panels and curves capped',
      {n: (len(side[n]['data']['axes']), max(len(ax['lines']) for ax in side[n]['data']['axes']))
       for n in EVO_CAP},
      {'fcdur_polcurve_evolution': (2, 4), 'fcdur_xover_evolution': (2, 4),
       'fcdur_nyquist_evolution': (2, 4), 'fcdur_ecsa_evolution': (1, 4)})
evo = {n: max(len(ln['x']) for ax in side[n]['data']['axes'] for ln in ax['lines']) for n in EVO_CAP}
true('... and thinned', all(evo[n] <= EVO_CAP[n] for n in EVO_CAP), evo)
# The OCV plot's share follows its bin (the 3,000-bin cap, checked with the
# point budget); everything else stays under the plan's guide.
size = sum(p.stat().st_size for p in (out / '_plot_data').glob('*.json')
           if p.stem != 'fcdur_ocv_vs_time')
true('... all sidecars but the OCV plot together under 150 KB', size < 150_000, size)

# ── 9. The portal: registry, form, job runner, store ─────────────────
print('9. portal and store')
from concurrent.futures import Future
os.environ['JOBS_DIR'] = str(tmp / 'jobs')
os.environ.update(JSONBIN_API_KEY='test', JSONBIN_COLLECTION_ID='COLL', JSONBIN_INDEX_BIN_ID='INDEX')
with contextlib.redirect_stdout(io.StringIO()):
    import scripts as portal_scripts
    import app.main as portal
    from scripts.helpers import jsonbin
    from scripts.helpers.record import _round_sig
    from scripts.fc_durability_analysis import PARAM_RANGES, run as fcdur_run

check('registered as FC Durability', portal_scripts.SCRIPT_REGISTRY.get('FC Durability') is fcdur_run,
      True)
check('... with its short label', portal_scripts.SCRIPT_SHORT.get('FC Durability'), 'FCDurability')
form = {f['key']: f for f in portal_scripts.SCRIPT_PARAMS.get('FC Durability', [])}
check('form: no "No Images", which would store nothing (V6)',
      [o['value'] for o in form.get('image_format', {}).get('options', [])],
      ['png', 'svg', 'pdf', 'tiff'])
check('form: the stand kept or declared, no Auto-detect (the CSVs are the same on both stands)',
      ([o['value'] for o in form.get('stand', {}).get('options', [])], form.get('stand', {}).get('default')),
      (['keep', 'Scribner 1', 'Scribner 2', 'FCTS 1', 'FCTS 2', 'FCTS 3', 'FCTS 4'], 'keep'))
numbers = {k: f for k, f in form.items() if f['type'] == 'number'}
check('form: a field for every parameter run() reads (bar the station step tail)',
      sorted(numbers), sorted(k for k in PARAM_RANGES if k != 'step_tail_s'))
check('form: defaults as run() has them',
      {k: f['default'] for k, f in numbers.items()},
      {k: PARAM_RANGES[k][0] for k in numbers if k in PARAM_RANGES})
check('form: bounds inside run()\'s ranges, defaults on each field\'s step grid',
      [k for k, f in numbers.items() if k not in PARAM_RANGES
       or not PARAM_RANGES[k][1] <= f['min'] <= f['default'] <= f['max'] <= PARAM_RANGES[k][2]
       or abs((f['default'] - f['min']) / f['step'] - round((f['default'] - f['min']) / f['step']))
       > 1e-6], [])


class Store:
    """JSONBin in memory: bins and the index persist across pushes."""

    def __init__(self):
        self.index, self.bins = {'schema': 2, 'runs': []}, {}

    def __call__(self, url, method='GET', body=None, extra_headers=None):
        copy = (lambda x: json.loads(json.dumps(x)))
        if method == 'POST':
            bid = f'BIN{len(self.bins) + 1}'
            self.bins[bid] = copy(body)
            return {'record': body, 'metadata': {'id': bid}}
        if method == 'GET':
            bid = next((b for b in self.bins if f'/{b}/' in url), None)
            return {'record': copy(self.bins[bid] if bid else self.index)}
        if method == 'PUT':
            bid = next((b for b in self.bins if url.endswith(f'/{b}')), None)
            if bid:
                self.bins[bid] = copy(body)
            else:
                self.index = copy(body)
            return {'record': body}
        raise AssertionError(method)


store = Store()
jsonbin._request = store


def submit(job_id, src, params, keep=lambda rel: True):
    """One job through the portal's own runner and completion callback."""
    d = portal.JOBS_DIR / job_id
    (d / 'output').mkdir(parents=True)
    for p in sorted(Path(src).rglob('*')):
        if p.is_file() and keep(p.relative_to(src)):
            (d / 'input' / p.relative_to(src)).parent.mkdir(parents=True, exist_ok=True)
            shutil.copy(p, d / 'input' / p.relative_to(src))
    names = [str(p.relative_to(d / 'input')) for p in sorted((d / 'input').rglob('*')) if p.is_file()]
    portal.jobs[job_id] = {'job_id': job_id, 'status': 'running', 'script': 'FC Durability',
                           'input_files': names, 'params': dict(params)}
    fut = Future()
    with contextlib.redirect_stdout(io.StringIO()):
        try:
            fut.set_result(portal._run_job(job_id, 'FC Durability', str(d / 'input'),
                                           str(d / 'output'), dict(params)))
        except Exception as e:                       # noqa: BLE001
            fut.set_exception(e)
        portal._on_job_done(job_id, fut)
        # The push runs after, on the portal's one push thread, in order: a
        # no-op queued behind it returns once it has reported.
        portal.push_executor.submit(lambda: None).result()
    return portal.jobs[job_id]


DECISION_7 = ('Hours tested', 'Hours to death', 'OCV initial', 'OCV decay rate', 'OCV loss',
              'j_xover initial', 'j_xover final', 'R_short final', 'RH cycles tested',
              'RH cycles to death')


def indexed(row):
    """What the index should hold for a summary row: the plan's decision-7
    keys it has (and the RH cycle counts, for a test that cycles RH), at six
    significant figures."""
    return {k: _round_sig(row[k]) for k in DECISION_7 if k in row}


SAMPLE = '260601_DUR-1'
params = {'sample_name': SAMPLE, 'image_format': 'png', 'stand': 'FCTS 2'}
# The first push is case 1 cut at 8 h, still alive; the second is the whole
# test, dead at 12 h: the same sample pushed again with more data.
early = lambda rel: (parse_step_name(rel.name) or {}).get('letter') == 'a' or \
    int((parse_step_name(rel.name) or {}).get('number') or 99) <= 11
j1 = submit('job-1', xd_dir, params, keep=early)
push = j1.get('metrics_push') or {}
check('first push: stored as a new sample', (j1['status'], push.get('pushed'), push.get('merged')),
      ('complete', True, False))
outs = j1['output_files']
check('every output named {sample}-FCDurability-..., none twice',
      (all(o.startswith(f'{SAMPLE}-FCDurability-') for o in outs), len(set(outs)) == len(outs),
       len(outs)), (True, True, 11))
entry = store.index['runs'][0]
row1 = j1['script_result']['summary'][0]
check('index: one step-less fc_durability unit',
      [(u['Analysis'], u['step']) for u in entry['Data']], [('fc_durability', '')])
check('... with the summary row\'s values', entry['Data'][0].get('key_values'), indexed(row1))
check('... and no Hours to death while alive',
      ('Hours to death' in entry['Data'][0].get('key_values', {}), row1['Verdict']),
      (False, 'alive'))
check('... stand as declared, run date from the name', (entry.get('stand'), entry.get('run_date')),
      ('FCTS 2', '2026-06-01'))
rec = store.bins[entry['bin_id']]
check('record: all ten plots in one bucket', {b: len(p) for b, p in rec['metrics'].items()},
      {'fc_durability': 10})
check('... one summary row', len(rec.get('summary') or []), 1)

j2 = submit('job-2', xd_dir, params)
push = j2.get('metrics_push') or {}
check('re-push: merged into the same sample', (push.get('pushed'), push.get('merged'),
                                               push.get('bin_id'), len(store.index['runs'])),
      (True, True, entry['bin_id'], 1))
entry = store.index['runs'][0]
row2 = j2['script_result']['summary'][0]
check('... still one unit, now with the dead run\'s values',
      [(u['Analysis'], u['step'], u.get('key_values')) for u in entry['Data']],
      [('fc_durability', '', indexed(row2))])
check('... Hours to death 12 h', entry['Data'][0].get('key_values', {}).get('Hours to death'), 12.0)
rec = store.bins[entry['bin_id']]
check('... one summary row, the new one, after two pushes',
      [(r.get('Label'), r.get('Verdict')) for r in rec.get('summary') or []], [('fcdur', 'dead')])
check('... ten plots, replaced rather than added', {b: len(p) for b, p in rec['metrics'].items()},
      {'fc_durability': 10})
check('... both jobs recorded, n_jobs 2', ([j['job_id'] for j in rec.get('jobs') or []],
                                          entry.get('n_jobs')), (['job-1', 'job-2'], 2))
check('... every sidecar kept, within the body limit',
      (push.get('sidecars_kept'), push.get('sidecars_dropped'),
       push.get('body_bytes', 1e9) < jsonbin.MAX_BODY_BYTES), (10, 0, True))

# Re-analysed with a 900 mA/cm² threshold the sample is alive again, pushed
# on the form's default, Keep stored stand: nothing of the dead verdict may
# survive, and the stand stays as stored.
j3 = submit('job-3', xd_dir, dict(params, xover_threshold='900', stand='keep'))
entry = store.index['runs'][0]
rec = store.bins[entry['bin_id']]
row3 = j3['script_result']['summary'][0]
check('dead back to alive: Hours to death gone from the index and the record',
      ([(u['Analysis'], u.get('key_values')) for u in entry['Data']],
       [r.get('Verdict') for r in rec.get('summary') or []]),
      ([('fc_durability', indexed(row3))], ['alive']))
check('... "Keep stored stand" keeps it, in both', (entry.get('stand'), rec.get('stand')),
      ('FCTS 2', 'FCTS 2'))
# A stand chosen on a re-push replaces the stored one (V5), across families.
submit('job-3b', xd_dir, dict(params, xover_threshold='900', stand='Scribner 2'))
entry = store.index['runs'][0]
rec = store.bins[entry['bin_id']]
check('a stand chosen on a re-push replaces it, in both (V5)',
      (entry.get('stand'), rec.get('stand'), len(store.index['runs'])), ('Scribner 2', 'Scribner 2', 1))
# A new sample pushed on "Keep stored stand" has no stand to keep, and its
# files cannot name one: none is recorded (the page asks for one first).
# The page asks the store whether a sample is in it just before uploading, so
# the runs list it reads must not be the 30 s cache, which here was filled
# before the push.
import asyncio
import inspect


def listed(force):
    with contextlib.redirect_stdout(io.StringIO()):
        out = portal.view_runs(sample=None, script=None, analysis=None, stand=None,
                               since=None, until=None, limit=None, force=force)
        if inspect.iscoroutine(out):
            out = asyncio.run(out)
    return {r['sample_name'] for r in out.get('runs', [])}


# Reading JSONBin blocks; the page asks before every upload on Keep, so the
# endpoint runs in FastAPI's thread pool rather than on the event loop.
check('the runs list is read off the event loop', inspect.iscoroutinefunction(portal.view_runs),
      False)


listed(False)
submit('job-new', xd_dir, dict(params, sample_name='260603_DUR-NEW', stand='keep'))
fresh = next((e for e in store.index['runs'] if e['sample_name'] == '260603_DUR-NEW'), {})
check('a new sample on "Keep stored stand": no stand recorded, none guessed from the CSVs',
      (fresh.get('stand'), store.bins.get(fresh.get('bin_id'), {}).get('stand'),
       bool(fresh.get('Data'))), (None, None, True))
import time
from scripts.helpers import viewstore as _vs
# The push drops the views' copy of the index, so even a read that is not
# forced finds a sample pushed a moment ago, however recently the copy was
# filled before it.
check('a push drops the views\' 30 s copy of the index: a sample pushed a moment ago is listed',
      (_vs._index_cache['data'], '260603_DUR-NEW' in listed(False)), (None, True))
# A copy filled from before the push (another process's, say) still hides
# it; read fresh on request, it is there.
_vs._index_cache.update(data={**store.index, 'runs': [e for e in store.index['runs']
                                                      if e['sample_name'] != '260603_DUR-NEW']},
                        at=time.time())
check('the runs list read fresh on request shows a sample pushed a moment ago; a stale copy does not',
      ('260603_DUR-NEW' in listed(False), '260603_DUR-NEW' in listed(True)), (False, True))

# The page reads the store again once an analysis it ran is stored. The push
# runs after the job is complete, so a poll must see that it is still to come:
# push_pending is set with "complete" and cleared beside the push's result,
# whatever the result.
import scripts.helpers.jsonbin as _jb
_real_push = _jb.push_job_metrics
_seen = {}


def _spy_push(**kw):
    _seen.update(portal.jobs[kw['job_id']])
    return _real_push(**kw)


def _broken_push(**kw):
    raise RuntimeError('store down')


try:
    _jb.push_job_metrics = _spy_push
    _done = submit('job-pending', xd_dir, dict(params, sample_name='260603_DUR-PENDING', stand='keep'))
    _jb.push_job_metrics = _broken_push
    _broke = submit('job-broken', xd_dir, dict(params, sample_name='260603_DUR-BROKEN', stand='keep'))
finally:
    _jb.push_job_metrics = _real_push
check('while its push runs a job is complete with push_pending set; after, cleared beside the result',
      (_seen.get('status'), _seen.get('push_pending'), _done.get('push_pending'),
       (_done.get('metrics_push') or {}).get('pushed')), ('complete', True, False, True))
check('... cleared too when the push fails outright',
      (_broke.get('status'), _broke.get('push_pending'), (_broke.get('metrics_push') or {}).get('pushed')),
      ('complete', False, False))
_empty = tmp / 'empty_upload'
_empty.mkdir(exist_ok=True)
check('a job that fails has no push to wait for',
      (lambda j: (j.get('status'), 'push_pending' in j))(
          submit('job-failed', _empty, dict(params, sample_name='260603_DUR-FAILED'))),
      ('failed', False))

# A read of the index out when a push drops the views' copy (forget_index)
# may be from before the push: it is answered, not kept.
_real_fetch = _jb.fetch_index
_old_index = {**store.index, 'runs': store.index['runs'][:1]}


def _read_across_a_push():
    _vs.forget_index()                     # the push lands while this read is out
    return _old_index


try:
    _vs.forget_index()
    _jb.fetch_index = _read_across_a_push
    _first = _vs.fetch_index()
    _jb.fetch_index = _real_fetch
    _second = _vs.fetch_index()
finally:
    _jb.fetch_index = _real_fetch
check('a read of the index out across a push is answered but not kept: the next read is fresh',
      (_first is _old_index, _vs._index_cache['data'] is _old_index,
       len(_second['runs']) == len(store.index['runs'])), (True, False, True))

# An admin change applied rewrites bins without moving their entries'
# timestamps: the views' caches are dropped; a dry run leaves them.
_vs.fetch_index()
_vs._detail_cache['BIN@t'] = {'stale': True}
portal._admin_call(lambda apply=False: {'applied': apply}, apply=False)
_kept = ('BIN@t' in _vs._detail_cache, _vs._index_cache['data'] is not None)
portal._admin_call(lambda apply=False: {'applied': apply}, apply=True)
check('an admin change applied drops the views\' caches; a dry run does not',
      (_kept, 'BIN@t' in _vs._detail_cache, _vs._index_cache['data']), ((True, True), False, None))

# Comparisons ask for a durability plot by its type: the unit is step-less and
# holds ten plots, and (analysis, step) alone answers with the first stored.
from scripts.helpers import viewstore
TYPES = ('fcdur_ocv_vs_time', 'fcdur_xover_vs_time', 'fcdur_hfr_vs_time', 'fcdur_ecsa_vs_time',
         'fcdur_perf_vs_time')
key = store.index['runs'][0]['bin_id']
with contextlib.redirect_stdout(io.StringIO()):
    got = viewstore.plot_series_batch(
        [{'key': key, 'analysis': 'fc_durability', 'step': '', 'plot_type': t} for t in TYPES]
        + [{'key': key, 'analysis': 'fc_durability', 'step': '', 'plot_type': 'fcdur_nothing'}])
check('series by plot type: each durability trajectory, the unknown type missing',
      ([(s['plot_type'], s['plot'].endswith(t)) for s, t in zip(got['series'], TYPES)],
       [m.get('plot_type') for m in got['missing']]),
      ([(t, True) for t in TYPES], ['fcdur_nothing']))
xo = next((s for s in got['series'] if s['plot_type'] == 'fcdur_xover_vs_time'), {'axes': [{}]})
mean = next((ln for ln in xo['axes'][0].get('lines', []) if 'window mean' in ln['label']), {})
check('... the crossover series against hold hours, BOL to the last characterization',
      (mean.get('label'), round((mean.get('x') or [0])[0], 6), round((mean.get('x') or [0])[-1], 6)),
      ('Crossover, 100% RH, window mean', 0.0, 12.0))

# A sample named ..._fcdur_... that also holds a stepped analysis: every plot
# name contains the row label 'fcdur', and the durability row must still be
# replaced on each push rather than kept beside the new one.
FCD = '260602_fcdur_MEA7'
pc = tmp / 'pc_out' / '_plot_data'
pc.mkdir(parents=True)
(pc / f'{FCD}-FCAnalysis-polcurve_c03_80C.json').write_text(json.dumps(
    {'plot_type': 'polcurve', 'data': {'axes': [{'texts': [{'text': 'OCV = 0.95 V'}],
                                                 'axhlines': [], 'axvlines': [], 'lines': []}]}}))
quiet(jsonbin.push_job_metrics, job_id='job-pc', sample_name=FCD, script='Fuel Cell Full Analysis',
      output_dir=tmp / 'pc_out', input_files=['c03.csv'], script_short='FCAnalysis',
      summary=[{'Analysis': 'polcurve', 'Label': 'polcurve_c03', 'OCV': 0.95}])
submit('job-4', xd_dir, dict(params, sample_name=FCD))
j5 = submit('job-5', xd_dir, dict(params, sample_name=FCD, xover_threshold='900'))
entry = next(e for e in store.index['runs'] if e['sample_name'] == FCD)
rec = store.bins[entry['bin_id']]
check('"fcdur" in the sample name: one durability row after re-analysis, the other kept',
      ([(r['Analysis'], r.get('Verdict')) for r in rec.get('summary') or []],
       sorted((u['Analysis'], u['step']) for u in entry['Data']),
       next((u for u in entry['Data'] if u['Analysis'] == 'fc_durability'), {}).get('key_values')),
      ([('polcurve', None), ('fc_durability', 'alive')], [('fc_durability', ''), ('polcurve', 'c03')],
       indexed(j5['script_result']['summary'][0])))

# A stored line past 400 points reaches the browser thinned to bucket extremes,
# unless asked for whole: Comparisons smooths the OCV trace from medians of
# the stored points, which extremes are not. Whole is still bounded.
LONG = '260604_DUR-LONG'
lp = tmp / 'long_out' / '_plot_data'
lp.mkdir(parents=True)
hold = [[round(k / 60, 4) for k in range(1000)], [0.6 if k % 10 == 0 else 0.9 for k in range(1000)]]
huge = [[round(k / 600, 4) for k in range(25000)], [0.9] * 25000]
(lp / f'{LONG}-FCDurability-fcdur_ocv_vs_time.json').write_text(json.dumps(
    {'plot_type': 'fcdur_ocv_vs_time',
     'data': {'axes': [{'ylabel': 'Hold OCV (V)', 'texts': [], 'axhlines': [], 'axvlines': [],
                        'lines': [{'x': hold[0], 'y': hold[1], 'label': 'hold'},
                                  {'x': huge[0], 'y': huge[1], 'label': 'huge'}]}]}}))
(lp / f'{LONG}-FCDurability-fcdur_hfr_vs_time.json').write_text(json.dumps(
    {'plot_type': 'fcdur_hfr_vs_time',
     'data': {'axes': [{'ylabel': 'HFR (mΩ·cm²)', 'texts': [], 'axhlines': [], 'axvlines': [],
                        'lines': [{'x': hold[0], 'y': hold[1], 'label': 'HFR'}]}]}}))
quiet(jsonbin.push_job_metrics, job_id='job-long', sample_name=LONG, script='FC Durability',
      output_dir=tmp / 'long_out', input_files=['c2.csv'], script_short='FCDurability',
      summary=[{'Analysis': 'fc_durability', 'Label': 'fcdur', 'Hours tested': 41.7}])
key = next(e['bin_id'] for e in store.index['runs'] if e['sample_name'] == LONG)
ask = {'key': key, 'analysis': 'fc_durability', 'step': '', 'plot_type': 'fcdur_ocv_vs_time'}
with contextlib.redirect_stdout(io.StringIO()):
    got = viewstore.plot_series_batch([ask, dict(ask, full=True), dict(ask, full='yes')])
shape = [[(len(ln['x']), ln.get('decimated', False)) for ln in s['axes'][0]['lines']]
         for s in got['series']]
check('series: past 400 points thinned and marked; asked for whole, the stored line, '
      'up to 20,000 points',
      (shape[0][0][1] and shape[0][0][0] <= 402, shape[1], shape[2] == shape[0]),
      (True, [(1000, False), (shape[1][1][0], True)], True))
true('... the stored points themselves, dips and all',
     got['series'][1]['axes'][0]['lines'][0]['y'] == hold[1]
     and shape[1][1][0] <= viewstore.MAX_FULL_SERIES_POINTS + 2, shape)
with contextlib.redirect_stdout(io.StringIO()):
    hfr = viewstore.plot_series_batch([dict(ask, plot_type='fcdur_hfr_vs_time', full=True)])
check('... whole only for the OCV plot: another type asked for whole is still thinned',
      [(len(ln['x']), ln.get('decimated', False)) for ln in hfr['series'][0]['axes'][0]['lines']],
      [shape[0][0]])

# The OCV band through the portal, from stored records: Comparisons' Export
# with the band on for one sample and off for another, and Recall's redraw.
import scripts.compare_polcurves as compare_mod
from scripts.helpers import plot_compare
from scripts.helpers.record import decode_sidecars
# Its legend and Same Conditions grouping read the step as the index does: a
# letter after the number kept (b17b read b1 before), and none taken from a
# gas flow (f3b's grouped as 'h', a zz1 purge's as 'n', from H2 and N2).
check('the comparison reads the step as the index does', [compare_mod._parse_full_conditions(n).get('step')
      for n in ('polcurve_b17b_IV_80C-95RH_OCV-0o3V_0o2H2-0o2Air_0kPa.png',
                'polcurve_f3b_S_PolarizationCurve_80.3Cell-95RH_0.2WH2-0.2WA.png',
                'polcurve_b14_IV_80C_100RH_OCV_0o3V_0o2H2_0o2Air_0kPa.png',
                'ocv_zz1_Purge_25C-Dry_0o5N2-0o5N2_0kPa.png')], ['b17b', 'f3b', 'b14', None])
check('... and groups by that step\'s letter', [compare_mod._parse_condition_signature(n) for n in (
          'polcurve_f3b_IV_80C_100RH_OCV_0o3V_0o2H2_0o2Air_0kPa.png',
          'polcurve_b14_IV_80C_100RH_OCV_0o3V_0o2H2_0o2Air_0kPa.png',
          'ocv_zz1_Purge_25C-Dry_0o5N2-0o5N2_0kPa.png')],
      ['f_80C_100RH_0kPa', 'b_80C_100RH_0kPa', '25C_0kPa'])
# Its label takes the step a script declared, as the index files the plot:
# the electrolyzer cell id, which a name like polcurve_a1-rev2 does not give
# the stricter reader; else the name's, lower-cased.
ELX = {'plot_type': 'polcurve', 'metadata': {'conditions': {'step': 'a1-rev2', 'T_C': 80.0}}}
HOLD = {'plot_type': 'polcurve', 'metadata': {'conditions': {'j_hold_Acm2': 1.0}}}
FC_B = 'S2-FullAnalysis-polcurve_F3B_S_PolarizationCurve_80.3Cell-95RH_0.2WH2-0.2WA.png'
check('the comparison\'s label takes a declared step, else the name\'s',
      [compare_mod._build_clean_label('S1', 'S1-ElxPolCurve-polcurve_a1-rev2.png', 'plot_type', ELX),
       compare_mod._build_clean_label('S1', 'S1-ElxPolCurve-polcurve_a1-rev2.png', 'plot_type'),
       compare_mod._build_clean_label('S2', FC_B, 'plot_type_conditions', HOLD),
       compare_mod._build_clean_label('S2', FC_B, 'plot_type_conditions')],
      ['S1_a1-rev2', 'S1', 'S2_f3b', 'S2_f3b'])
lab_seen = []
real_lab_render, real_lab_xlsx = compare_mod.render_overlay_comparison, compare_mod.export_comparison_excel
compare_mod.render_overlay_comparison = lambda items, **kw: lab_seen.append([it['label'] for it in items])
compare_mod.export_comparison_excel = lambda *a, **k: None
try:
    with contextlib.redirect_stdout(io.StringIO()):
        compare_mod._compare_generic([
            {'filename': 'S1-ElxPolCurve-polcurve_a1-rev2.png', 'sample_name': 'S1', 'sidecar': ELX},
            {'filename': FC_B, 'sample_name': 'S2', 'sidecar': HOLD}],
            str(tmp / 'lab_out'), {'image_format': 'png', 'grouping_mode': 'plot_type_conditions'})
finally:
    compare_mod.render_overlay_comparison = real_lab_render
    compare_mod.export_comparison_excel = real_lab_xlsx
check('... in the comparison drawn', lab_seen, [['S1_a1-rev2', 'S2_f3b']])
k1, k2 = (next(e['bin_id'] for e in store.index['runs'] if e['sample_name'] == s)
          for s in (SAMPLE, '260603_DUR-NEW'))
with contextlib.redirect_stdout(io.StringIO()):
    ocv = {s['key']: s['plot'] for s in viewstore.plot_series_batch(
        [{'key': k, 'analysis': 'fc_durability', 'step': '', 'plot_type': 'fcdur_ocv_vs_time'}
         for k in (k1, k2)])['series']}
stored_ocv = decode_sidecars(store.bins[k1])[ocv[k1]]['data']['axes'][0]
holds = sum(1 for ln in stored_ocv['lines'] if ln.get('label') == BAND_LOW)
true('stored record: the OCV plot keeps its band edges', holds >= 2, holds)
with contextlib.redirect_stdout(io.StringIO()):
    sources = viewstore.materialize_for_compare(
        [{'key': k1, 'plot': ocv[k1], 'label': 'with band', 'band': True},
         {'key': k2, 'plot': ocv[k2], 'label': 'without', 'band': False},
         {'key': k2, 'plot': ocv[k2], 'label': 'unsaid', 'band': 'yes'}], tmp / 'cmp_in')
check('staging for Export: a literal band flag is passed on, anything else left out',
      [s.get('band', 'absent') for s in sources], [True, False, 'absent'])
seen = []
real_render = compare_mod.render_overlay_comparison


def spy_render(items, **kw):
    fig = real_render(items, **kw)
    seen.append(([(it['label'], it.get('band')) for it in items],
                 sum(isinstance(c, PolyCollection) for c in fig.axes[0].collections)))
    return fig


compare_mod.render_overlay_comparison = spy_render
try:
    with contextlib.redirect_stdout(io.StringIO()):
        got = compare_mod.run(str(tmp / 'cmp_in'), str(tmp / 'cmp_out'),
                              {'sources': json.dumps(sources), 'image_format': 'png'})
finally:
    compare_mod.render_overlay_comparison = real_render
check('Export: each sample\'s choice reaches the figure, unsaid drawn as stored; only the '
      'two banded samples are shaded',
      seen, [([('with band', True), ('without', False), ('unsaid', True)], 2 * holds)])
wl = trace_labels(next((tmp / 'cmp_out').rglob('*.xlsx')))
check('... and the workbook: the band\'s edges for those two only',
      {k: len(edge(v)) for k, v in wl.items()}, {'with band': 2 * holds, 'without': 0, 'unsaid': 2 * holds})
true('... the figure written', any(f.endswith('.png') for f in got['files_produced']), got)
# Recall's run detail redraws one stored plot through the same renderer.
redraws = []
real_overlay = plot_compare.render_overlay_comparison


def spy_overlay(items, **kw):
    fig = real_overlay(items, **kw)
    redraws.append(sum(isinstance(c, PolyCollection) for c in fig.axes[0].collections))
    return fig


plot_compare.render_overlay_comparison = spy_overlay
try:
    with contextlib.redirect_stdout(io.StringIO()):
        resp = portal.view_render({'key': k1, 'plot': ocv[k1]})
        if inspect.iscoroutine(resp):
            resp = asyncio.run(resp)
finally:
    plot_compare.render_overlay_comparison = real_overlay
check('Recall\'s redraw of the stored OCV plot shades every hold\'s band',
      (redraws, Path(resp.path).exists()), ([holds], True))

# A band edge level across a hold's two bins is stored as two points (a level
# two-point line would read back as a reference line). The redraw pairs it
# all the same, and the series sent to a chart carry it as a line, so the
# Comparisons chart pairs it too; other points are not sent.
LEVEL = '260605_DUR-LEVEL'
lv = tmp / 'level_out' / '_plot_data'
lv.mkdir(parents=True)
level_sc = {'plot_type': 'fcdur_ocv_vs_time',
            'data': {'axes': [{'ylabel': 'Hold OCV (V)', 'texts': [], 'axhlines': [], 'axvlines': [],
                               'lines': [{'x': [0.5, 1.5], 'y': [0.95, 0.96], 'label': BAND_HIGH},
                                         {'x': [0.5, 1.5], 'y': [0.91, 0.92], 'label': '30% RH'}],
                               'scatters': [{'x': [0.5, 1.5], 'y': [0.9, 0.9], 'sizes': [10.0],
                                             'label': BAND_LOW},
                                            {'x': [0.2, 1.8], 'y': [0.93, 0.92], 'sizes': [34.0],
                                             'label': 'settled start and end (30-min medians)'}]}]}}
(lv / f'{LEVEL}-FCDurability-fcdur_ocv_vs_time.json').write_text(json.dumps(level_sc))
quiet(jsonbin.push_job_metrics, job_id='job-level', sample_name=LEVEL, script='FC Durability',
      output_dir=tmp / 'level_out', input_files=['c2.csv'], script_short='FCDurability',
      summary=[{'Analysis': 'fc_durability', 'Label': 'fcdur', 'Hours tested': 2.0}])
key = next(e['bin_id'] for e in store.index['runs'] if e['sample_name'] == LEVEL)
with contextlib.redirect_stdout(io.StringIO()):
    got = viewstore.plot_series_batch([{'key': key, 'analysis': 'fc_durability', 'step': '',
                                        'plot_type': 'fcdur_ocv_vs_time', 'full': True}])
    fig4 = plot_compare.render_overlay_comparison([{'label': 'level', 'sidecar': level_sc}])
check('a band edge stored as two level points: sent to the chart as a line, other points not',
      sorted((ln['label'], ln['x'], ln['y']) for ln in got['series'][0]['axes'][0]['lines']),
      sorted([(BAND_HIGH, [0.5, 1.5], [0.95, 0.96]), ('30% RH', [0.5, 1.5], [0.91, 0.92]),
              (BAND_LOW, [0.5, 1.5], [0.9, 0.9])]))
check('... and shaded by the redraw', sum(isinstance(c, PolyCollection) for c in fig4.axes[0].collections), 1)
_plt.close(fig4)

# Recall's workbook export lists each line beside its own x: an FC Durability
# plot draws each hold on its own hours, where most plots share one x column.
from scripts.helpers import xlsx_export


def expected_columns(ax):
    cols, cur = [], None
    for ln in ax['lines']:
        if not ln.get('x'):
            continue
        if ln['x'] != cur:
            cols.append((ax.get('xlabel') or 'x', ln['x']))
            cur = ln['x']
        cols.append((ln.get('label') or 'series', ln['y']))
    return cols


rec1 = store.bins[k1]
want_cols = expected_columns(stored_ocv)
block = next(c for lab, c in xlsx_export._data_blocks(rec1, decode_sidecars(rec1), 'fc_durability')
             if lab.endswith('fcdur_ocv_vs_time'))
check('analysis workbook: each hold\'s columns beside its own hours, one x column a hold',
      (block == want_cols, sum(1 for h, _ in want_cols if h == (stored_ocv.get('xlabel') or 'x'))),
      (True, holds))
wbx = openpyxl.load_workbook(io.BytesIO(xlsx_export.build_plot_workbook(rec1, ocv[k1])))
wsx = next(w for w in wbx.worksheets if w.cell(2, 2).value == BAND_LOW)
n_rows = max(len(v) for _, v in want_cols)
sheet_cols = [(wsx.cell(2, c).value, [wsx.cell(r, c).value for r in range(3, 3 + n_rows)])
              for c in range(1, len(want_cols) + 1)]
check('... and the plot workbook alike',
      sheet_cols, [(h, v + [None] * (n_rows - len(v))) for h, v in want_cols])
check('... while lines on one x keep a single x column',
      xlsx_export._columns('j', [('V', [1, 2], [3, 4]), ('P', [1, 2], [5, 6])]),
      [('j', [1, 2]), ('V', [3, 4]), ('P', [5, 6])])
# A later panel on the same x label and length joins the analysis block, each
# line still beside its own x (the evolution plots: a panel per RH, one
# length, different x).
two_panel = {'p': {'plot_type': 'x', 'data': {'axes': [
    {'xlabel': 'h', 'lines': [{'x': [0, 1], 'y': [5, 6], 'label': 'a'},
                              {'x': [0, 1], 'y': [7, 8], 'label': 'b'}]},
    {'xlabel': 'h', 'title': 'rate', 'lines': [{'x': [0, 1], 'y': [9, 9.5], 'label': 'r'}]}]}}}
detail2 = {'metrics': {'b': {'p': {'conditions': {}}}}}
heads = []
for bx, rx in (([0, 1], [0, 1]), ([0, 1], [5, 6]), ([2, 3], [5, 6])):
    two_panel['p']['data']['axes'][0]['lines'][1]['x'] = bx
    two_panel['p']['data']['axes'][1]['lines'][0]['x'] = rx
    heads.append(xlsx_export._data_blocks(detail2, two_panel, 'b')[0][1])
check('... a later panel joins the block, each line beside its own x',
      heads,
      [[('h', [0, 1]), ('a', [5, 6]), ('b', [7, 8]), ('rate · r', [9, 9.5])],
       [('h', [0, 1]), ('a', [5, 6]), ('b', [7, 8]), ('h', [5, 6]), ('rate · r', [9, 9.5])],
       [('h', [0, 1]), ('a', [5, 6]), ('h', [2, 3]), ('b', [7, 8]), ('h', [5, 6]),
        ('rate · r', [9, 9.5])]])

# ── 10. Scribner FuelCell (.fcd) steps ───────────────────────────────
# The Scribner's own software writes one .fcd file per step, all in one
# folder: no sample in the name, 'o' for the decimal point, CVs named by
# scan rate, a header before the data, and Time (Sec) running on from the
# start of the station session.
print('10. Scribner .fcd steps')
import math
from types import SimpleNamespace
with contextlib.redirect_stdout(io.StringIO()):
    import scripts.fc_durability_analysis as FD
    from scripts.fc_durability_analysis import (
        parse_fcd_step_name, fcd_layout, fcd_columns, read_time_bounds, analyze_ecsa_file,
        analyze_eis_file, analyze_polcurve_file, step_formats, step_tail_text)

# The 36 columns FuelCell5 logs, as on the KM8-S6 run (October 2026).
FCD_COLS = ['Time (Sec)', 'I (A)', 'I (mA/cm²)', 'Power (Watts)', 'Power (mW/cm²)', 'E_Stack (V)',
            'E_Avg (V)', 'E_Comp_Stack (V)', 'E_Comp_Avg (V)', 'E_iR_Stack (V)', 'E_iR_Avg (V)',
            'E_iR_Stack (mOhm)', 'E_iR_Avg (mOhm*cm²)', 'Cell (C)', 'Temp_Anode (C)',
            'Flow_Anode  (l/min)', 'Wet_Anode', 'Temp_Cathode (C)', 'Flow_Cathode  (l/min)',
            'Wet_Cathode', 'HFR (mOhm)', 'HFR (mOhm*cm²)', 'Z_Freq (Hz)', 'Z_Real (Ohm)',
            'Z_Imag (Ohm)', 'Z_Mag (Ohm)', 'Z_Angle (Deg)', 'Ctrl_Status', 'Ctrl_Mode', 'Ctrl_Value',
            'Ctrl_Repeat', 'Ctrl_Temp_Cell (C)', 'Ctrl_Temp_Anode (C)', 'Ctrl_Temp_Cathode (C)',
            'Ctrl_Flow_Anode  (l/min)', 'Ctrl_Flow_Cathode  (l/min)']
TERM = 'Open Circuit; 30 Seconds; Fixed 1 Sec/Pt; Term:E<0.7 NextStep'


def write_fcd(path, t, data, desc='Open Circuit; 30 Seconds; Fixed 1 Sec/Pt; ', area=5,
              columns=None):
    """A FuelCell5 file as the station writes one: 'Key: value' header lines,
    the column line, 'End Comments', then tab-separated rows, CRLF throughout.
    `t` is Time (Sec) on the session clock; `data` maps columns to values,
    and the rest log 0; `columns`, the columns logged, all of FCD_COLS by
    default. Returns how many lines come before the data."""
    t = np.asarray(t, float)
    logged = columns or FCD_COLS
    cols = {c: np.zeros(len(t)) for c in logged}
    cols['Time (Sec)'] = t
    for c, v in data.items():
        cols[c] = np.broadcast_to(np.asarray(v, float), t.shape)
    head = ['FUELCELL DATA', 'Format: FuelCell5', 'Time: "10/1/2026 9:00:00 AM"',
            'DateTime: 46296.375', 'comment 1: ""', 'comment 2: ""',
            f'Description: "{desc}; Single File: {Path(path).name}(Append)"',
            f'SurfaceArea: {area:g}', 'StackNumber: 1', f'DataColumnCount: {len(logged)}']
    head += [f'DataColumn{k}: {k}' for k in range(1, len(logged) + 1)]
    head += ['\t'.join(logged), 'End Comments']
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    with open(path, 'w', encoding='utf-8', newline='') as fh:
        fh.write('\r\n'.join(head) + '\r\n')
        fh.writelines('\t'.join(f'{cols[c][k]:.10g}' for c in logged) + '\r\n'
                      for k in range(len(t)))
    return len(head)


def csv_to_fcd(src, dst, t0=6601.1, zero_hz_rows=0):
    """The same measurement as a .fcd file: E_Stack and I (A) for the
    working-electrode voltage and current (or mV and mA), Z_* for a GEIS
    sweep, which the station opens with rows at 0 Hz; the first sample at t0
    on the session clock."""
    df = pd.read_csv(src)
    t = df[df.columns[0]].to_numpy(float)
    if 'mV' in df.columns:
        data = {'E_Stack (V)': df['mV'] / 1000.0, 'I (A)': df['mA'] / 1000.0}
    elif 'Frequency' in df.columns:
        data = {'E_Stack (V)': df['DC Electrode Voltage'], 'I (A)': df['DC Current'],
                'Z_Freq (Hz)': df['Frequency'], 'Z_Real (Ohm)': df['Real Impedance'],
                'Z_Imag (Ohm)': df['Imaginary Impedance']}
    else:
        data = {'E_Stack (V)': df['Working Electrode Voltage'], 'I (A)': df['Current']}
    data = {k: np.asarray(v, float) for k, v in data.items()}
    if zero_hz_rows:           # with stale readings in Z_*, which must not count
        data = {k: np.r_[np.full(zero_hz_rows, 0.5 if k.startswith('Z_') else v[0]), v]
                for k, v in data.items()}
        data['Z_Freq (Hz)'][:zero_hz_rows] = 0.0
        t = np.r_[t[0] - np.arange(zero_hz_rows, 0, -1), t]
    return write_fcd(dst, t0 + t - t[0], data)


# Names.
r = parse_fcd_step_name('c4_OCV-RHC-1000cyc_90C-100RH(90C)-0RH_0o1H2-0o1Air_0kPa.fcd')
check('fcd name: step, kind, no sample', (r['step'], r['letter'], r['number'], r['kind'],
                                          r['kind_token'], r['sample'], r['source']),
      ('c4', 'c', 4, 'ocv', 'OCV-RHC-1000cyc', '', 'fcd'))
check('... RH cycling: the wet RH, the cycle said in words', (r['T_C'], r['RH_pct'], r['rh_text']),
      (90.0, 100.0, '100↔0% RH cycling'))
check('... flows with o for the point, the back pressure',
      (r['anode_gas'], r['anode_slpm'], r['cathode_gas'], r['cathode_slpm'], r['P_kPa'], r['suffix']),
      ('H2', 0.1, 'Air', 0.1, 0.0, ''))
KM8 = {'b20': 'b20_CV-2mVs_0o12to0o8V_80C-100RH_0o4H2-0o4N2_0kPa.fcd',
       'b24': 'b24_CV-50mVs_0o08to1V_80C-100RH_0o1H2-0o1N2_0kPa.fcd',
       'a4': 'a4_CV-500mVs_0o05to1o2V_80C-100RH_0o2H2-0o2N2_0kPa.fcd',
       'b14': 'b14_IV_80C-95RH_OCV-0o3V_0o2H2-0o2Air_0kPa.fcd',
       'b26': 'b26_GEIS_80C-100RH_0o1H2-0o1N2_0kPa.fcd',
       'b1': 'b1_activation_0o6V_80C-100RH_0o2H2-0o2Air_0kPa.fcd',
       'b3': 'b3_Purge_80C-100RH_0o1H2-0o4N2_0kPa.fcd',
       'c3': 'c3_OCV_N2-to-Air_90C-100RH_0o1H2-0o1Air_0kPa.fcd',
       'b8': 'b8_Cell-ramp-up_90C-100RH_0o05H2-0o05N2_0kPa.fcd'}
check('fcd kinds: CVs by scan rate, IV, GEIS, activation, Purge an OCV, ramps other',
      {k: parse_fcd_step_name(n)['kind'] for k, n in KM8.items()},
      {'b20': 'crossover', 'b24': 'ecsa', 'a4': 'eclean', 'b14': 'polcurve', 'b26': 'eis',
       'b1': 'breakin', 'b3': 'ocv', 'c3': 'ocv', 'b8': 'other'})
check('... a CV range, an IV span or a gas change kept as the suffix',
      [parse_fcd_step_name(KM8[k])['suffix'] for k in ('b20', 'b14', 'c3')],
      ['0o12to0o8V', 'OCV-0o3V', 'N2-to-Air'])
r = parse_fcd_step_name('c1_T-ramp-up_93C-(90C)RH_0o1H2-0o1N2_0kPa.fcd')
close('a dew point gives the RH it makes at the cell temperature', r['RH_pct'], 89.2, 0.05)
check('... and is said as one', r['rh_text'], 'dew point 90 °C')
r = parse_fcd_step_name('zz2_shut-down_25C-Dry_0o4N2-0o4N2_0kPa.fcd')
check('a two-letter block; Dry', (r['step'], r['letter'], r['number'], r['RH_pct'], r['anode_humid'],
                                  r['anode_gas'], r['cathode_gas']),
      ('zz2', 'zz', 2, 0.0, 'dry', 'N2', 'N2'))
for n, want in (('f3b_OCV_80C-100RH_0o1H2-0o1N2_0kPa.fcd', ('f3b', 'f', 3, 'b')),
                ('ZZ1B_Purge_25C-Dry_0o5N2-0o5N2_0kPa.fcd', ('zz1b', 'zz', 1, 'b'))):
    r = parse_fcd_step_name(n)
    check(f'fcd name, a letter after the step number: {want[0]}',
          (r['step'], r['letter'], r['number'], r['trail']), want)
check('... two letters after it is no step', parse_fcd_step_name('f3bc_OCV_80C-100RH_0o1H2-0o1N2_0kPa.fcd'),
      None)
for k, step in enumerate(('f2', 'F3B', 'f3', 'f3a', 'f4')):
    write_fcd(tmp / 'trail_fcd' / '261019_SYN-F' / f'{step}_OCV_80C-100RH_0o1H2-0o1N2_0kPa.fcd',
              1000.0 * (k + 1) + np.arange(10, 601.0, 10), {'E_Stack (V)': 0.1},
              desc='Open Circuit; 10 Minutes; Fixed 10 Sec/Pt; ')
tl = build_timeline(str(tmp / 'trail_fcd'))
check('a .fcd run with f3, f3a and F3B: all read, in step order whatever the case, nothing skipped',
      ([f.step for f in tl.files], tl.skipped), (['f2', 'f3', 'f3a', 'f3b', 'f4'], []))
r = parse_fcd_step_name('a0_T-cell-ramp-up_83C_0kPa.fcd')
check('a temperature alone: no RH, no gases', (r['T_C'], math.isnan(r['RH_pct']), r['anode_gas'],
                                              r['P_kPa']), (83.0, True, None, 0.0))
check('not .fcd step files', [parse_fcd_step_name(n) for n in (
    'notes_80C-100RH.fcd', 'b3_OCV.fcd', 'abc1_OCV_80C-100RH_0o1H2-0o1Air.fcd',
    'b3_OCV_95Cell-30RH_0.1WH2-0.2WA.fcd')], [None] * 4)
check('kinds by their first word; OCV among the later words only when the first names nothing',
      [parse_fcd_step_name(f'b5_{k}_80C-100RH_0o1H2-0o1N2_0kPa.fcd')['kind'] for k in (
          'GEIS-150mAcm2', 'IV-H2Air', 'CV-2mVs-x2', 'RHC-OCV-1000cyc', 'Purge-Cell-ramp-down',
          'Break-in', 'PolarizationCurve', 'T-ramp-up', 'IV-OCV-0o3V', 'GEIS-OCV',
          'activation-OCV-0o6V', 'Break-in-OCV', 'N2-to-Air-OCV')],
      ['eis', 'polcurve', 'crossover', 'ocv', 'ocv', 'breakin', 'polcurve', 'other', 'polcurve',
       'eis', 'breakin', 'breakin', 'ocv'])
r = parse_fcd_step_name('c4_OCV_0o1H2-0o2Air_90C-100RH_150kPa.fcd')
check('flows and pressure read wherever they stand after the kind',
      (r['kind'], r['anode_gas'], r['cathode_gas'], r['cathode_slpm'], r['P_kPa'], r['suffix']),
      ('ocv', 'H2', 'Air', 0.2, 150.0, ''))

# The file: header, columns, the step's own clock, units, the 0 Hz rows.
fu = tmp / 'fcd_unit'
n_head = write_fcd(fu / 'a3_OCV_80C-100RH_0o2H2-0o2N2_0kPa.fcd', 6601.1 + np.arange(12.0),
                   {'E_Stack (V)': 0.1 + 0.001 * np.arange(12), 'I (A)': 0.25})
head, cols, skip = fcd_layout(str(fu / 'a3_OCV_80C-100RH_0o2H2-0o2N2_0kPa.fcd'))
check('fcd header: its fields, the column line, data after End Comments',
      (head['Format'], head['SurfaceArea'], head['Description'][:21], len(cols), cols[5], skip),
      ('FuelCell5', '5', 'Open Circuit; 30 Seco', 36, 'E_Stack (V)', n_head))
check('a step 6601 s into the session is on its own clock, from one interval in',
      read_time_bounds(str(fu / 'a3_OCV_80C-100RH_0o2H2-0o2N2_0kPa.fcd')), (12, 1.0, 12.0))
c = fcd_columns(str(fu / 'a3_OCV_80C-100RH_0o2H2-0o2N2_0kPa.fcd'),
                ('Working Electrode Voltage', 'Current', 'mV', 'mA'))
check('E_Stack and I (A) stand in for the voltage and current, mV and mA scaled',
      [round(float(c[k].iloc[3]), 9) for k in ('Working Electrode Voltage', 'Current', 'mV', 'mA')],
      [0.103, 0.25, 103.0, 250.0])
write_fcd(fu / 'b26_GEIS_80C-100RH_0o1H2-0o1N2_0kPa.fcd', 50.0 + np.arange(5.0),
          {'Z_Freq (Hz)': [0, 0, 1000, 100, 10], 'Z_Real (Ohm)': [9, 9, 0.01, 0.012, 0.02],
           'Z_Imag (Ohm)': [9, 9, 0.001, -0.002, -0.004]})
g = fcd_columns(str(fu / 'b26_GEIS_80C-100RH_0o1H2-0o1N2_0kPa.fcd'),
                ('Frequency', 'Real Impedance', 'Imaginary Impedance'))
check('GEIS: the rows at 0 Hz before the sweep carry no impedance',
      [np.isnan(g[k].to_numpy()[:2]).all() and np.isfinite(g[k].to_numpy()[2:]).all()
       for k in ('Frequency', 'Real Impedance', 'Imaginary Impedance')], [True] * 3)
touch(fu / 'bad' / 'b9_OCV_80C-100RH_0o1H2-0o1N2_0kPa.fcd', b'FUELCELL DATA\r\n1\t2\r\n')
check('a .fcd without End Comments is skipped as unreadable',
      build_timeline(str(fu / 'bad')).skipped,
      [('b9_OCV_80C-100RH_0o1H2-0o1N2_0kPa.fcd',
        "unreadable: ValueError: no 'End Comments' line, so not a FuelCell data file")])

# The analyzers read a .fcd file as they read the station CSV of the same
# measurement.
eq = tmp / 'fcd_equiv'
write_cv(eq / 'x.csv', 3.0, R=2000.0)
write_ecsa_cv(eq / 'e.csv', 300.0)
write_geis_sweep(eq / 'g.csv', 0.05)
write_stepped_pol(eq / 'p.csv', 0.02)
for src, zero, fn in (('x', 0, analyze_crossover_file), ('e', 0, analyze_ecsa_file),
                      ('g', 2, analyze_eis_file), ('p', 0, analyze_polcurve_file)):
    csv_to_fcd(eq / f'{src}.csv', eq / f'{src}.fcd', zero_hz_rows=zero)
    a = quiet(fn, str(eq / f'{src}.csv'), CharSettings())
    b = quiet(fn, str(eq / f'{src}.fcd'), CharSettings())
    same_values = sorted(a[0]) == sorted(b[0]) and all(
        (math.isnan(a[0][k]) and math.isnan(b[0][k])) or abs(a[0][k] - b[0][k]) <= 1e-9 * max(1.0, abs(a[0][k]))
        for k in a[0])
    same_curve = all(np.allclose(x, y, rtol=1e-9, atol=1e-12) for x, y in zip(a[1], b[1])) and \
        [len(x) for x in a[1]] == [len(y) for y in b[1]]
    check(f'{fn.__name__}: a .fcd file gives the station CSV\'s values and curve',
          (same_values, same_curve, a[2] == b[2]), (True, True, True))


def fcd_run(root):
    """A Scribner run as FuelCell writes it, the KM8-S6 shape cut short: BOL
    characterization, two 3 h OCV holds under RH cycling, each followed by
    a crossover (and the first by a pol curve), the second hold's run ending
    in a wet half under the station's own E < 0.7 V limit, then the
    shut-down; one folder, with the run's notes and a stray .fcd beside the
    steps."""
    d = Path(root) / '261001_SYN-FCD-1'
    src = Path(root).parent / (Path(root).name + '_sources')
    hold = 'OCV-RHC-1000cyc_90C-100RH(90C)-0RH_0o1H2-0o1Air_0kPa.fcd'
    t0 = iter(np.arange(1, 20) * 40000.3)            # each step later in the session

    def step(name, t_el, data, **kw):
        write_fcd(d / name, next(t0) + t_el, data, **kw)

    def measured(name, writer, *a, zero_hz_rows=0):
        writer(src / f'{name}.csv', *a)
        csv_to_fcd(src / f'{name}.csv', d / name, t0=next(t0), zero_hz_rows=zero_hz_rows)

    step('a1_T-hum-ramp-up_83C-100RH_0o1H2-0o1N2_0kPa.fcd', np.arange(10, 601.0, 10),
         {'E_Stack (V)': 0.1}, desc='Open Circuit; 10 Minutes; Fixed 10 Sec/Pt; ')
    measured('a6_CV-2mVs_0o12to0o8V_80C-100RH_0o4H2-0o4N2_0kPa.fcd', write_cv, 3.0)
    measured('a10_CV-50mVs_0o08to1V_80C-100RH_0o1H2-0o1N2_0kPa.fcd', write_ecsa_cv, 300.0)
    measured('a12_GEIS_80C-100RH_0o1H2-0o1N2_0kPa.fcd', write_geis_sweep, 0.05, zero_hz_rows=2)
    measured('b14_IV_80C-95RH_OCV-0o3V_0o2H2-0o2Air_0kPa.fcd', write_stepped_pol, 0.0)
    step('c3_OCV_N2-to-Air_90C-100RH_0o1H2-0o1Air_0kPa.fcd', np.arange(1, 121.0),
         {'E_Stack (V)': 0.93}, desc='Open Circuit; 2 Minutes; Fixed 1 Sec/Pt; ')
    t_el = np.arange(15, 10801.0, 15)          # 30 s wet, 45 s dry: 144 RH cycles
    wet = ((t_el % 75) < 30).astype(float)
    step('c4_' + hold, t_el, {'E_Stack (V)': 0.95 - 0.004 * t_el / 3600, 'Wet_Anode': wet,
                               'Wet_Cathode': wet}, desc=TERM)
    measured('c6_CV-2mVs_0o12to0o8V_80C-100RH_0o4H2-0o4N2_0kPa.fcd', write_cv, 3.5)
    measured('c8_IV_80C-95RH_OCV-0o3V_0o2H2-0o2Air_0kPa.fcd', write_stepped_pol, 0.02)
    t_el = np.arange(15, 9001.0, 15)           # 120 RH cycles, then under the limit
    wet = ((t_el % 75) < 30).astype(float)
    step('d4_' + hold, t_el, {'E_Stack (V)': np.where(t_el < 9000, 0.94, 0.65), 'Wet_Anode': wet,
                               'Wet_Cathode': wet}, desc=TERM)
    measured('d6_CV-2mVs_0o12to0o8V_80C-100RH_0o4H2-0o4N2_0kPa.fcd', write_cv, 4.0)
    step('zz2_shut-down_25C-Dry_0o4N2-0o4N2_0kPa.fcd', np.arange(1, 301.0), {'E_Stack (V)': 0.01},
         desc='Open Circuit; 5 Minutes; Fixed 1 Sec/Pt; ')
    touch(d / '261001_SYN-FCD-1.txt', b'run notes\r\n')
    touch(d / 'readme.fcd', b'not a step\r\n')
    return d


fd = fcd_run(tmp / 'fcd_run').parent        # the upload: the run's folder inside it
tl = build_timeline(str(fd))
check('fcd run: every step read, the stray .fcd skipped, the notes ignored',
      (len(tl.files), tl.skipped, tl.ignored),
      (12, [('261001_SYN-FCD-1/readme.fcd', 'name does not follow the step-file pattern for .fcd '
             'files')], {'.txt': 1}))
check('... roles', {f.step: f.role for f in tl.files},
      {'a1': 'other', 'a6': 'crossover', 'a10': 'ecsa', 'a12': 'eis', 'b14': 'polcurve',
       'c3': 'air-step', 'c4': 'hold', 'c6': 'crossover', 'c8': 'polcurve', 'd4': 'hold',
       'd6': 'crossover', 'zz2': 'other'})
check('... the header read: format, area, the hold\'s own E limit',
      ({f.source for f in tl.files}, {f.area_cm2 for f in tl.files},
       [(f.step, f.term_v) for f in tl.files if f.term_v is not None]),
      ({'fcd'}, {5.0}, [('c4', 0.7), ('d4', 0.7)]))
check('... two holds, said with their RH cycling',
      [(s.conditions, [f.step for f in s.files]) for s in tl.segments],
      [('90 °C, 100↔0% RH cycling, H2/Air 0.1/0.1 slpm, 0 kPa', ['c4']),
       ('90 °C, 100↔0% RH cycling, H2/Air 0.1/0.1 slpm, 0 kPa', ['d4'])])
check('... the clock: each hold ends with its last sample, no tail added',
      ([round(f.length_s, 6) for f in tl.holds], round(tl.hours_tested, 6)), ([10800.0, 9000.0], 5.5))
check('... RH cycles from the Wet flags: 144 and 120, the cycle 75 s',
      ([(f.rh_cycles, f.C_start, f.C_end) for f in tl.holds], tl.cycles_tested, tl.rh_period_s),
      ([(144, 0, 144), (120, 144, 264)], 264, 75.0))
check('... cycles completed by hold hours: none at the start, a cycle ends 30 s in, 144 by '
      'the first hold\'s end and through the characterization after it',
      [tl.cycles_at(h) for h in (0.0, 29 / 3600, 30 / 3600, 3.0, 3.0 + 29 / 3600, 5.5, 6.0)],
      [0, 0, 1, 144, 144, 264, 264])
check('... and no file read as starting late', [f.step for f in tl.files if f.note], [])

titles = []
_timeline = FD.plot_timeline
FD.plot_timeline = lambda tl_, path, title='', verdict=None: (
    titles.append(title), _timeline(tl_, path, title=title, verdict=verdict))[1]
try:
    res = quiet(run, str(fd), str(tmp / 'out_fcd'), {'image_format': 'png'})
finally:
    FD.plot_timeline = _timeline
row = res['summary'][0]
check('fcd run(): dead where the station ended the second hold, hours, files',
      (res['verdict'], row['Hold files'], row['Step files']),
      ('dead at 5.50 h, OCV (station limit)', 2, 12))
check('... the summary row: hours and RH cycles, tested and to death, the cause said',
      (row['Hours to death'], row['RH cycles tested'], row['RH cycles to death'], row['Death cause'],
       row['OCV death']), (5.5, 264, 264, 'OCV (station limit)', 5.5))
close('... OCV initial: the first hold\'s settled start', row['OCV initial'], 0.95 - 0.004 * 25 / 60, 2e-4)
close('... OCV final: the last 30 min of the second', row['OCV final'], 0.94, 1e-9)
close('... crossover initial from the BOL .fcd CV', row['j_xover initial'], 3.0, 1e-3)
close('... and final from the last hold\'s', row['j_xover final'], 4.0, 1e-3)
check('... the second hold\'s run ending in a wet half under its E limit is an OCV death, and says so',
      ([n for n in res['notes'] if 'own limit' in n], res['ocv_death_h']),
      (['OCV death at 5.50 h in hold 2 (d4), step d4: the station\'s own limit (Term:E<0.7) tripped '
        'in a wet half as the run ended; the OCV was under it 15 s in all (lowest reading 0.6500 V); '
        'the 10-min median never fell under 0.7 V in it'],
       5.5))
check('... the plots named for the run\'s folder, as .fcd names carry no sample', titles,
      ['Durability test timeline — 261001_SYN-FCD-1'])
check('... every plot written, the RH cycles one too',
      sorted(p.stem for p in (tmp / 'out_fcd').glob('*.png')),
      sorted(['fcdur_timeline', 'fcdur_ocv_vs_time', 'fcdur_xover_vs_time', 'fcdur_hfr_vs_time',
              'fcdur_ecsa_vs_time', 'fcdur_perf_vs_time', 'fcdur_polcurve_evolution',
              'fcdur_xover_evolution', 'fcdur_nyquist_evolution', 'fcdur_ecsa_evolution',
              'fcdur_rh_cycles']))
rc = json.loads((tmp / 'out_fcd' / '_plot_data' / 'fcdur_rh_cycles.json').read_text(),
                parse_constant=_no_nan)
rc_line = next(ln for ln in rc['data']['axes'][0]['lines'] if ln.get('label') == 'RH cycles completed')
check('... its stored map: cycles against hold hours, each hold\'s ends, its first and last cycle '
      'and every 25th',
      (rc_line['x'][:4], rc_line['y'][:4], rc_line['x'][7:10], rc_line['y'][7:10], rc_line['x'][-1],
       rc_line['y'][-1], len(rc_line['x']), round(float(np.interp(4.0, rc_line['x'], rc_line['y'])), 2)),
      ([0.0, 0.00833, 0.50833, 1.02917], [0.0, 1.0, 25.0, 50.0], [2.9875, 3.0, 3.00833],
       [144.0, 144.0, 145.0], 5.5, 264.0, 16, round(144 + float(np.interp(3600, [1830, 3705], [25, 50])), 2)))
rv = extract_values(rc)
plain = lambda v: v.get('value') if isinstance(v, dict) else v
check('... and its readout', (plain(rv.get('RH cycles tested')), plain(rv.get('RH cycle')),
                              plain(rv.get('RH cycles to death'))), (264, 75, 264))
ov = extract_values(json.loads((tmp / 'out_fcd' / '_plot_data' / 'fcdur_ocv_vs_time.json').read_text(),
                               parse_constant=_no_nan))
check('... the OCV plot in whole RH cycles, its bin 75 s', (ov.get('OCV plot bin') or {}).get('value'),
      1.25)
wb = openpyxl.load_workbook(tmp / 'out_fcd' / 'fc_durability.xlsx', read_only=True)
summary = {r_[0]: r_[1] for r_ in wb['Summary'].iter_rows(min_row=2, values_only=True)}
check('workbook: the format, no tail, the notes file ignored',
      (summary['Step file format'], summary['Step tail added per hold file (s)'],
       summary['Ignored files']),
      ('Scribner FuelCell (.fcd)', '0 (.fcd steps end with their last sample)', '1 .txt'))
tlx = json.loads((tmp / 'out_fcd' / '_plot_data' / 'fcdur_timeline.json').read_text(),
                 parse_constant=_no_nan)['data']['axes'][0]
tl_lines = {ln['label']: ln['x'] for ln in tlx['lines']}
check('... the time counted: OCV hold time only, said in the Summary, and the timeline drawn on it '
      '(Corey, 13 Oct): the holds end to end, each crossover a tick at its hold hour, the death a '
      'red line',
      (summary['Hours tested (h)'], summary['Time counted'], tlx['xlabel'], tl_lines['OCV hold'],
       [round(x, 4) for x in pulse_mid(tl_lines['H₂ crossover'])],
       [(v_['x'], v_['label']) for v_ in tlx['axvlines']]),
      (5.5, 'OCV hold time only: characterizations and every other step between holds add none',
       'Hold hours (h): OCV hold time only; every other step is a tick at the hold hour it ran at',
       [-0.033, 0.0, 0.0, 3.0, 3.0, 3.0, 3.0, 5.5, 5.5, 5.533], [0.0, 3.0, 5.5],
       [(5.5, 'death at 5.50 h')]))
check('... and the OCV death rule says the station\'s limit counts, and when',
      summary['OCV death rule'].split('. Also ')[1],
      "a death, however short: a .fcd hold's own voltage limit (Term:E<...) acting -- the station "
      "ending a step whose RH is steady on it, or, in an RH-cycling loop, the limit tripping in a "
      "half of its sub-step: a reading under it ending the half, or the half ending a sample and a "
      "half or more short of the header's time with a reading within 20 mV of the limit at its end. "
      "Dated from the stretch under the limit; a trip counts once the OCV has reached the limit in "
      "the hold and settled, or from the first reading if it never does. Samples above the current "
      "limit are left out of it too")
man = list(wb['Manifest'].iter_rows(values_only=True))
rh = {m[2]: m[man[0].index('RH (%)')] for m in man[1:]}
check('... the manifest says the RH a number cannot', (rh['c4'], rh['a1'], rh['zz2']),
      ('100↔0% RH cycling', 100, 0))


def sheet(name):
    rows_ = list(wb[name].iter_rows(values_only=True))
    return [dict(zip(rows_[0], r_)) for r_ in rows_[1:]]


trace_rows = sheet('OCV trace')
trace_cyc = [r_['RH cycles'] for r_ in trace_rows]
check('... RH cycles beside the hours: Summary and verdict, each hold and its files, each '
      'characterization, the station death, the manifest',
      ([(k, summary.get(k)) for k in ('RH cycles tested', 'RH cycle (s)', 'RH cycles to death',
                                      'Death cause')],
       [(r_['RH cycles'], r_['RH cycles before']) for r_ in sheet('Holds')],
       [r_['RH cycles'] for r_ in sheet('Hold files')],
       [(r_['Characterization'], r_['RH cycles before']) for r_ in sheet('Cycles')],
       sorted({(r_['Step'], r_['RH cycles before']) for r_ in sheet('Characterizations')}),
       [(r_['Verdict'], r_['Start (RH cycles)'])
        for r_ in sheet('OCV events')],
       {m['Step']: m['RH cycles'] for m in sheet('Manifest') if m['RH cycles'] is not None}),
      ([('RH cycles tested', 264), ('RH cycle (s)', 75), ('RH cycles to death', 264),
        ('Death cause', 'OCV (station limit)')],
       [(144, 0), (120, 144)], [144, 120],
       [('BOL', 0), ('3 h', 144), ('5.5 h', 264)],
       [('a10', 0), ('a12', 0), ('a6', 0), ('b14', 0), ('c6', 144), ('c8', 144), ('d6', 264)],
       [('death: the station\'s own limit tripped in a wet half as the run ended (lowest sample '
         'shown)', 264)],
       {'c4': 144, 'd4': 120}))
# Cycle n of a hold ends 30 s + (n - 1) × 75 s into it: c4 has 144, d4 120.
done_by = lambda h: (0 if h < 3 else 144) + max(0, min(144 if h < 3 else 120,
                                                       math.floor(((h % 3 if h < 3 else h - 3) * 3600 - 30)
                                                                  / 75 + 1e-9) + 1))
want_cyc = [done_by(r_['Hold hours (h)']) for r_ in trace_rows]
check('... and the OCV trace: cycles completed at each bin, the second column, from 0 to 264',
      (list(trace_rows[0])[1], trace_cyc[0], max(trace_cyc), trace_cyc == want_cyc),
      ('RH cycles', 0, 264, True))
wb.close()
res = quiet(run, str(fd), str(tmp / 'out_fcd25'), {'image_format': 'none', 'geo_area': '25'})
check('an area on the form unlike the headers\' SurfaceArea is noted',
      [n for n in res['notes'] if 'SurfaceArea' in n],
      ['the .fcd headers give an active area of 5 cm² (SurfaceArea); this analysis used 25 cm², '
       'the geometric area parameter'])
res = quiet(run, str(fd), str(tmp / 'out_fcd5'), {'image_format': 'none', 'geo_area': '5'})
check('... and not when they agree', [n for n in res['notes'] if 'SurfaceArea' in n], [])

# A .fcd hold logged every 2 min still starts with its step; RH cycling and
# a steady RH at the same temperature and gases are two conditions.
slow = tmp / 'fcd_slow'
write_fcd(slow / 'c4_OCV_90C-100RH_0o1H2-0o1Air_0kPa.fcd', 70000.0 + np.arange(120, 9001.0, 120),
          {'E_Stack (V)': 0.95})
write_fcd(slow / 'c5_OCV-RHC_90C-100RH(90C)-0RH_0o1H2-0o1Air_0kPa.fcd',
          80000.0 + np.arange(10, 7201.0, 10), {'E_Stack (V)': 0.94})
tl = build_timeline(str(slow))
check('a .fcd hold logged every 2 min: its length is its last stamp, no late-start note',
      [(f.role, f.length_s, f.note) for f in tl.files], [('hold', 9000.0, ''), ('hold', 7200.0, '')])
tr = load_hold_trace(tl.segments[0])
check('... and its samples sit on the hold clock from the first interval',
      (float(tr.t_s[0]), float(tr.t_s[-1]), len(tr.t_s)), (120.0, 9000.0, 75))
check('RH cycling is a condition of its own', [s.conditions for s in tl.segments],
      ['90 °C, 100% RH, H2/Air 0.1/0.1 slpm, 0 kPa',
       '90 °C, 100↔0% RH cycling, H2/Air 0.1/0.1 slpm, 0 kPa'])
check('... a hold named for RH cycling whose Wet flags never change counts none, and says so',
      ([f.rh_cycles for f in tl.holds], tl.cycles_tested, [n for n in tl.notes if 'Wet flags' in n]),
      ([0, 0], 0, ['c5: named for RH cycling, but its Wet flags never switch from wet to dry, so no '
                   'RH cycles are counted for it']))
# One that logs no Wet flag at all is said so, not as flags that never switch.
noflag = tmp / 'fcd_noflag'
write_fcd(noflag / 'c5_OCV-RHC_90C-100RH(90C)-0RH_0o1H2-0o1Air_0kPa.fcd',
          80000.0 + np.arange(5, 7201.0, 5), {'E_Stack (V)': 0.94},
          columns=[c for c in FCD_COLS if not c.startswith('Wet_')])
tl_nf = build_timeline(str(noflag))
check('... one that logs no Wet flag at all, and whose OCV holds steady, counts none, and says that',
      ([(f.role, f.wet_logged) for f in tl_nf.files], tl_nf.cycles_tested,
       [n for n in tl_nf.notes if 'Wet flag' in n]),
      ([('hold', False)], 0, ['c5: named for RH cycling, but it logs no Wet flag (Wet_Cathode, '
                              'Wet_Anode), and its OCV never swings 2 mV or more at a steady pace for '
                              '10 cycles, so no RH cycles are counted for it']))
# Why a test with holds counts no RH cycles, for the Summary: none when it
# counts some, or has no holds. A hold named for RH cycling that shows no
# cycle is named.
cyc_dir = tmp / 'fcd_absent_cycling'
write_fcd(cyc_dir / 'c4_OCV-RHC_90C-100RH(90C)-0RH_0o1H2-0o1Air_0kPa.fcd', 100.0 + np.arange(15, 7201.0, 15),
          {'E_Stack (V)': 0.95, 'Wet_Cathode': (np.arange(15, 7201.0, 15) % 75 < 30).astype(float)})
nohold = tmp / 'fcd_absent_nohold'
write_fcd(nohold / 'c4_OCV_90C-100RH_0o1H2-0o1Air_0kPa.fcd', 100.0 + np.arange(15, 1801.0, 15),
          {'E_Stack (V)': 0.95})
check('why no RH cycles: a hold named for RH cycling with no cycle, beside a steady one or alone; none '
      'to say with cycles or without holds',
      [rh_cycles_absent(build_timeline(str(d_))) for d_ in (slow, noflag, cyc_dir, nohold)],
      ['none counted: c5: named for RH cycling, but no cycle is seen in it; the other holds give none',
       'none counted: c5: named for RH cycling, but no cycle is seen in it', None, None])
# A Wet flag column with no reading in it is no flag read: the hold's OCV is
# read for its cycles instead.
blank = tmp / 'fcd_blank_flags'
write_fcd(blank / 'c5_OCV-RHC_90C-100RH(90C)-0RH_0o1H2-0o1Air_0kPa.fcd',
          80000.0 + np.arange(5, 7201.0, 5), {'E_Stack (V)': 0.94, 'Wet_Anode': np.nan,
                                                 'Wet_Cathode': np.nan})
tl_bl = build_timeline(str(blank))
check('... nor one whose Wet flag columns hold no reading',
      ([f.wet_logged for f in tl_bl.holds], [n for n in tl_bl.notes if 'Wet flag' in n],
       rh_cycles_absent(tl_bl)),
      ([False], ['c5: named for RH cycling, but it logs no Wet flag (Wet_Cathode, Wet_Anode), and its '
                 'OCV never swings 2 mV or more at a steady pace for 10 cycles, so no RH cycles are '
                 'counted for it'],
       'none counted: c5: named for RH cycling, but no cycle is seen in it'))
mixed = tmp / 'fcd_absent_mixed'
write_four(mixed / 'c3_S_OCV_90Cell-30RH_0.1WH2-0.1WA.csv', 7200, 5)
write_fcd(mixed / 'c4_OCV_90C-100RH_0o1H2-0o1Air_0kPa.fcd', 100.0 + np.arange(15, 7201.0, 15),
          {'E_Stack (V)': 0.95, 'Wet_Cathode': 1.0})
write_fcd(mixed / 'c5_OCV-RHC_90C-100RH(90C)-0RH_0o1H2-0o1Air_0kPa.fcd',
          9000.0 + np.arange(5, 7201.0, 5), {'E_Stack (V)': 0.95},
          columns=[c for c in FCD_COLS if not c.startswith('Wet_')])
tl_mx = build_timeline(str(mixed))
check('none counted in a mix: the hold named for RH cycling with no cycle named, the steady ones not',
      ([f.role for f in tl_mx.files], rh_cycles_absent(tl_mx)),
      (['hold', 'hold', 'hold'],
       'none counted: c5: named for RH cycling, but no cycle is seen in it; the other holds give none'))
write_fcd(mixed / 'c4_OCV_90C-100RH_0o1H2-0o1Air_0kPa.fcd', 100.0 + np.arange(15, 7201.0, 15),
          {'E_Stack (V)': 0.95, 'Wet_Cathode': (np.arange(15, 7201.0, 15) % 75 < 30).astype(float)})
write_fcd(mixed / 'c6_OCV_90C-100RH_0o1H2-0o1Air_0kPa.fcd', 20000.0 + np.arange(15, 7201.0, 15),
          {'E_Stack (V)': 0.95, 'Wet_Cathode': 1.0})          # held wet: none to count, none missed
res_mx = quiet(run, str(mixed), str(tmp / 'out_absent_mixed'), {'image_format': 'none'})
wb_mx = openpyxl.load_workbook(tmp / 'out_absent_mixed' / 'fc_durability.xlsx', read_only=True)
sum_mx = {r_[0]: r_[1] for r_ in wb_mx['Summary'].iter_rows(min_row=2, values_only=True)}
wb_mx.close()
check('... and beside cycles counted, the holds that could not be counted are named; steady ones are not',
      (res_mx['summary'][0].get('RH cycles tested'), sum_mx.get('RH cycles not counted in')),
      (96, 'c5: named for RH cycling, but no cycle is seen in it'))

# RH cycles from the Wet flags: the cathode's when it cycles, else the
# anode's; every fall from wet to dry is a cycle, at a steady RH too (Corey:
# the stations never switch the flag there, and if one did, it would count);
# and the cycle is the median time between cycle ends, so a pause in the
# cycling does not stretch it.
wf = tmp / 'fcd_wet'
t_w = np.arange(15, 7201.0, 15)                     # 2 h at 15 s
cyc75, cyc150 = ((t_w % 75) < 30).astype(float), ((t_w % 150) < 60).astype(float)
once = (t_w < 3600).astype(float)                   # wet an hour, then dry
paused = np.where((t_w > 1800) & (t_w < 3600), 0.0, cyc75)
STEADY_N, CYCLING_N = 'c4_OCV_90C-100RH_0o1H2-0o1Air_0kPa.fcd', \
    'c4_OCV-RHC_90C-100RH(90C)-0RH_0o1H2-0o1Air_0kPa.fcd'
twice = ((t_w < 1800) | ((t_w >= 2100) & (t_w < 5400)) | (t_w >= 5700)).astype(float)
odd = np.ones_like(t_w)
for a in (10, 20, 60, 70, 150, 160, 300, 310, 400, 405, 440, 470):
    odd[a:a + 2] = 0.0                                # a dozen switches off, at no pace
cyc75_late = ((t_w % 75) >= 15) & ((t_w % 75) < 45)  # the same cycle, 15 s later
wet_cases = [('cathode', CYCLING_N, cyc150, cyc75), ('anode', CYCLING_N, cyc150, np.ones_like(t_w)),
             ('most falls', CYCLING_N, cyc75, once), ('tie', CYCLING_N, cyc75_late.astype(float), cyc75),
             ('once', STEADY_N, once, once), ('once named', CYCLING_N, once, once),
             ('twice', STEADY_N, twice, twice), ('at no pace', STEADY_N, odd, odd),
             ('unnamed', STEADY_N, cyc75, cyc75), ('paused', CYCLING_N, paused, paused)]
got_w = []
for label, nm, wa, wc in wet_cases:
    write_fcd(wf / label / nm, 100.0 + t_w, {'E_Stack (V)': 0.95, 'Wet_Anode': wa, 'Wet_Cathode': wc})
    tl_w = build_timeline(str(wf / label))
    got_w.append((label, tl_w.cycles_tested, tl_w.rh_period_s,
                  float(tl_w.holds[0].rh_ends_s[0]) if tl_w.cycles_tested else None))
check('Wet flags: the flag that falls most, the cathode\'s on a tie; every fall a cycle, whatever the '
      'hold is named and however few (Corey: a switch at a steady RH would count); a pause leaves '
      'the cycle 75 s',
      got_w, [('cathode', 96, 75.0, 30.0), ('anode', 48, 150.0, 60.0), ('most falls', 96, 75.0, 30.0),
              ('tie', 96, 75.0, 30.0), ('once', 1, None, 3600.0), ('once named', 1, None, 3600.0),
              ('twice', 2, 3600.0, 1800.0), ('at no pace', 12, 450.0, 165.0),
              ('unnamed', 96, 75.0, 30.0), ('paused', 73, 75.0, 30.0)])
# Cycles at no steady pace are counted, but give no cycle to bin the OCV
# plot by: it takes the finest bin, and the workbook and readout say why.
res_np = quiet(run, str(wf / 'at no pace'), str(tmp / 'out_no_pace'), {'image_format': 'png'})
wb_np = openpyxl.load_workbook(tmp / 'out_no_pace' / 'fc_durability.xlsx', read_only=True)
sum_np = {r_[0]: r_[1] for r_ in wb_np['Summary'].iter_rows(min_row=2, values_only=True)}
wb_np.close()
rc_np = json.loads((tmp / 'out_no_pace' / '_plot_data' / 'fcdur_rh_cycles.json').read_text(),
                   parse_constant=_no_nan)
check('... cycles at no steady pace: counted, but the plot bin is the finest, and both say why',
      (res_np['summary'][0]['RH cycles tested'], sum_np['RH cycle (s)'],
       sum_np['OCV plot bin (min)'].split(' (')[1].split(');')[0],
       any('RH cycle: no steady pace' in t_['text'] for t_ in rc_np['data']['axes'][0]['texts'])),
      (12, '450, the median: the cycles keep no steady pace',
       'automatic: the finest that keeps about 3,000 bins over the test, at least 1 min; no steady RH '
       'cycle to bin by', True))
check('... while a steady cycle, a pause in it and all, is the one the bin follows, in a hold not '
      'named for it too',
      [(lab, build_timeline(str(wf / lab)).rh_steady) for lab in ('cathode', 'paused', 'unnamed', 'at no pace')],
      [('cathode', True), ('paused', True), ('unnamed', True), ('at no pace', False)])
# A pause in the cycling is a flat stretch on the stored map, both its ends
# kept, so the map is within a cycle of the count through it.
tl_p = build_timeline(str(wf / 'paused'))
mx, my = rh_cycle_map(tl_p)
grid = np.linspace(0.0, 2.0, 241)
check('... and the map through a pause: its ends kept, within a cycle of the count at every 30 s',
      ((0.50417, 25.0) in list(zip(mx.round(5), my)), (1.00833, 26.0) in list(zip(mx.round(5), my)),
       float(np.max(np.abs(np.interp(grid, mx, my) - [tl_p.cycles_at(h) for h in grid]))) < 1.0),
      (True, True, True))

# Cycles to death are those completed by the death, not the test's: a hold
# cycling RH for 3 h whose OCV falls under 0.7 V for good after 2 h.
dc = tmp / 'fcd_cyc_death'
t_d = np.arange(15, 10801.0, 15)
write_fcd(dc / 'c4_OCV-RHC_90C-100RH(90C)-0RH_0o1H2-0o1Air_0kPa.fcd', 100.0 + t_d,
          {'E_Stack (V)': np.where(t_d < 7200, 0.95, 0.6), 'Wet_Anode': ((t_d % 75) < 30).astype(float),
           'Wet_Cathode': ((t_d % 75) < 30).astype(float)})
res_d = quiet(run, str(dc), str(tmp / 'out_cyc_death'), {'image_format': 'none', 'sample_name': 'cyc'})
row_d = res_d['summary'][0]
tl_d = build_timeline(str(dc))
check('a death 2 h into 3 h of RH cycling: cycles to death those done by then, tested all 144',
      (round(row_d['Hours to death'], 2), row_d['Death cause'], row_d['RH cycles to death'],
       tl_d.cycles_at(row_d['Hours to death']), row_d['RH cycles tested']),
      (2.0, 'OCV', 96, 96, 144))

# A column name logged twice does not stop the read.
cols2 = FCD_COLS[:6] + ['E_Stack (V)'] + FCD_COLS[7:]
(fu / 'dup').mkdir()
lines = ['FUELCELL DATA', 'Format: FuelCell5', '\t'.join(cols2), 'End Comments']
lines += ['\t'.join(['5', '0'] + ['0.5'] * 4 + ['0.7'] + ['0.5'] * (len(cols2) - 7))] * 6
(fu / 'dup' / 'a3_OCV_80C-100RH_0o2H2-0o2N2_0kPa.fcd').write_text('\r\n'.join(lines) + '\r\n')
c = fcd_columns(str(fu / 'dup' / 'a3_OCV_80C-100RH_0o2H2-0o2N2_0kPa.fcd'),
                ('Time Stamp', 'Working Electrode Voltage'))
check('a repeated column name: the first is read', (len(c), float(c['Working Electrode Voltage'].iloc[0])),
      (6, 0.5))

# Re-runs appended to one file: FuelCell starts Time (Sec) over in it after a
# restart; the runs are joined end to end, in order, and noted.
rr = tmp / 'fcd_rerun'
t_rr = np.r_[5000.0 + np.arange(10, 10801.0, 10), np.arange(10, 7201.0, 10)]
write_fcd(rr / 'c4_OCV_90C-100RH_0o1H2-0o1Air_0kPa.fcd', t_rr,
          {'E_Stack (V)': np.r_[np.full(1080, 0.95), np.full(720, 0.90)]})
tl = build_timeline(str(rr))
check('appended re-runs: joined end to end, the step as long as both, noted',
      ([(f.role, f.length_s, f.restarts) for f in tl.files], tl.notes),
      ([('hold', 18000.0, 1)],
       ['c4: Time (Sec) starts over 1 time(s): runs of the step appended to the file, joined end '
        'to end']))
tr = load_hold_trace(tl.segments[0])
check('... on the hold clock in order: the first run, then the second',
      (bool(np.all(np.diff(tr.t_s) > 0)), float(tr.t_s[1079]), float(tr.v[1079]),
       float(tr.t_s[1080]), float(tr.v[1080])), (True, 10800.0, 0.95, 10810.0, 0.90))
write_fcd(rr / 'jit' / 'a3_OCV_80C-100RH_0o1H2-0o1N2_0kPa.fcd',
          [100.0, 101.0, 100.999, 102.0, 103.0], {'E_Stack (V)': 0.1})
tl = build_timeline(str(rr / 'jit'))
check('... while a stamp a hair out of order is no restart',
      [(f.restarts, f.length_s, f.note) for f in tl.files], [(0, 4.0, '')])

# A hold the station ends on its own limit before the hold rule's 2 h is
# still a hold: its hours stay on the clock, and what follows it is after it.
cut = tmp / 'fcd_cut'
write_fcd(cut / 'b20_CV-2mVs_0o12to0o8V_80C-100RH_0o4H2-0o4N2_0kPa.fcd', 1.0 + np.arange(50.0),
          {'E_Stack (V)': 0.4})
t_cut = np.arange(10, 5401.0, 10)
write_fcd(cut / 'c4_OCV_90C-100RH_0o1H2-0o1Air_0kPa.fcd', 9000.0 + t_cut,
          {'E_Stack (V)': np.where(t_cut < 5400, 0.93, 0.69)}, desc=TERM)
write_fcd(cut / 'c3_OCV_N2-to-Air_90C-100RH_0o1H2-0o1Air_0kPa.fcd', 8000.0 + np.arange(1, 121.0),
          {'E_Stack (V)': 0.93})
write_fcd(cut / 'c14_CV-2mVs_0o12to0o8V_80C-100RH_0o4H2-0o4N2_0kPa.fcd', 20000.0 + np.arange(50.0),
          {'E_Stack (V)': 0.4})
tl = build_timeline(str(cut))
check('a hold ended on its own limit at 1.5 h: still a hold, its hours on the clock',
      ({f.step: f.role for f in tl.files}, round(tl.hours_tested, 6),
       [(f.step, f.interval) for f in tl.files if f.kind == 'crossover']),
      ({'b20': 'crossover', 'c3': 'air-step', 'c4': 'hold', 'c14': 'crossover'}, 1.5,
       [('b20', 0), ('c14', 1)]))
check('... and noted', [n for n in tl.notes if 'own limit' in n],
      ['c4: counted as a hold though it ran 1.50 h, short of the 2 h hold rule: the station ended '
       'it on its own limit (Term:E<0.7) at 0.690 V'])
write_fcd(tmp / 'fcd_cut_n2' / 'c2_OCV_80C-100RH_0o1H2-0o1N2_0kPa.fcd', 300.0 + np.arange(1, 61.0),
          {'E_Stack (V)': 0.1}, desc=TERM)
check('... an H2/N2 OCV step under the same limit is neither a hold nor cut short, though its OCV '
      'sits under it',
      [(f.step, f.role, f.cut_short) for f in build_timeline(str(tmp / 'fcd_cut_n2')).files],
      [('c2', 'h2n2', False)])
ocv = analyze_ocv(tl)
check('... an OCV death at its last sample, though the 10-min median never got under 0.7 V',
      (round(ocv.death.start_h, 6), ocv.death.trigger, ocv.death.step,
       decide(tl, ocv, analyze_characterizations(tl)).text),
      (1.5, 'station', 'c4', 'dead at 1.50 h, OCV (station limit)'))
cut2 = tmp / 'fcd_cut_now'
write_fcd(cut2 / 'c4_OCV_90C-100RH_0o1H2-0o1Air_0kPa.fcd', 500.0 + np.arange(1, 3.0),
          {'E_Stack (V)': [0.62, 0.61]}, desc=TERM)
write_fcd(cut2 / 'c5_OCV_90C-100RH_0o1H2-0o1Air_0kPa.fcd', 900.0 + np.arange(10, 7201.0, 10),
          {'E_Stack (V)': 0.95}, desc=TERM)
tl = build_timeline(str(cut2))
ocv = analyze_ocv(tl)
check('a hold the station ends at once, its OCV under the limit from the start, is a hold and the '
      'death at its start; a later good hold is a recovery, flagged',
      ([(f.step, f.role, f.length_s) for f in tl.files], round(ocv.death.start_h, 6),
       ocv.death.trigger, ocv.recovered_h is not None),
      ([('c4', 'hold', 2.0), ('c5', 'hold', 7200.0)], round(1 / 3600, 6), 'station', True))
# A step named for a gas change comes up from N2's OCV: under its limit it has
# not failed, so it is neither a hold nor a death, and the note says so.
gc = tmp / 'fcd_gas_change'
write_fcd(gc / 'c3_OCV_N2-to-Air_90C-100RH_0o1H2-0o1Air_0kPa.fcd', 500.0 + np.arange(1, 3.0),
          {'E_Stack (V)': [0.12, 0.11]}, desc=TERM)
write_fcd(gc / 'c4_OCV_90C-100RH_0o1H2-0o1Air_0kPa.fcd', 900.0 + np.arange(10, 7201.0, 10),
          {'E_Stack (V)': 0.95}, desc=TERM)
tl = build_timeline(str(gc))
ocv = analyze_ocv(tl)
check('... but a step named for a gas change (N2-to-Air) ended under its limit is neither, and '
      'says so',
      ([(f.step, f.role, f.gas_change, f.cut_short) for f in tl.files], ocv.death,
       [n_ for n_ in tl.notes if 'gas change' in n_]),
      ([('c3', 'air-step', True, False), ('c4', 'hold', False, False)], None,
       ['c3: the station ended it on its own limit (Term:E<0.7) at 0.110 V, but it is named for a '
        'gas change, so it is neither a hold nor an OCV death']))
gc2 = tmp / 'fcd_gas_change_long'
write_fcd(gc2 / 'c3_OCV_N2-to-Air_90C-100RH_0o1H2-0o1Air_0kPa.fcd', 500.0 + np.arange(10, 7201.0, 10),
          {'E_Stack (V)': np.r_[np.full(719, 0.95), 0.65]}, desc=TERM)
tl = build_timeline(str(gc2))
check('... nor a station death when it runs long enough to be a hold',
      ([(f.step, f.role) for f in tl.files], analyze_ocv(tl).death), ([('c3', 'hold')], None))
# Appended re-runs: each run the station ended is a stop. One that never came
# up after a restart is a death at its start; a stop followed by a good
# re-run is a death all the same, the re-run its recovery.
rr2 = tmp / 'fcd_rerun_stops'
write_fcd(rr2 / 'up' / 'c4_OCV_90C-100RH_0o1H2-0o1Air_0kPa.fcd',
          np.r_[5000.0 + np.arange(10, 3601.0, 10), [10.0, 11.0]],
          {'E_Stack (V)': np.r_[np.full(360, 0.95), [0.11, 0.10]]}, desc=TERM)
write_fcd(rr2 / 'back' / 'c4_OCV_90C-100RH_0o1H2-0o1Air_0kPa.fcd',
          np.r_[5000.0 + np.arange(10, 3601.0, 10), np.arange(10, 7201.0, 10)],
          {'E_Stack (V)': np.r_[np.full(359, 0.95), 0.65, np.full(720, 0.95)]}, desc=TERM)
got_rr = []
for case in ('up', 'back'):
    tl_r = build_timeline(str(rr2 / case))
    o_r = analyze_ocv(tl_r)
    got_rr.append((case, [(f.role, f.restarts) for f in tl_r.files], round(o_r.death.start_h, 4),
                   o_r.death.trigger, o_r.recovered_h is not None))
check('... a re-run that never came up a death at its start; a run the station stopped before a '
      'good re-run a death at the stop, the re-run its recovery',
      got_rr, [('up', [('hold', 1)], round(3610 / 3600, 4), 'station', False),
               ('back', [('hold', 1)], 1.0, 'station', True)])
# The OCV through the limit just at a rollover: the next file of the hold,
# under the limit from its first sample and ended at once, is the death.
ro = tmp / 'fcd_rollover_trip'
write_fcd(ro / 'c4_OCV_90C-100RH_0o1H2-0o1Air_0kPa.fcd', 100.0 + np.arange(10, 10801.0, 10),
          {'E_Stack (V)': np.r_[np.full(1079, 0.95), 0.701]}, desc=TERM)
write_fcd(ro / 'c5_OCV_90C-100RH_0o1H2-0o1Air_0kPa.fcd', 20000.0 + np.arange(1, 3.0),
          {'E_Stack (V)': [0.699, 0.69]}, desc=TERM)
tl = build_timeline(str(ro))
ocv = analyze_ocv(tl)
check('... and the OCV through the limit at a rollover: the next file, ended at once, the death',
      ([(f.step, f.role) for f in tl.files], [len(s_.files) for s_ in tl.segments],
       round(ocv.death.start_h, 6), ocv.death.step),
      ([('c4', 'hold'), ('c5', 'hold')], [2], round(3 + 1 / 3600, 6), 'c5'))
# Once at the limit and then under it, a step the station ends at once is a
# hold, and the death at its fall; a later good hold is a recovery, flagged.
cut2b = tmp / 'fcd_cut_soon'
write_fcd(cut2b / 'c4_OCV_90C-100RH_0o1H2-0o1Air_0kPa.fcd', 500.0 + np.arange(1, 4.0),
          {'E_Stack (V)': [0.71, 0.62, 0.61]}, desc=TERM)
write_fcd(cut2b / 'c5_OCV_90C-100RH_0o1H2-0o1Air_0kPa.fcd', 900.0 + np.arange(10, 7201.0, 10),
          {'E_Stack (V)': 0.95}, desc=TERM)
tl = build_timeline(str(cut2b))
ocv = analyze_ocv(tl)
check('... one that was at it and fell under it at once is a hold, the death at the fall, a later '
      'good hold a recovery, flagged',
      ([(f.step, f.role, f.length_s) for f in tl.files], round(ocv.death.start_h, 6),
       ocv.death.trigger, ocv.recovered_h is not None),
      ([('c4', 'hold', 3.0), ('c5', 'hold', 7200.0)], round(2 / 3600, 6), 'station', True))
# A station death in a hold's first file comes before a median-rule stretch
# in its next: runs are judged in clock order, whichever rule found them.
# A 2 h hold under its limit from its first sample is a hold by its length,
# but the station did not end it as the OCV failed: the OCV rule judges it,
# from the end of the settle, not a station death at its start.
low = tmp / 'fcd_low_hold'
write_fcd(low / 'c4_OCV_90C-100RH_0o1H2-0o1Air_0kPa.fcd', 100.0 + np.arange(10, 7201.0, 10),
          {'E_Stack (V)': 0.6}, desc=TERM)
ocv = analyze_ocv(build_timeline(str(low)))
check('a 2 h hold under its limit throughout: a station death from its first sample, the OCV '
      'rule\'s stretch after it',
      ([(r.trigger, r.verdict) for r in ocv.runs], round(ocv.death.start_h * 3600, 3)),
      ([('station', 'death'), ('median', 'after death')], 10.0))
# A Wet flag that cannot be read is said so, and not also said never to switch.
_ends = FD.fcd_rh_cycle_ends
FD.fcd_rh_cycle_ends = lambda *a, **k: (_ for _ in ()).throw(ValueError('unreadable'))
try:
    tl_e = build_timeline(str(wf / 'once named'))
finally:
    FD.fcd_rh_cycle_ends = _ends
check('... a Wet flag that cannot be read is said so, once, and its OCV is not read in its place',
      [n_ for n_ in tl_e.notes if 'Wet' in n_ or 'OCV' in n_],
      ['c4: its Wet flags could not be read (ValueError: unreadable), so its RH cycles are not counted'])
cut3 = tmp / 'fcd_cut_then_dip'
write_fcd(cut3 / 'c4_OCV_90C-100RH_0o1H2-0o1Air_0kPa.fcd', 500.0 + np.arange(10, 721.0, 10),
          {'E_Stack (V)': np.r_[np.full(60, 0.93), np.full(12, 0.65)]}, desc=TERM)
t_c3 = np.arange(10, 7201.0, 10)
write_fcd(cut3 / 'c5_OCV_90C-100RH_0o1H2-0o1Air_0kPa.fcd', 900.0 + t_c3,
          {'E_Stack (V)': np.where((t_c3 > 1800) & (t_c3 <= 4200), 0.6, 0.95)}, desc=TERM)
ocv = analyze_ocv(build_timeline(str(cut3)))
check('... and comes first when the next file has a median-rule stretch later',
      [(r.trigger, r.verdict, r.step) for r in ocv.runs], [('station', 'death', 'c4'), ('median', 'after death', 'c5')])
st = tmp / 'fcd_two_ends'
write_fcd(st / 'c4_OCV_90C-100RH_0o1H2-0o1Air_0kPa.fcd', 100.0 + np.arange(10, 7201.0, 10),
          {'E_Stack (V)': np.r_[np.full(500, 0.95), np.full(200, 0.66), np.full(19, 0.72), 0.69]},
          desc=TERM)
tl = build_timeline(str(st))
ocv = analyze_ocv(tl)
check('the station death is the run under the limit the hold ends on; the median rule\'s '
      'earlier stretch comes first',
      [(r.trigger, r.verdict, round(r.start_h * 3600)) for r in ocv.runs],
      [('median', 'death', 5010), ('station', 'after death', 7200)])
# A hold the station ends at its last sample, then the crossover after it
# over the limit: both at the same hour, though the hold clock sums the hold
# lengths (7,230 s and 7,200 s) one way and the OCV trace the other.
sb = tmp / 'fcd_station_both'
sb_src = tmp / 'fcd_station_both_src'
for name, j_x, t0_ in (('a6', 3.0, 1000.0), ('c6', 3.5, 30000.0), ('d6', 20.0, 60000.0)):
    write_cv(sb_src / f'{name}.csv', j_x)
    csv_to_fcd(sb_src / f'{name}.csv', sb / f'{name}_CV-2mVs_0o12to0o8V_80C-100RH_0o4H2-0o4N2_0kPa.fcd', t0=t0_)
write_fcd(sb / 'c4_OCV_90C-100RH_0o1H2-0o1Air_0kPa.fcd', 10000.0 + np.arange(10, 7231.0, 10),
          {'E_Stack (V)': 0.95}, desc=TERM)
write_fcd(sb / 'd4_OCV_90C-100RH_0o1H2-0o1Air_0kPa.fcd', 40000.0 + np.arange(10, 7201.0, 10),
          {'E_Stack (V)': np.r_[np.full(719, 0.95), 0.65]}, desc=TERM)
tl = build_timeline(str(sb))
ocv = analyze_ocv(tl)
ca_sb = analyze_characterizations(tl)
v = decide(tl, ocv, ca_sb)
check('a station death at a hold\'s end and the crossover over the limit after it: both, at the '
      'same hour summed two ways',
      (ocv.death.trigger, round(ocv.death.start_h, 9), round(ca_sb.xover_death.H, 9), v.text),
      ('station', round(14430 / 3600, 9), round(14430 / 3600, 9), 'dead at 4.01 h, both (station limit)'))
causes = []
for eps in (-1e-15, 1e-15, 1e-6):
    ocv.death = dc_replace(ocv.death, start_h=ca_sb.xover_death.H + eps)
    causes.append(decide(tl, ocv, ca_sb).cause)
check('... a last-bit difference either way still both; an OCV death 3.6 ms after the crossover '
      'is after it', causes, ['both', 'both', 'crossover'])


# The station's limit inside an RH-cycling loop (KM8-S6: Term:E<0.7 NextStep
# on the 30 s wet sub-step): a wet half it trips in, the loop running on, is
# an OCV death too (Corey, 13 Oct).
def rhc_hold(path, cycles, v_of=lambda c, wet, k: 0.95, cut=lambda c: None, flag_of=None,
             desc=TERM, t0=100.0, cathode_wet=False, cut_dry=lambda c: None, i_of=None, t_start=0.0,
             drop=()):
    """An RH-cycling hold as FuelCell logs it at 1 s: each cycle a 30 s wet
    half, then 45 s dry. v_of(c, wet, k) is the OCV at cycle c's k-th sample
    of that half, cut(c) and cut_dry(c) the sample the station's limit ended
    the wet or dry half at (None: all of it), flag_of(c, wet, k) a Wet flag
    other than the half's, i_of(c, wet, k) a current other than 0, `drop`
    the (c, wet, k) rows not logged; `cathode_wet`: the cathode's flag wet
    throughout, only the anode's cycling. Returns each sample's time on the
    step's clock and the data."""
    ts, vs, ws, cs = [], [], [], []
    t = t_start
    for c in range(cycles):
        for wet, n_ in ((1, 30 if cut(c) is None else cut(c)), (0, 45 if cut_dry(c) is None else cut_dry(c))):
            for k in range(1, n_ + 1):
                t += 1.0
                if (c, wet, k) in drop:
                    continue
                ts.append(t)
                vs.append(v_of(c, wet, k))
                f_ = flag_of(c, wet, k) if flag_of else None
                ws.append(float(wet) if f_ is None else f_)
                cs.append(i_of(c, wet, k) if i_of else 0.0)
    data = {'E_Stack (V)': np.asarray(vs), 'Wet_Anode': np.asarray(ws),
            'Wet_Cathode': 1.0 if cathode_wet else np.asarray(ws), 'I (A)': np.asarray(cs)}
    if path is not None:
        write_fcd(path, t0 + np.asarray(ts), data, desc=desc)
    return np.asarray(ts), data


RHC_HOLD = 'c4_OCV-RHC_90C-100RH(90C)-0RH_0o1H2-0o1Air_0kPa.fcd'
trip = tmp / 'fcd_trips'
one_at = lambda c0, k0, v0=0.65: (lambda c, wet, k: v0 if (c, wet, k) == (c0, 1, k0) else 0.95)
rhc_hold(trip / 'once' / RHC_HOLD, 100, one_at(40, 10), cut=lambda c: 10 if c == 40 else None)
tl = build_timeline(str(trip / 'once'))
ocv = analyze_ocv(tl)
d = ocv.death
check('a wet half the station\'s limit tripped in, the loop running on: a hold, its cycles all '
      'counted, and an OCV death at the trip, however short, flagged as recovered at once',
      ([(f.role, f.length_s, f.cut_short) for f in tl.files], tl.cycles_tested,
       (d.trigger, d.trips, d.stopped, d.run_ended, d.half, round(d.start_h * 3600, 6),
        round(d.minutes, 4), d.lowest_v),
       round(ocv.recovered_h * 3600, 6), decide(tl, ocv, analyze_characterizations(tl)).text),
      ([('hold', 7480.0, False)], 100, ('station', 1, False, False, 'wet', 3010.0, 0.0167, 0.65),
       3011.0, 'dead at 0.84 h, OCV (station limit)'))
check('... and says what the station did, and when the OCV was back',
      ([n_ for n_ in ocv.notes if n_.startswith(('OCV death', 'OCV recovered'))],
       [r_['Verdict'] for r_ in FD.ocv_event_rows(ocv, tl)]),
      (['OCV death at 0.84 h in hold 1 (c4), step c4: the station\'s own limit (Term:E<0.7) tripped '
        'in a wet half and the RH loop ran on; the OCV was under it 1 s in all (lowest reading '
        '0.6500 V); the 10-min median never fell under 0.7 V in it',
        'OCV recovered: the 10-min median never fell under 0.7 V while the station\'s limit acted, '
        'and was at 0.7 V or more after it at 0.84 h. Flagged; the death at 0.84 h stands'],
       ['death: the station\'s own limit tripped in a wet half and the RH loop ran on (lowest sample '
        'shown)']))
res_once = quiet(run, str(trip / 'once'), str(tmp / 'out_trip_once'), {'image_format': 'none'})
wb_once = openpyxl.load_workbook(tmp / 'out_trip_once' / 'fc_durability.xlsx', read_only=True)
sum_once = {r_[0]: r_[1] for r_ in wb_once['Summary'].iter_rows(min_row=2, values_only=True)}
ev_once = list(wb_once['OCV events'].iter_rows(values_only=True))
wb_once.close()
check('... and the workbook: the recovery said to be after the station acted, the median never '
      'having fallen; the events sheet\'s column so named',
      (sum_once['OCV recovered after death at (h)'], ev_once[0][-2:], ev_once[1][-2]),
      ('0.8364, after the station\'s limit last acted: the 10-min median never fell under the '
       'threshold', ('Median at or over threshold after (h)', 'Start (RH cycles)'), 0.8364))
# Under the limit for every wet half of 40 cycles, over it while dry: each
# wet half ended at its first sample, 46 s apart, is one event, not forty; its
# minutes are those the OCV was under the limit, its lowest reading the
# lowest of the event (a dry one here); the 10-min median never falls under
# 0.7 V. The loop's programmed cycle is still the plot bin's, though a third
# of the cycles came short.
rhc_hold(trip / 'every' / RHC_HOLD, 130,
         lambda c, wet, k: (0.60 if (c, wet, k) == (45, 0, 20) else
                            ((0.62 if c == 30 else 0.68) if wet else 0.75) if 20 <= c < 60 else 0.95),
         cut=lambda c: 1 if 20 <= c < 60 else None)
tl = build_timeline(str(trip / 'every'))
ocv = analyze_ocv(tl)
check('wet halves tripped in every cycle for 40 cycles: one event, from the first trip to the last, '
      'its minutes those under the limit; the median rule silent; the plot bin still the loop\'s cycle',
      ([(r_.trigger, r_.verdict, r_.trips, round(r_.start_h * 3600, 6), round(r_.end_h * 3600, 6),
         round(r_.minutes, 4), r_.lowest_v) for r_ in ocv.runs], round(ocv.recovered_h * 3600, 6),
       tl.cycles_tested, tl.rh_steady, ocv.plot_bin_s),
      ([('station', 'death', 40, 1501.0, 3295.0, round(41 / 60, 4), 0.6)], 3296.0, 130, True, 75.0))
check('... and says so', ([n_ for n_ in ocv.notes if n_.startswith('OCV death')],
                          [r_['Verdict'] for r_ in FD.ocv_event_rows(ocv, tl)]),
      (['OCV death at 0.42 h in hold 1 (c4), step c4: the station\'s own limit (Term:E<0.7) tripped '
        'in 40 wet halves up to 0.92 h, the RH loop running on; the OCV was under it 41 s in all '
        '(lowest reading 0.6000 V); the 10-min median never fell under 0.7 V in it'],
       ['death: the station\'s own limit tripped in 40 wet halves, the RH loop running on (lowest '
        'sample shown)']))
# Where the station did not act: a wet reading under the limit as the wet
# half runs its 30 s, a dry one under it (mid-half or ending it), a full wet
# half whose last row has no voltage, one ending a sample short (jitter), or
# one ending early on load (not at open circuit). The 10-min median stays
# over 0.7 V; no death.
quiet_cases = {
    'not cut': dict(v_of=one_at(40, 10)),
    'dry dip': dict(v_of=lambda c, wet, k: 0.65 if (c, wet, k) == (40, 0, 20) else 0.95),
    'dry end': dict(v_of=lambda c, wet, k: 0.65 if (c, wet, k) == (40, 0, 45) else 0.95),
    'no voltage last': dict(v_of=lambda c, wet, k: {(40, 1, 29): 0.65, (40, 1, 30): math.nan}.get(
        (c, wet, k), 0.95)),
    'a sample short': dict(v_of=one_at(40, 29, 0.705), cut=lambda c: 29 if c == 40 else None),
    'on load': dict(v_of=lambda c, wet, k: 0.50 if (c, wet, k) == (40, 1, 10) else
                    0.71 if (c, wet) == (40, 1) else 0.95, cut=lambda c: 10 if c == 40 else None,
                    i_of=lambda c, wet, k: 0.5 if (c, wet, k) == (40, 1, 10) else 0.0),
    'a wet flag in a dry half': dict(flag_of=lambda c, wet, k: 1.0 if (c, wet, k) == (40, 0, 20) else None),
    'a dry flag in a wet half': dict(flag_of=lambda c, wet, k: 0.0 if (c, wet, k) == (40, 1, 15) else None),
    'blank flags ending a wet half': dict(
        flag_of=lambda c, wet, k: math.nan if (c, wet) == (40, 1) and k > 28 else None),
    'a wet half of 28 s': dict(cut=lambda c: 28 if c == 40 else None),
    'rows not logged': dict(drop={(40, 1, 29), (40, 1, 30)}),
}
got_q = []
for label, kw in quiet_cases.items():
    rhc_hold(trip / label / RHC_HOLD, 100, **kw)
    o_q = analyze_ocv(build_timeline(str(trip / label)))
    got_q.append((label, o_q.death, [r_.trigger for r_ in o_q.runs]))
check('... but no death where the station did not act: a wet reading under the limit as the half runs '
      'on, a dry one, a full half without a last voltage, a half a sample short, a half ended on load; '
      'nor where a half only looks short with the OCV far over the limit: a flag switched for a sample, '
      'blank flags, a half that ran 28 s, rows not logged',
      got_q, [(label, None, []) for label in quiet_cases])
# The OCV between 0.712 and 0.717 V, never under the limit: a half that only
# looks short -- a wet flag for a sample, blank flags at either end of a wet
# half -- is no trip, even this close to the limit. Nor is a half that ended
# early in a run's first half (a rollover's, past any settle), or its last
# (the run ending in it), or where no whole half lasts the header's time.
near_v = lambda c, wet, k: 0.712 if wet else 0.717
near_cases = {
    'near: a wet flag in a dry half': dict(flag_of=lambda c, wet, k: 1.0 if (c, wet, k) == (40, 0, 20) else None),
    'near: blank flags ending a wet half': dict(
        flag_of=lambda c, wet, k: math.nan if (c, wet) == (40, 1) and k > 28 else None),
    'near: blank flags starting a wet half': dict(
        flag_of=lambda c, wet, k: math.nan if (c, wet) == (40, 1) and k < 3 else None),
    'near: the run ending in a short wet half': dict(cut=lambda c: 12 if c == 99 else None,
                                                     cut_dry=lambda c: 0 if c == 99 else None),
    'near: no half lasting the header\'s time': dict(cut=lambda c: 12 if c == 40 else None,
                                                     desc='Open Circuit; 60 Seconds; Fixed 1 Sec/Pt; '
                                                          'Term:E<0.7 NextStep'),
}
got_n = []
for label, kw in near_cases.items():
    rhc_hold(trip / label / RHC_HOLD, 100, near_v, **kw)
    got_n.append((label, analyze_ocv(build_timeline(str(trip / label))).death))
near_guards = {
    'near: rows missing at a short half\'s start': dict(cut=lambda c: 12 if c == 40 else None,
                                                        drop={(40, 1, 1), (40, 1, 2)}),
    'near: a dry flag for a sample in a wet half': dict(
        flag_of=lambda c, wet, k: 0.0 if (c, wet, k) == (40, 1, 15) else None),
    'near: a blank flag inside a short half': dict(cut=lambda c: 12 if c == 40 else None,
                                                   flag_of=lambda c, wet, k: math.nan if (c, wet, k) == (40, 1, 6) else None),
}
for label, kw in near_guards.items():
    rhc_hold(trip / label / RHC_HOLD, 100, near_v, **kw)
    got_n.append((label, analyze_ocv(build_timeline(str(trip / label))).death))
rhc_hold(trip / 'near: first half' / RHC_HOLD, 100)
rhc_hold(trip / 'near: first half' / RHC_HOLD.replace('c4_', 'c5_'), 100, near_v,
         cut=lambda c: 12 if c == 0 else None, t0=20000.0)
got_n.append(('near: a rollover\'s first half', analyze_ocv(build_timeline(str(trip / 'near: first half'))).death))
check('... nor where the OCV is near the limit but the half only looks short, or lies at a run\'s '
      'first or last, or no whole half lasts the header\'s time',
      got_n, [(label, None) for label in near_cases] + [(label, None) for label in near_guards]
      + [('near: a rollover\'s first half', None)])
# A loop programmed for over 2 h whose every wet half the limit tripped in
# at once ran 1.28 h: a hold all the same, by its programmed time, and said
# so; a steady hold the station ended with one flag switched in it is still
# a hold the station ended; a hold named for RH cycling whose flags never
# switch is not judged on its limit, and says so.
rhc_hold(trip / 'short loop' / RHC_HOLD, 100, lambda c, wet, k: 0.60, cut=lambda c: 1)
tl = build_timeline(str(trip / 'short loop'))
check('... a loop shortened under 2 h by its tripped wet halves: still a hold, dead from its first '
      'reading, and noted',
      ([(f.role, round(f.length_s / 3600, 4), round(f.lost_s / 60, 2)) for f in tl.files],
       round(analyze_ocv(tl).death.start_h * 3600, 6), [n_ for n_ in tl.notes if 'hold rule' in n_]),
      ([('hold', 1.2778, 48.33)], 1.0,
       ['c4: counted as a hold though it ran 1.28 h, short of the 2 h hold rule: its wet halves fell '
        '48.3 min short of the header\'s 30 s, as they do when its own limit (Term:E<0.7) trips in '
        'them']))
glitch_ = tmp / 'fcd_steady_glitch'
t_g = np.arange(10, 3601.0, 10)
write_fcd(glitch_ / 'c4_OCV_90C-100RH_0o1H2-0o1Air_0kPa.fcd', 100.0 + t_g,
          {'E_Stack (V)': np.r_[np.full(359, 0.95), 0.69], 'Wet_Anode': 1.0,
           'Wet_Cathode': np.where(t_g == 1800, 0.0, 1.0)}, desc=TERM)
tl = build_timeline(str(glitch_))
check('... a 1 h steady hold the station ended, one flag switched in it: still a hold the station ended',
      ([(f.role, f.loop) for f in tl.files], analyze_ocv(tl).death.station_acted()),
      ([('hold', False)], 'the station ended the step on its own limit'))
write_fcd(tmp / 'fcd_steady_dry_glitch' / 'c4_OCV_90C-0RH_0o1H2-0o1Air_0kPa.fcd', 100.0 + t_g,
          {'E_Stack (V)': np.r_[np.full(359, 0.95), 0.69], 'Wet_Anode': 0.0,
           'Wet_Cathode': np.where(t_g == 1800, 1.0, 0.0)}, desc=TERM)
tl = build_timeline(str(tmp / 'fcd_steady_dry_glitch'))
check('... and a steady dry one with a wet flag for a sample: the station\'s end judged all the same',
      ([(f.role, f.loop) for f in tl.files], analyze_ocv(tl).death.station_acted()),
      ([('hold', False)], 'the station ended the step on its own limit'))
write_fcd(tmp / 'fcd_load_end' / 'c4_OCV_90C-100RH_0o1H2-0o1Air_0kPa.fcd', 100.0 + np.arange(10, 7201.0, 10),
          {'E_Stack (V)': np.r_[np.full(719, 0.95), 0.50], 'I (A)': np.r_[np.zeros(719), 0.5]}, desc=TERM)
check('... a steady hold whose last row draws current: not at open circuit, so no stop',
      analyze_ocv(build_timeline(str(tmp / 'fcd_load_end'))).death, None)
# A loop's header giving the whole step's time, not a sub-step's, is no
# lost time; nor is one giving the dry halves' time.
rhc_hold(trip / 'whole-step header' / RHC_HOLD, 8, desc='Open Circuit; 2 Hours; Fixed 1 Sec/Pt; Term:E<0.7 NextStep')
rhc_hold(trip / 'dry header' / RHC_HOLD, 90, desc='Open Circuit; 45 Seconds; Fixed 1 Sec/Pt; Term:E<0.7 NextStep')
check('... a loop\'s header giving the whole step\'s time, or the dry halves\', loses it no time: a '
      '10-min loop and a 1.9 h one stay short of a hold',
      [(f.role, f.lost_s) for lab in ('whole-step header', 'dry header')
       for f in build_timeline(str(trip / lab)).files],
      [('air-step', 0.0), ('air-step', 0.0)])
write_fcd(tmp / 'fcd_named_flat_short' / RHC_HOLD, 100.0 + np.arange(10, 3601.0, 10),
          {'E_Stack (V)': np.r_[np.full(359, 0.95), 0.65]}, desc=TERM)
check('... a 1 h step named for RH cycling, its flags flat, ending under its limit: a loop, so not '
      'promoted to a hold', [(f.role, f.loop) for f in build_timeline(str(tmp / 'fcd_named_flat_short')).files],
      [('air-step', True)])
write_fcd(tmp / 'fcd_named_flat' / RHC_HOLD, 100.0 + np.arange(10, 7201.0, 10),
          {'E_Stack (V)': np.r_[np.full(719, 0.95), 0.65]}, desc=TERM)
tl = build_timeline(str(tmp / 'fcd_named_flat'))
check('... a hold named for RH cycling whose flags never switch: its limit not judged, and said so',
      (analyze_ocv(tl).death, [n_ for n_ in tl.notes if 'never switch' in n_]),
      (None, ['c4: named for RH cycling, but its Wet flags never switch from wet to dry, so no RH '
              'cycles are counted for it, nor is its own voltage limit judged']))
# While the OCV is still coming up at the start of a hold, a trip is no death:
# not before a reading has reached the limit, nor within the settle. A cell
# that never comes up is dead from its first reading.
rise_cases = {
    'rising': (lambda c, wet, k: 0.60 if c == 0 else 0.93, lambda c: 1 if c == 0 else None, 100),
    'rising past the settle': (lambda c, wet, k: 0.60 if c < 16 else 0.93,
                               lambda c: 1 if c < 16 else None, 110),
    'noise in the settle': (lambda c, wet, k: 0.65 if (c, wet, k) == (3, 1, 5) else 0.93,
                            lambda c: 5 if c == 3 else None, 100),
}
got_r = []
for label, (v_of, cut, n_c) in rise_cases.items():
    rhc_hold(trip / label / RHC_HOLD, n_c, v_of, cut=cut)
    got_r.append((label, analyze_ocv(build_timeline(str(trip / label))).death is None))
check('... nor a trip while the OCV is still coming up: before it first reaches the limit, or within '
      'the settle', got_r, [(label, True) for label in rise_cases])
rhc_hold(trip / 'dead from the start' / RHC_HOLD, 160, lambda c, wet, k: 0.60, cut=lambda c: 1)
ocv = analyze_ocv(build_timeline(str(trip / 'dead from the start')))
d = ocv.death
check('... while a cell that never comes up to the limit is dead from its first reading, the median '
      'rule\'s stretch after it',
      (d.trigger, d.trips, round(d.start_h * 3600, 6), round(d.end_h * 3600, 6),
       [(r_.trigger, r_.verdict) for r_ in ocv.runs]),
      ('station', 160, 1.0, 7315.0, [('station', 'death'), ('median', 'after death')]))
# A blank Wet flag belongs to neither half: after a trip, the next flagged
# row still ends the half; on the trip's own row, the half reads as ending
# early, a sample before it.
rhc_hold(trip / 'unread after' / RHC_HOLD, 100, one_at(40, 10), cut=lambda c: 10 if c == 40 else None,
         flag_of=lambda c, wet, k: math.nan if (c, wet, k) == (40, 0, 1) else None)
rhc_hold(trip / 'unread on it' / RHC_HOLD, 100, one_at(40, 10), cut=lambda c: 10 if c == 40 else None,
         flag_of=lambda c, wet, k: math.nan if (c, wet, k) == (40, 1, 10) else None)
check('... a blank flag after a trip, or on it: the trip still read, on its row',
      [(lab, d_.trips, round(d_.start_h * 3600, 6)) for lab in ('unread after', 'unread on it')
       for d_ in [analyze_ocv(build_timeline(str(trip / lab))).death]],
      [('unread after', 1, 3010.0), ('unread on it', 1, 3010.0)])
# A wet half ended early with no reading under the limit logged: FuelCell may
# act on a reading it does not log, so a half two or more samples short of
# the header's 30 s is a trip too.
rhc_hold(trip / 'early' / RHC_HOLD, 100, one_at(40, 12, 0.705), cut=lambda c: 12 if c == 40 else None)
tl = build_timeline(str(trip / 'early'))
ocv = analyze_ocv(tl)
d = ocv.death
check('... a wet half ended 18 s early with no reading under the limit logged, but one within 20 mV '
      'of it: a trip at its end',
      (d.trips, round(d.start_h * 3600, 6), round(d.minutes, 4), d.lowest_v,
       [n_ for n_ in ocv.notes if n_.startswith('OCV death')]),
      (1, 3012.0, 0.0, 0.705,
       ['OCV death at 0.84 h in hold 1 (c4), step c4: the station\'s own limit (Term:E<0.7) tripped '
        'in a wet half and the RH loop ran on; the OCV was under it 0 s in all (lowest reading '
        '0.7050 V); the 10-min median never fell under 0.7 V in it']))
rhc_hold(trip / 'near after' / RHC_HOLD, 100,
         lambda c, wet, k: 0.71 if (c, wet, k) == (40, 0, 1) else 0.73 if (c, wet) == (40, 1) else 0.95,
         cut=lambda c: 12 if c == 40 else None)
d = analyze_ocv(build_timeline(str(trip / 'near after'))).death
check('... the evidence near the limit may be the next half\'s first reading: a trip at the short half\'s end',
      (d.trips, round(d.start_h * 3600, 6)), (1, 3012.0))
# A half ended early whose last row has no voltage, or no current: the trip
# is at its last reading at open circuit, within 20 mV of the limit.
got_b = []
for label, kw in {'no voltage': dict(v_of=lambda c, wet, k: {(40, 1, 11): 0.71, (40, 1, 12): math.nan}.get(
                      (c, wet, k), 0.95)),
                  'no current': dict(v_of=one_at(40, 11, 0.71),
                                     i_of=lambda c, wet, k: math.nan if (c, wet, k) == (40, 1, 12) else 0.0)
                  }.items():
    rhc_hold(trip / label / RHC_HOLD, 100, cut=lambda c: 12 if c == 40 else None, **kw)
    d_ = analyze_ocv(build_timeline(str(trip / label))).death
    got_b.append((label, d_.trips, round(d_.start_h * 3600, 6), d_.lowest_v))
check('... a half ended early on a row with no voltage or no current: the trip at its last reading',
      got_b, [('no voltage', 1, 3011.0, 0.71), ('no current', 1, 3011.0, 0.71)])
# The RH cycle is the loop's, from the cycles whose wet half ran its time: a
# healthy hold, then one whose every wet half was tripped in at once.
rhc_hold(trip / 'healthy then dead' / RHC_HOLD, 100)
rhc_hold(trip / 'healthy then dead' / RHC_HOLD.replace('c4_', 'c5_'), 160, lambda c, wet, k: 0.6,
         cut=lambda c: 1, t0=20000.0)
tl = build_timeline(str(trip / 'healthy then dead'))
check('... the RH cycle from cycles whose wet half ran its time: 75 s, not the tripped ones\' 46 s',
      (tl.cycles_tested, tl.rh_period_s, tl.rh_steady, analyze_ocv(tl).plot_bin_s), (260, 75.0, True, 75.0))
# The limit belongs to the sub-step the header describes: one whose halves
# last its time. Here it gives 45 s, the dry halves'.
rhc_hold(trip / 'dry limit' / RHC_HOLD, 100,
         lambda c, wet, k: {(40, 0, 20): 0.65, (50, 1, 10): 0.65}.get((c, wet, k), 0.95),
         cut=lambda c: 10 if c == 50 else None, cut_dry=lambda c: 20 if c == 40 else None,
         desc='Open Circuit; 45 Seconds; Fixed 1 Sec/Pt; Term:E<0.7 NextStep')
ocv = analyze_ocv(build_timeline(str(trip / 'dry limit')))
check('... a header giving the dry halves\' time puts the limit on them: a dry half tripped in is the '
      'death, a wet half ended under the limit nothing',
      [(r_.trigger, r_.verdict, r_.trips, r_.half, round(r_.start_h * 3600, 6)) for r_ in ocv.runs],
      [('station', 'death', 1, 'dry', 3050.0)])
# Under the limit from the end of a dry half into the wet half it tripped in
# at once: dated from the dry samples, where the OCV went under.
rhc_hold(trip / 'dry into wet' / RHC_HOLD, 100,
         lambda c, wet, k: 0.66 if (c == 40 and not wet and k > 40) or (c == 41 and wet) else 0.95,
         cut=lambda c: 1 if c == 41 else None)
d = analyze_ocv(build_timeline(str(trip / 'dry into wet'))).death
check('... under the limit from the end of a dry half into the wet half tripped in at once: dated '
      'from the dry half', (d.trips, round(d.start_h * 3600, 6), round(d.end_h * 3600, 6)),
      (1, 3071.0, 3076.0))
# Trips more than the median window apart are separate events, the death and
# one after it; closer, one event. Cycles 40 and 47 trip 8.3 min apart, 40 and
# 50 12.1 min apart.
two_at = lambda c1, c2, v2=0.65: (lambda c, wet, k: {(c1, 1, 5): 0.65, (c2, 1, 5): v2}.get((c, wet, k), 0.95))
got_m = []
for label, (c1, c2, v2) in {'8 min': (40, 47, 0.65), '12 min': (40, 50, 0.65),
                             '37 min': (30, 60, 0.60)}.items():
    rhc_hold(trip / label / RHC_HOLD, 110, two_at(c1, c2, v2), cut=lambda c, c1=c1, c2=c2: 5 if c in (c1, c2) else None)
    got_m.append((label, [(r_.verdict, r_.trips, round(r_.start_h * 3600, 6), round(r_.end_h * 3600, 6),
                           r_.lowest_v) for r_ in analyze_ocv(build_timeline(str(trip / label))).runs]))
check('... trips 8 min apart one event; 12 or 37 min apart two, the death and one after it, each '
      'its own lowest',
      got_m, [('8 min', [('death', 2, 3005.0, 3505.0, 0.65)]),
              ('12 min', [('death', 1, 3005.0, 3005.0, 0.65), ('after death', 1, 3730.0, 3730.0, 0.65)]),
              ('37 min', [('death', 1, 2255.0, 2255.0, 0.65), ('after death', 1, 4480.0, 4480.0, 0.6)])])
# Across the files of one hold, and across a re-run appended to one file,
# trips close together are still one event, under the first file's step.
rhc_hold(trip / 'rollover' / RHC_HOLD, 100, one_at(99, 5), cut=lambda c: 5 if c == 99 else None)
rhc_hold(trip / 'rollover' / RHC_HOLD.replace('c4_', 'c5_'), 100, one_at(1, 5),
         cut=lambda c: 5 if c == 1 else None, t0=20000.0)
tl = build_timeline(str(trip / 'rollover'))
ocv = analyze_ocv(tl)
check('... trips either side of a rollover, 2 min apart: one event, the first file\'s step',
      ([len(s_.files) for s_ in tl.segments],
       [(r_.trips, r_.step, round(r_.start_h * 3600, 6), round(r_.end_h * 3600, 6)) for r_ in ocv.runs]),
      ([2], [(2, 'c4', 7430.0, 7555.0)]))
# A restart (a re-run appended to the file) comes up again from its own
# start: a trip as its OCV rises, or within its settle, is no death; a later
# one is.
t_r0, d_r0 = rhc_hold(None, 50)
t_r1, d_r1 = rhc_hold(None, 80, lambda c, wet, k: 0.60 if c == 0 else 0.65 if (c, wet, k) == (20, 1, 5) else 0.93,
                      cut=lambda c: 1 if c == 0 else 5 if c == 20 else None)
write_fcd(trip / 'rerun' / RHC_HOLD, np.r_[100.0 + t_r0, t_r1],
          {k_: np.r_[np.broadcast_to(d_r0[k_], t_r0.shape), np.broadcast_to(d_r1[k_], t_r1.shape)]
           for k_ in d_r0}, desc=TERM)
tl = build_timeline(str(trip / 'rerun'))
ocv = analyze_ocv(tl)
check('... a restart\'s own rise is no death, though the hold was up before it; a trip after its '
      'settle is',
      ([f.restarts for f in tl.files],
       [(r_.verdict, r_.trips, round(r_.start_h * 3600, 6)) for r_ in ocv.runs]),
      ([1], [('death', 1, 5226.0)]))
# A rollover goes on as the file before left off: tripping at the end of one
# file and the start of the next is one event, every trip counted.
rhc_hold(trip / 'tripping over' / RHC_HOLD, 100, lambda c, wet, k: 0.65 if c >= 97 else 0.95,
         cut=lambda c: 1 if c >= 97 else None)
rhc_hold(trip / 'tripping over' / RHC_HOLD.replace('c4_', 'c5_'), 100,
         lambda c, wet, k: 0.65 if c < 2 else 0.95, cut=lambda c: 1 if c < 2 else None, t0=20000.0)
ocv = analyze_ocv(build_timeline(str(trip / 'tripping over')))
check('... tripping across a rollover: one event, the next file\'s first trips counted too',
      [(r_.trips, r_.step, round(r_.start_h * 3600, 6), round(r_.end_h * 3600, 6)) for r_ in ocv.runs],
      [(5, 'c4', 7276.0, 7460.0)])
# The anode's flag when it is the one that cycles, the cathode's wet throughout.
rhc_hold(trip / 'anode' / RHC_HOLD, 100, one_at(40, 10), cut=lambda c: 10 if c == 40 else None,
         cathode_wet=True)
tl = build_timeline(str(trip / 'anode'))
d = analyze_ocv(tl).death
check('... the wet halves read from the flag that cycles, the anode\'s when the cathode\'s stays wet',
      (tl.cycles_tested, d.trips, round(d.start_h * 3600, 6)), (100, 1, 3010.0))
# Tripped in twice, then the run ends in a dry half still under the limit:
# the dry sub-step has no limit known, so that is no stop; the RH loop ran on
# after each trip.
rhc_hold(trip / 'dry end of the run' / RHC_HOLD, 110,
         lambda c, wet, k: 0.66 if c >= 108 else 0.95, cut=lambda c: 1 if c >= 108 else None)
ocv = analyze_ocv(build_timeline(str(trip / 'dry end of the run')))
d = ocv.death
check('... tripped in twice, the run then ending in a dry half under the limit: two trips, no stop',
      (d.trips, d.stopped, d.run_ended, round(d.start_h * 3600, 6), round(d.end_h * 3600, 6),
       [n_ for n_ in ocv.notes if n_.startswith('OCV death')]),
      (2, False, False, 8101.0, 8147.0,
       ['OCV death at 2.25 h in hold 1 (c4), step c4: the station\'s own limit (Term:E<0.7) tripped '
        'in 2 wet halves up to 2.26 h, the RH loop running on; the OCV was under it 47 s in all '
        '(lowest reading 0.6600 V); the 10-min median never fell under 0.7 V in it']))
# The station's act first where a median stretch starts with it: the OCV
# under 0.7 V from the end of the settle for 12 min, every wet half tripped
# in. The stretch is after the death, not a dip.
rhc_hold(trip / 'tie' / RHC_HOLD, 110,
         lambda c, wet, k: 0.65 if 8 <= c < 24 or (c, wet, k) == (7, 0, 45) else 0.95,
         cut=lambda c: 1 if 8 <= c < 24 else None)
ocv = analyze_ocv(build_timeline(str(trip / 'tie')))
check('... a median stretch starting with the station\'s act is after the death, not a dip',
      [(r_.trigger, r_.verdict, round(r_.start_h * 3600)) for r_ in ocv.runs],
      [('station', 'death', 600), ('median', 'after death', 600)])
# What the station did, in words.
acts = [FD.OcvRun(1, 'c4', 1.0, 2.0, 0.0, 0.6, trigger='station', trips=n_, stopped=st_,
                  run_ended=en_, half='wet') for n_, st_, en_ in
        ((0, True, True), (1, False, False), (3, False, False), (1, False, True), (3, False, True),
         (2, True, True))]
check('... and says what the station did',
      [a_.station_acted() for a_ in acts] + [acts[2].station_acted(' (Term:E<0.7)', until=True)],
      ['the station ended the step on its own limit',
       'the station\'s own limit tripped in a wet half and the RH loop ran on',
       'the station\'s own limit tripped in 3 wet halves, the RH loop running on',
       'the station\'s own limit tripped in a wet half as the run ended',
       'the station\'s own limit tripped in 3 wet halves, the last as the run ended',
       'the station\'s own limit tripped in 2 wet halves, then the station ended the step',
       'the station\'s own limit (Term:E<0.7) tripped in 3 wet halves up to 2.00 h, the RH loop '
       'running on'])
# A step named for a gas change is never judged on its limit, nor a hold
# without one in its header.
rhc_hold(trip / 'gas change' / 'c3_OCV-RHC_N2-to-Air_90C-100RH(90C)-0RH_0o1H2-0o1Air_0kPa.fcd', 100,
         one_at(40, 10), cut=lambda c: 10 if c == 40 else None)
tl = build_timeline(str(trip / 'gas change'))
check('... nor is a wet half tripped in in a step named for a gas change',
      ([(f.role, f.gas_change) for f in tl.files], analyze_ocv(tl).death), ([('hold', True)], None))
rhc_hold(trip / 'no limit' / RHC_HOLD, 100, one_at(40, 10), cut=lambda c: 10 if c == 40 else None,
         desc='Open Circuit; 30 Seconds; Fixed 1 Sec/Pt; ')
check('... and none without the limit in the header',
      analyze_ocv(build_timeline(str(trip / 'no limit'))).death, None)

# The timeline on hold hours, at its edges: a hold the station ended within
# seconds is a dark tick at its hour, its end labelled without crowding the
# one before; folders starting late on the axis are named back from their
# boundary only where the name would run off it; and a run without a hold is
# drawn over the time its logs ran.
edge_tl = tmp / 'fcd_timeline_edges'
write_fcd(edge_tl / 'run' / 'c4_OCV_90C-100RH_0o1H2-0o1Air_0kPa.fcd', 100.0 + np.arange(10, 7201.0, 10),
          {'E_Stack (V)': 0.95}, desc=TERM)
write_fcd(edge_tl / 'run' / 'c5_OCV_90C-100RH_0o1H2-0o1N2_0kPa.fcd', 9000.0 + np.arange(1, 61.0),
          {'E_Stack (V)': 0.1})
write_fcd(edge_tl / 'run' / 'c6_OCV_90C-100RH_0o1H2-0o1Air_0kPa.fcd', 9500.0 + np.arange(1, 3.0),
          {'E_Stack (V)': [0.62, 0.61]}, desc=TERM)
write_fcd(edge_tl / 'run' / 'c7_OCV_90C-100RH_0o1H2-0o1N2_0kPa.fcd', 9800.0 + np.arange(1, 61.0),
          {'E_Stack (V)': 0.1})
write_fcd(edge_tl / 'run' / 'c8_OCV_90C-100RH_0o1H2-0o1Air_0kPa.fcd', 10000.0 + np.arange(10, 7201.0, 10),
          {'E_Stack (V)': 0.95}, desc=TERM)
write_fcd(edge_tl / 'zH2' / 'd4_OCV_90C-100RH_0o1H2-0o1Air_0kPa.fcd', 20000.0 + np.arange(10, 1801.0, 10),
          {'E_Stack (V)': np.r_[np.full(179, 0.95), 0.69]}, desc=TERM)
write_fcd(edge_tl / 'zzEOL' / 'e2_CV-2mVs_0o12to0o8V_80C-100RH_0o4H2-0o4N2_0kPa.fcd',
          1.0 + np.arange(50.0), {'E_Stack (V)': 0.4})
tl = build_timeline(str(edge_tl))
_fig = plot_timeline(tl)
lines_x = {ln.get_label(): list(np.round(ln.get_xdata(), 4)) for ln in _fig.axes[0].get_lines()}
tick_ = FD.TIMELINE_TICK * tl.hours_tested
folders_ = sorted((t_.get_text().strip(), t_.get_ha(), round(t_.get_position()[0], 4))
                  for t_ in _fig.texts if not t_.get_text().endswith(' h'))
check('a hold the station ended in 2 s between two others: a dark tick at its hour, the segment ends '
      'labelled where they do not crowd; folders late on the axis each named from its own boundary, '
      'the last back from it, not off the axis',
      ([round(s_.H_end_h - s_.H_start_h, 6) for s_ in tl.segments],
       lines_x['hold shorter than a tick'],
       sorted(t_.get_text() for t_ in _fig.texts if t_.get_text().endswith(' h')), folders_),
      ([2.0, round(2 / 3600, 6), 2.0, 0.5],
       list(np.round([2.0 - tick_ / 2, 2.0 - tick_ / 2, 2.0 + tick_ / 2, 2.0 + tick_ / 2], 4)),
       ['2.00 h', '4.00 h', '4.50 h'],
       [('run', 'left', 0.0), ('zH2', 'left', round(4.0 + 2 / 3600, 4)),
        ('zzEOL', 'right', round(tl.hours_tested, 4))]))
_plt.close(_fig)
tl = build_timeline(str(tmp / 'nohold'))
_fig = plot_timeline(tl)
check('... and a run without a hold drawn over the time its step logs ran, saying why',
      (_fig.axes[0].get_xlabel(), [round(x, 4) for x in _fig.axes[0].get_xlim()]),
      ('No OCV hold found, so no hold hours: elapsed time over every step, from the step logs (h)',
       [0.0, round(tl.files[0].length_s / 3600.0, 4)]))
_plt.close(_fig)

# Back pressure tells holds apart, and says so; crossovers without an RH are
# not judged, and say so; a run of both formats says its holds never match.
bp = tmp / 'fcd_bp'
write_fcd(bp / 'c4_OCV_90C-100RH_0o1H2-0o1Air_0kPa.fcd', 100.0 + np.arange(10, 7201.0, 10),
          {'E_Stack (V)': 0.95})
write_fcd(bp / 'c5_OCV_90C-100RH_0o1H2-0o1Air_150kPa.fcd', 9000.0 + np.arange(10, 7201.0, 10),
          {'E_Stack (V)': 0.80})
tl = build_timeline(str(bp))
check('holds at two back pressures: two conditions, each said with its pressure',
      [s.conditions for s in tl.segments],
      ['90 °C, 100% RH, H2/Air 0.1/0.1 slpm, 0 kPa', '90 °C, 100% RH, H2/Air 0.1/0.1 slpm, 150 kPa'])
noh = tmp / 'fcd_norh'
write_fcd(noh / 'a6_CV-2mVs_80C_0o4H2-0o4N2_0kPa.fcd', 1.0 + np.arange(50.0), {'E_Stack (V)': 0.4})
tl = build_timeline(str(noh))
ca = analyze_characterizations(tl)
check('a crossover named without RH: labelled so, and left out of the rule with a note',
      (ca.series, ca.death_series, [n for n in ca.notes if 'give no RH' in n]),
      (['Crossover, RH not given'], [],
       ['Crossover, RH not given: the file names give no RH, so these crossovers are not judged '
        'by the 100% RH crossover rule']))
mix = tmp / 'fcd_mix'
write_four(mix / 'c5_S_OCV_90Cell-100RH_0.1WH2-0.1WA.csv', 7200, 30, v=0.95)
write_fcd(mix / 'c6_OCV_90C-100RH_0o1H2-0o1Air_0kPa.fcd', 50.0 + np.arange(10, 7201.0, 10),
          {'E_Stack (V)': 0.80})
res = quiet(run, str(mix), str(tmp / 'out_mix'), {'image_format': 'none'})
check('a run of both formats says its holds never count as the same conditions',
      [n for n in res['notes'] if 'mixes' in n],
      ['this run mixes station CSVs and .fcd files; their names give conditions differently, so '
       'holds of the two never count as the same conditions'])

# Rows Full Analysis leaves out (Ctrl_Mode 5) stay out of a .fcd crossover
# and pol curve: every 10th row put in mode 5 with a junk current reads as
# the same file without those rows.
mode = tmp / 'fcd_mode'
write_cv(mode / 'x.csv', 3.0)
write_stepped_pol(mode / 'p.csv', 0.0)
k_mode, k_i = FCD_COLS.index('Ctrl_Mode'), FCD_COLS.index('I (A)')


def same_values(a, b):
    return sorted(a) == sorted(b) and all(
        (math.isnan(a[k]) and math.isnan(b[k])) or a[k] == b[k] for k in a)


for src, fn in (('x', analyze_crossover_file), ('p', analyze_polcurve_file)):
    csv_to_fcd(mode / f'{src}.csv', mode / f'{src}.fcd')
    lines_m = (mode / f'{src}.fcd').read_bytes().decode('utf-8').split('\r\n')
    first_row = lines_m.index('End Comments') + 1
    marked = set([k for k in range(first_row, len(lines_m)) if lines_m[k]][::10])
    for k in marked:
        row = lines_m[k].split('\t')
        row[k_mode], row[k_i] = '5', '0.5'
        lines_m[k] = '\t'.join(row)
    (mode / f'{src}5.fcd').write_bytes('\r\n'.join(lines_m).encode('utf-8'))
    (mode / f'{src}0.fcd').write_bytes('\r\n'.join(
        ln for k, ln in enumerate(lines_m) if k not in marked).encode('utf-8'))
    a = quiet(fn, str(mode / f'{src}5.fcd'), CharSettings())[0]
    b = quiet(fn, str(mode / f'{src}0.fcd'), CharSettings())[0]
    true(f'{fn.__name__}: rows in Ctrl_Mode 5 read as if they were not there', same_values(a, b),
         (a, b))

# Files a station or a hand edit can leave: 'End Comments' inside a comment,
# rows ending in a tab.
odd = tmp / 'fcd_odd'
write_fcd(odd / 'a3_OCV_80C-100RH_0o2H2-0o2N2_0kPa.fcd', 100.0 + np.arange(12.0), {'E_Stack (V)': 0.1})
text = (odd / 'a3_OCV_80C-100RH_0o2H2-0o2N2_0kPa.fcd').read_bytes().decode('utf-8')
text = text.replace('comment 1: ""', 'comment 1: "see End Comments below"')
text = text.replace('\r\n1', '\t\r\n1')          # every data row gets a trailing tab
(odd / 'a3_OCV_80C-100RH_0o2H2-0o2N2_0kPa.fcd').write_bytes(text.encode('utf-8'))
tl = build_timeline(str(odd))
check('\'End Comments\' in a comment and rows ending in a tab are read past',
      ([(f.step, f.n_rows, f.t_last_s) for f in tl.files], tl.skipped), ([('a3', 12, 12.0)], []))

# Nothing readable: the error says what FC Durability reads and what it found.
none_dir = tmp / 'fcd_none'
touch(none_dir / 'readme.fcd', b'x\r\n')
touch(none_dir / 'notes.txt', b'x\r\n')
try:
    quiet(run, str(none_dir), str(tmp / 'out_none'), {'image_format': 'none'})
    msg = ''
except RuntimeError as e:
    msg = str(e)
check('no step files: both formats named, and what was found',
      msg, 'No step files found. FC Durability reads station CSVs named like '
           'c5_<sample>_OCV_95Cell-30RH_0.1WH2-0.2WA.csv and Scribner FuelCell files named like '
           'c4_OCV_90C-100RH_0o1H2-0o1Air_0kPa.fcd. Found 1 file(s) skipped, e.g. readme.fcd: '
           'name does not follow the step-file pattern for .fcd files; ignored: 1 .txt.')

# The format and tail rows when a run mixes the two, or has no holds.
csv_f, fcd_f = SimpleNamespace(source='csv'), SimpleNamespace(source='fcd')
check('Summary format and tail rows: mixed, station CSV only, .fcd only without holds',
      [(step_formats(SimpleNamespace(files=fs)),
        step_tail_text(SimpleNamespace(files=fs, holds=hs, step_tail_s=30.0)))
       for fs, hs in (([csv_f, fcd_f], [csv_f, fcd_f]), ([csv_f], [csv_f]), ([fcd_f], []))],
      [('station CSV and Scribner FuelCell (.fcd)', '30 (station CSVs); 0 (.fcd)'),
       ('station CSV', 30.0),
       ('Scribner FuelCell (.fcd)', '0 (.fcd steps end with their last sample)')])

# Through the portal: the .fcd run stored with its RH cycles beside its hours,
# twice (two samples) and beside a test that holds RH steady; Comparisons draws
# against cycles off the stored map, and Export stages each plot against them,
# leaving out the steady test as the chart does.
SYN, SYN2, STEADY = '261001_SYN-FCD-1', '261006_SYN-FCD-2', '261005_SYN-STEADY'
jf = submit('job-fcd', fd, {'sample_name': SYN, 'image_format': 'png', 'stand': 'Scribner 1'})
submit('job-fcd2', fd, {'sample_name': SYN2, 'image_format': 'png', 'stand': 'Scribner 1'})
submit('job-steady', slow, {'sample_name': STEADY, 'image_format': 'png', 'stand': 'Scribner 1'})
keys = {s: next(e['bin_id'] for e in store.index['runs'] if e['sample_name'] == s)
        for s in (SYN, SYN2, STEADY)}
kv_of = lambda s: next(e for e in store.index['runs'] if e['sample_name'] == s)['Data'][0]['key_values']
check('stored: RH cycles tested and to death in the index beside the hours; none for a steady RH',
      ({k: v for k, v in kv_of(SYN).items() if k.startswith(('Hours', 'RH'))},
       sorted(k for k in kv_of(STEADY) if k.startswith('RH'))),
      ({'Hours tested': 5.5, 'Hours to death': 5.5, 'RH cycles tested': 264, 'RH cycles to death': 264},
       []))
push_f = jf.get('metrics_push') or {}
check('... all eleven plots, the RH cycles map among them, every sidecar kept',
      ({b: len(p) for b, p in store.bins[keys[SYN]]['metrics'].items()}, push_f.get('sidecars_kept'),
       push_f.get('sidecars_dropped')), ({'fc_durability': 11}, 11, 0))
ask_c = {'analysis': 'fc_durability', 'step': '', 'plot_type': 'fcdur_rh_cycles', 'full': True}
with contextlib.redirect_stdout(io.StringIO()):
    got = viewstore.plot_series_batch([dict(ask_c, key=keys[SYN]), dict(ask_c, key=keys[STEADY])])
sent = next((ln for s in got['series'] for ax in s['axes'] for ln in ax['lines']
             if ln['label'] == 'RH cycles completed'), {})
check('the map sent to the chart as stored; none for the steady test',
      ((sent.get('x'), sent.get('y')), [m['key'] for m in got['missing']]),
      ((rc_line['x'], rc_line['y']), [keys[STEADY]]))
with contextlib.redirect_stdout(io.StringIO()):
    ocv_plot = {s['key']: s['plot'] for s in viewstore.plot_series_batch(
        [{'key': k, 'analysis': 'fc_durability', 'step': '', 'plot_type': 'fcdur_ocv_vs_time'}
         for k in keys.values()])['series']}
    sources = viewstore.materialize_for_compare(
        [{'key': keys[s], 'plot': ocv_plot[keys[s]], 'label': s, 'band': True, 'x': 'cycles'}
         for s in (SYN, SYN2, STEADY)], tmp / 'cyc_in')
check('Export against RH cycles: both cycling tests staged, the steady one left out',
      [s['label'] for s in sources], [SYN, SYN2])
stored = decode_sidecars(store.bins[keys[SYN]])[ocv_plot[keys[SYN]]]['data']['axes'][0]
staged_ax = plot_compare.load_sidecar(plot_compare.find_sidecar(
    sources[0]['output_dir'], sources[0]['filename']))['data']['axes'][0]
on_map = lambda xs: [round(float(np.interp(x, rc_line['x'], rc_line['y'])), 6) for x in xs]
vlines = on_map([v['x'] for v in stored['axvlines']])
true('... every line, band edge and marker read off the stored map; the characterization and the '
     'death at 144 and 264 cycles',
     len(staged_ax['lines']) == len(stored['lines']) > 0
     and all(np.allclose(a['x'], on_map(b['x']), atol=1e-6) for a, b in zip(
         staged_ax['lines'] + staged_ax['scatters'], stored['lines'] + stored['scatters']))
     and [round(v['x'], 6) for v in staged_ax['axvlines']] == vlines
     and staged_ax['xlabel'] == 'RH cycles' and {144.0, 264.0} <= set(vlines),
     ([v['x'] for v in stored['axvlines']], [v['x'] for v in staged_ax['axvlines']]))
seen_c = []
real_render = compare_mod.render_overlay_comparison


def spy_cyc(items, **kw):
    fig = real_render(items, **kw)
    ax0 = fig.axes[0]
    seen_c.append((ax0.get_xlabel(), [it['label'] for it in items],
                   max(float(np.max(ln.get_xdata())) for ln in ax0.get_lines() if len(ln.get_xdata()))))
    return fig


compare_mod.render_overlay_comparison = spy_cyc
try:
    with contextlib.redirect_stdout(io.StringIO()):
        got = compare_mod.run(str(tmp / 'cyc_in'), str(tmp / 'cyc_out'),
                              {'sources': json.dumps(sources), 'image_format': 'png'})
finally:
    compare_mod.render_overlay_comparison = real_render
check('... the figure drawn against RH cycles, to the death at 264',
      [(x, labels, round(m, 6)) for x, labels, m in seen_c], [('RH cycles', [SYN, SYN2], 264.0)])
wb = openpyxl.load_workbook(next((tmp / 'cyc_out').rglob('*.xlsx')), read_only=True)
heads = {c for ws in wb.worksheets for row in ws.iter_rows(min_row=3, max_row=3, values_only=True)
         for c in row if c}
wb.close()
true('... and its workbook\'s traces against RH cycles', 'RH cycles' in heads and 'Hold time (h)' not in heads,
     heads)

# Export pairs a durability plot's panels by label: a test stored before
# j @ 0.65 V replaced j @ 0.6 V draws nothing on the j @ 0.65 V panel, and one
# missing its first panel still lands its second on the second.
perf_name = next(n for n, sc in decode_sidecars(store.bins[keys[SYN]]).items()
                 if sc.get('plot_type') == 'fcdur_perf_vs_time')
perf = decode_sidecars(store.bins[keys[SYN]])[perf_name]
old = json.loads(json.dumps(perf))
old_j = next(a for a in old['data']['axes'] if str(a.get('ylabel', '')).startswith('j @'))
old_j['ylabel'] = 'j @ 0.6 V (A/cm²)'
old_v = next(a for a in old['data']['axes'] if str(a.get('ylabel', '')).startswith('V @'))
for ln in old_j['lines'] + old_v['lines']:    # its own values, to tell it apart when drawn
    ln['y'] = [v * 3 if isinstance(v, (int, float)) else v for v in ln['y']]
lone = json.loads(json.dumps(perf))
lone['data']['axes'] = [a for a in lone['data']['axes'] if str(a.get('ylabel', '')).startswith('j @')]
paired = compare_mod._panels_by_label([{'label': n, 'sidecar': sc} for n, sc in
                                       (('new', perf), ('old', old), ('lone', lone))])
panel_of = lambda it: [(a.get('ylabel'), len(a['lines']) + len(a.get('scatters') or []))
                       for a in it['sidecar']['data']['axes']]
ref_panels = [(a.get('ylabel'), len(a['lines']) + len(a.get('scatters') or []))
              for a in perf['data']['axes'] if not a.get('is_twin') and (a['lines'] or a.get('scatters'))]
check('the performance panels: V @ 1 A/cm² and j @ 0.65 V, read by label',
      [p for p, _ in ref_panels], ['V @ 1 A/cm² (V)', 'j @ 0.65 V (A/cm²)'])
check('... the old test\'s j @ 0.6 V panel left empty; the lone j panel lands second',
      [panel_of(it) for it in paired],
      [ref_panels, [ref_panels[0], ('j @ 0.65 V (A/cm²)', 0)], [('V @ 1 A/cm² (V)', 0), ref_panels[1]]])
check('... the stored sidecars untouched', (old_j['ylabel'], len(lone['data']['axes'])),
      ('j @ 0.6 V (A/cm²)', 1))
# The decay panel's label says the rate window set on the form: two samples
# worked out over different windows still share it.
ocv_sc = {'plot_type': 'fcdur_ocv_vs_time', 'data': {'axes': [
    {'subplot_pos': [1, 1], 'ylabel': 'Hold OCV (V)', 'lines': [{'x': [0.5, 1.5], 'y': [0.95, 0.94]}]},
    {'subplot_pos': [2, 1], 'ylabel': 'OCV decay (μV/h)\n10 h windows',
     'lines': [{'x': [10.0, 11.0], 'y': [120.0, 130.0]}]}]}}
win5 = json.loads(json.dumps(ocv_sc))
decay5 = next(a for a in win5['data']['axes'] if str(a.get('ylabel', '')).startswith('OCV decay'))
decay5['ylabel'] = decay5['ylabel'].replace('10 h windows', '5 h windows')
paired_o = compare_mod._panels_by_label([{'label': n, 'sidecar': sc} for n, sc in
                                         (('10 h', ocv_sc), ('5 h', win5))])
check('... and a decay panel worked out over another rate window paired with it all the same',
      [[(str(a.get('ylabel')).split('\n')[-1], len(a['lines']) > 0) for a in it['sidecar']['data']['axes']]
       for it in paired_o],
      [[('Hold OCV (V)', True), ('10 h windows', True)], [('Hold OCV (V)', True), ('5 h windows', True)]])
handed = []
real_render = compare_mod.render_overlay_comparison
compare_mod.render_overlay_comparison = lambda items, **kw: (handed.append(items),
                                                             real_render(items, **kw))[1]
try:
    with contextlib.redirect_stdout(io.StringIO()):
        compare_mod._compare_generic([{'label': 'new', 'filename': 'p.png', 'sidecar': perf},
                                      {'label': 'old', 'filename': 'p.png', 'sidecar': old}],
                                     str(tmp / 'pair_out'), {'image_format': 'png'})
finally:
    compare_mod.render_overlay_comparison = real_render
check('... and Export\'s comparison draws them so',
      [panel_of(it) for it in handed[0]] if handed else None,
      [ref_panels, [ref_panels[0], ('j @ 0.65 V (A/cm²)', 0)]])
with contextlib.redirect_stdout(io.StringIO()):
    fig_p = plot_compare.render_overlay_comparison(paired)
ys_of = lambda ys: tuple(round(float(v), 6) for v in ys if v is not None and np.isfinite(float(v)))
drawn_j = {ys_of(ln.get_ydata()) for ln in fig_p.axes[1].get_lines()}
drawn_v = {ys_of(ln.get_ydata()) for ln in fig_p.axes[0].get_lines()}
new_j = next(a for a in perf['data']['axes'] if str(a.get('ylabel', '')).startswith('j @'))
_plt.close(fig_p)
check('... drawn so: the j @ 0.65 V panel carries the new test\'s values, never the old\'s; the old '
      'test still on the V @ 1 A/cm² panel',
      ({ys_of(ln['y']) for ln in new_j['lines']} <= drawn_j,
       bool({ys_of(ln['y']) for ln in old_j['lines']} & drawn_j),
       {ys_of(ln['y']) for ln in old_v['lines']} <= drawn_v),
      (True, False, True))

# ── 11. RH cycles read off the OCV ───────────────────────────────────
# A hold that logs no Wet flag -- every station CSV, and a .fcd hold whose Wet
# columns are blank -- has its RH cycles read off its OCV (Corey, 15 Oct:
# FCTS's 0% RH steps are "the 10h wet/dry cycling segment"). Wet gas lowers
# the OCV, so under RH cycling it steps between a low level and a high one;
# each rise from the one to the other ends a cycle, at its first dry reading.
print('11. RH cycles read off the OCV')
with contextlib.redirect_stdout(io.StringIO()):
    from scripts.fc_durability_analysis import (ocv_rh_cycle_ends, rh_cycle_rule, rh_cycles_uncounted,
                                                log_zero_s)


def rh_square(t, period=76.0, wet=44.0, amp=0.010, base=0.900, phase=0.0):
    """An OCV under RH cycling, as in Kolon's 0% RH steps: low (wet) for `wet`
    s of each `period` counted from `phase`, `amp` V higher (dry) the rest."""
    t = np.asarray(t, float)
    return base + amp * (((t - phase) % period) >= wet)


def square_rises(t, period=76.0, wet=44.0, phase=0.0):
    """The readings at which rh_square first reads dry after a wet half."""
    t = np.asarray(t, float)
    dry = ((t - phase) % period) >= wet
    return t[1:][~dry[:-1] & dry[1:]]


rng_o = np.random.default_rng(11)
t2h, t10h = np.arange(1.0, 7201.0), np.arange(1.0, 36001.0)     # 2 h and 10 h at 1 s
want_r = square_rises(t2h)
r_o = ocv_rh_cycle_ends(t2h, rh_square(t2h) + rng_o.normal(0, 0.0003, len(t2h)))
check('a 76 s cycle, 44 s of it wet, 10 mV, 0.3 mV of noise: a cycle ends at each first dry reading; '
      'its period and swing, swinging throughout',
      (None if r_o.ends is None else (len(r_o.ends), r_o.ends.tolist() == want_r.tolist()), r_o.period_s,
       round(r_o.swing_v, 4), r_o.share, r_o.why),
      ((95, True), 76.0, 0.01, 1.0, ''))
ends_d = ocv_rh_cycle_ends(t2h, rh_square(t2h, phase=-50.0)).ends
check('... starting in a dry half: the first cycle ends once the first wet half is over',
      (float(ends_d[0]), ends_d.tolist() == square_rises(t2h, phase=-50.0).tolist()), (70.0, True))
glitch = rh_square(t2h)
glitch[np.searchsorted(t2h, [44 + 76 * k + 15 for k in range(3, 90, 7)])] -= 0.015   # out of a dry half
glitch[np.searchsorted(t2h, [76 * k + 20 for k in range(5, 90, 11)])] -= 0.020      # deep into a wet one
glitch[np.searchsorted(t2h, [76 * k + 10 for k in range(8, 90, 13)])] += 0.015      # up, in a wet one
check('a reading dropping out of a dry half, deep into a wet one, or up out of it, ends no cycle and '
      'moves none', ocv_rh_cycle_ends(t2h, glitch).ends.tolist() == want_r.tolist(), True)
dip10 = rh_square(t2h)
for k in range(3, 90, 7):
    dip10[(t2h >= 44 + 76 * k + 12) & (t2h < 44 + 76 * k + 22)] -= 0.015
check('... nor 10 s out of a dry half, a rise from it under half a period after the last',
      ocv_rh_cycle_ends(t2h, dip10).ends.tolist() == want_r.tolist(), True)
spike = rh_square(t2h)
spike[(t2h >= 76 * 40 + 10) & (t2h < 76 * 40 + 20)] += 0.015                     # 10 s up, in a wet half
ends_sp = ocv_rh_cycle_ends(t2h, spike).ends
check('... and 10 s up out of a wet half ends that cycle there, adding none',
      (len(ends_sp), sorted(set(ends_sp.tolist()) ^ set(want_r.tolist()))),
      (95, [76.0 * 40 + 10, 76.0 * 40 + 44]))
ph_o = t2h % 76
shoulder = np.where(ph_o < 30, 0.900, np.where(ph_o < 40, 0.9045, np.where(ph_o < 44, 0.900, 0.910)))
check('a rise that stalls 10 s part-way, short of the high threshold, ends its cycle once over it',
      ocv_rh_cycle_ends(t2h, shoulder).ends.tolist() == want_r.tolist(), True)
between = rh_square(t2h)
between[:43] = 0.9035                    # the first wet half between the thresholds, nearer the low
check('starting between the thresholds, nearer the low level: low, so its first rise ends a cycle',
      len(ocv_rh_cycle_ends(t2h, between).ends), 95)
noise_o = 0.9 + rng_o.normal(0, 0.002, len(t10h))
r_n = ocv_rh_cycle_ends(t10h, noise_o)
check('noise alone counts none, though it swings over 2 mV: it swings at no period, so RH held steady',
      (float(np.subtract(*np.percentile(noise_o, [75, 25]))) > 0.002, r_n.ends, r_n.period_s, r_n.why),
      (True, None, None, ''))
r_st = ocv_rh_cycle_ends(t10h, 0.92 - 1e-7 * t10h + rng_o.normal(0, 0.0002, len(t10h)))
check('a steady hold drifting down under 0.2 mV of noise: none, RH held steady', (r_st.ends, r_st.why),
      (None, ''))
r_sm = ocv_rh_cycle_ends(t2h, rh_square(t2h, amp=0.0015) + rng_o.normal(0, 1e-4, len(t2h)))
check('a 1.5 mV swing at a steady 76 s is under the 2 mV the rule asks: none, its period still read',
      (r_sm.ends, r_sm.period_s, r_sm.swing_v, r_sm.share, r_sm.why), (None, 76.0, 0.0, 0.0, ''))
burst = 0.92 + np.where((t10h >= 18000) & (t10h < 18400), rh_square(t10h) - 0.9, 0.0) \
    + rng_o.normal(0, 0.0002, len(t10h))                   # 400 s of cycling in 10 h
r_b = ocv_rh_cycle_ends(t10h, burst)
check('a steady hold that swings for 400 s: too few rises to count, and RH said held steady, the hold '
      'being steady but for them', (r_b.ends, r_b.period_s, r_b.share < 0.05, r_b.why), (None, 76.0, True, ''))
nine, ten = np.arange(1.0, 728.0), np.arange(1.0, 729.0)        # rises at 44 + 76k s: nine by 727 s
r_9, r_10 = ocv_rh_cycle_ends(nine, rh_square(nine)), ocv_rh_cycle_ends(ten, rh_square(ten))
check('ten cycles are enough to count; nine, the OCV swinging throughout, are too few, and said so: not '
      'steady', (r_9.ends, r_9.period_s, r_9.why, len(r_10.ends)),
      (None, 76.0, 'its OCV swings 10 mV at a 76 s period, but rises from its low level to its high one only '
                   '9 times', 10))
t800 = np.arange(1.0, 7201.0)
r_800 = ocv_rh_cycle_ends(t800, rh_square(t800, 800.0, 320.0))
real_rises = FD._rises
FD._rises = lambda state: np.array([], dtype=int)
try:
    r_nv = ocv_rh_cycle_ends(t2h, rh_square(t2h))
finally:
    FD._rises = real_rises
check('... and one swinging throughout that never rises from low to high, said so',
      (r_nv.ends, r_nv.why), (None, 'its OCV swings 10.1 mV at a 76 s period, but never rises from its low level '
                                    'to its high one'))
check('... nor is a 2 h hold swinging 10 mV every 800 s throughout, nine times',
      (r_800.ends, r_800.period_s, r_800.share, r_800.why),
      (None, 800.0, 1.0, 'its OCV swings 10 mV at a 800 s period, but rises from its low level to its high '
                         'one only 9 times'))
nine_dip = rh_square(nine)
nine_dip[(nine >= 44 + 76 * 4 + 12) & (nine < 44 + 76 * 4 + 22)] -= 0.015     # 10 s out of a dry half
r_9d = ocv_rh_cycle_ends(nine, nine_dip)
check('... and nine with a dip out of a dry half rise ten times, one of them no cycle: too few to count, '
      'and said so', (r_9d.ends, r_9d.why.endswith('10 times, only 9 of them a period or so apart')),
      (None, True))
jit = np.random.default_rng(0).uniform(-20.0, 20.0, 500)
k_j = np.floor(t10h / 76.0).astype(int)
jittery = 0.9 + 0.010 * ((t10h - 76.0 * k_j) >= 44.0 + jit[k_j])     # wet halves from 24 to 64 s
r_j = ocv_rh_cycle_ends(t10h, jittery)
check('rises at no steady pace count none, though the OCV swings 10 mV every 76 s, and are said so',
      (r_j.ends, r_j.period_s,
       r_j.why.startswith('its OCV swings 10 mV at a 76 s period, but rises')
       and r_j.why.endswith('times, at no steady pace a period apart')), (None, 76.0, True))
# Cycling in under half a hold whose rises fail the rules is unreadable too,
# not steady: it rises often enough to be RH cycling.
r_jp = ocv_rh_cycle_ends(t10h, np.where(t10h < 14400, jittery, 0.9) + rng_o.normal(0, 0.0002, len(t10h)))
t4p = np.arange(4.0, 36001.0, 4.0)
r_4p = ocv_rh_cycle_ends(t4p, np.where(t4p < 14400, rh_square(t4p, 30.0, 17.4), 0.9)
                         + rng_o.normal(0, 0.0002, len(t4p)))
check('... and so is such cycling in 4 h of a 10 h hold, as is a 30 s cycle there logged every 4 s: neither '
      'said to hold RH steady',
      [(r_.ends, r_.share < 0.5, r_.why.endswith(tail_))
       for r_, tail_ in ((r_jp, 'at no steady pace a period apart'), (r_4p, 'of them a period or so apart'))],
      [(None, True, True), (None, True, True)])
# A dip out of a dry half below the low threshold, held long enough to count
# as a level, rises again from it; in half the cycles, picked at random, the
# dips' rises are put aside (a third of them) and every cycle is counted, once
# each.
dips_o = rh_square(t10h)
dipped = sorted(np.random.default_rng(5).choice(np.arange(2, 470), 234, replace=False).tolist())
for k in dipped:
    dips_o[(t10h >= 44 + 76 * k + 12) & (t10h < 44 + 76 * k + 22)] -= 0.015
r_dp = ocv_rh_cycle_ends(t10h, dips_o + rng_o.normal(0, 0.0002, len(t10h)))
check('a 10 s dip out of a dry half in half the cycles: every cycle counted, at its first dry reading',
      (len(dipped), None if r_dp.ends is None else r_dp.ends.tolist() == square_rises(t10h).tolist(),
       r_dp.why), (234, True, ''))
# Cycles are read only off a log every 5 s or finer -- six readings to the
# shortest cycle read. Logged more coarsely, a cycle can show as a slower one
# (a 76 s cycle logged every 24 s rises every 6 cycles), and a steady hold
# cannot be told from one that cycles in step with the log: not read at all.
coarse_why = []
for step_ in (6.0, 10.0, 24.0, 28.0, 30.0, 31.0, 60.0, 300.0):
    t_c = np.arange(step_, 37501.0, step_)
    r_c = ocv_rh_cycle_ends(t_c, rh_square(t_c) + rng_o.normal(0, 0.0002, len(t_c)))
    coarse_why.append((r_c.ends, r_c.period_s, r_c.why))
check('logged every 6, 10, 24, 28, 30, 31, 60 or 300 s: too coarse to read, said so, and not said to hold '
      'RH steady',
      coarse_why, [(None, None, f'its OCV is logged only every {s_:g} s, too coarsely to read')
                   for s_ in (6, 10, 24, 28, 30, 31, 60, 300)])
t_10s = np.arange(10.0, 36001.0, 10.0)
r_st10 = ocv_rh_cycle_ends(t_10s, 0.92 - 1e-7 * t_10s + rng_o.normal(0, 0.0002, len(t_10s)))
check('... nor is a steady hold logged every 10 s', (r_st10.ends, r_st10.why),
      (None, 'its OCV is logged only every 10 s, too coarsely to read'))
drift5 = 5.0 * (1 + 1e-8)
t_d5, t_51 = np.arange(1, 1441) * drift5, np.arange(1, 1441) * 5.01
r_d5, r_51 = ocv_rh_cycle_ends(t_d5, rh_square(t_d5)), ocv_rh_cycle_ends(t_51, rh_square(t_51))
check('a stamp drifting a hair over every 5 s is read as logged every 5 s; every 5.01 s, too coarse',
      (None if r_d5.ends is None else r_d5.ends.tolist() == square_rises(t_d5).tolist(), r_51.why),
      (True, 'its OCV is logged only every 5.01 s, too coarsely to read'))
# The gate on rises put aside as glitches still matters within the limit: a
# 30 s cycle logged every 4 s repeats as logged every 60 s, its rises a period
# of 60 s apart only every other time (all of them, counted, would be half).
t4s = np.arange(4.0, 36001.0, 4.0)
r_4s = ocv_rh_cycle_ends(t4s, rh_square(t4s, 30.0, 17.4) + rng_o.normal(0, 0.0002, len(t4s)))
check('a 30 s cycle logged every 4 s, read at a 60 s period: its rises put aside, none counted, and said so',
      (r_4s.ends, r_4s.period_s, r_4s.why.endswith('of them a period or so apart')), (None, 60.0, True))
# The swing is measured with the drift taken out: a steady hold whose OCV
# wobbles 1.6 mV at a steady period as it falls 5 mV an hour, or recovers 20 mV
# with a 1.6 mV wobble, held RH steady.
t_dr = np.arange(1.0, 36001.0)
fall_w = (0.92 - 0.005 * t_dr / 3600 + 0.0008 * np.sin(2 * np.pi * t_dr / 1100.0)
          + rng_o.normal(0, 0.00025, len(t_dr)))
rec_w = (0.90 + 0.020 * (1 - np.exp(-t2h / 2400.0)) + 0.0008 * np.sin(2 * np.pi * t2h / 1100.0)
         + rng_o.normal(0, 0.00025, len(t2h)))
r_fw, r_rw = ocv_rh_cycle_ends(t_dr, fall_w), ocv_rh_cycle_ends(t2h, rec_w)
check('a 1.6 mV wobble at a steady period as the OCV falls 5 mV an hour, or recovers 20 mV: RH held steady, '
      'the wobble a period found but no swing of 2 mV',
      [(r_.ends, r_.period_s is not None, r_.why) for r_ in (r_fw, r_rw)], [(None, True, '')] * 2)
# The drift is a rolling mean: a median over 30 min of a 60 s cycle half of
# it wet, logged every second, holds half its readings in each level give or
# take one, follows the cycle, and would leave no swing.
half60 = ocv_rh_cycle_ends(t2h, rh_square(t2h, 60.0, 30.0) + rng_o.normal(0, 0.0002, len(t2h)))
check('a 60 s cycle, half of it wet, logged every second: every cycle counted',
      None if half60.ends is None else half60.ends.tolist() == square_rises(t2h, 60.0, 30.0).tolist(), True)
slow_sine = 0.9 + 0.003 * np.sin(2 * np.pi * t10h / 600.0) + rng_o.normal(0, 0.0015, len(t10h))
r_ss = ocv_rh_cycle_ends(t10h, slow_sine)
check('a swing too noisy to rise at a steady pace: none, said so', (r_ss.ends, 'at no steady pace' in r_ss.why),
      (None, True))
real_period = FD._ocv_period_s
FD._ocv_period_s = lambda x, dt: 2 * real_period(x, dt)
try:
    doubled = ocv_rh_cycle_ends(t2h, rh_square(t2h))
finally:
    FD._ocv_period_s = real_period
check('rises a steady 76 s apart are no cycles of a 152 s period: cycles come a period apart',
      (doubled.ends, doubled.period_s, doubled.why.endswith('95 times, at no steady pace a period apart')),
      (None, 152.0, True))
smooth = 0.9 + 0.005 * np.sin(2 * np.pi * t10h / 1500.0) + rng_o.normal(0, 0.001, len(t10h))
r_smooth = ocv_rh_cycle_ends(t10h, smooth)
e_smooth, P_smooth = r_smooth.ends, r_smooth.period_s
check('a slow, smooth swing under noise (10 mV every 25 min, 1 mV of noise): read at its period, not '
      'at the wiggles noise leaves near lag 0: 24 cycles, one on the swing it starts on',
      (None if e_smooth is None else len(e_smooth) in (24, 25),
       P_smooth is not None and abs(P_smooth - 1500) <= 75,
       e_smooth is not None and abs(float(np.median(np.diff(e_smooth))) - 1500) <= 30),
      (True, True, True))
under = 0.9 + np.where(t10h < 14400, rh_square(t10h) - 0.9, 0.0) + rng_o.normal(0, 0.0002, len(t10h))
r_u = ocv_rh_cycle_ends(t10h, under)
check('cycling for 4 h of a 10 h hold: the cycles in those 4 h, and the share it swings in',
      (r_u.ends is not None and r_u.ends.tolist() == square_rises(t10h)[:189].tolist(),
       0.38 < r_u.share < 0.45),
      (True, True))
fading = rh_square(t10h) - 0.100 * t10h / t10h[-1] + rng_o.normal(0, 0.0003, len(t10h))
e_fade = ocv_rh_cycle_ends(t10h, fading).ends
check('cycling while the OCV falls 100 mV through the hold: every cycle, its drift taken out',
      e_fade is not None and e_fade.tolist() == square_rises(t10h).tolist(), True)
t5 = np.arange(5.0, 7201.0, 5.0)
r5 = ocv_rh_cycle_ends(t5, rh_square(t5, 75.0, 30.0, 0.02))
check('read every 5 s: a 75 s cycle with 30 s wet halves',
      (None if r5.ends is None else len(r5.ends), r5.period_s,
       r5.ends is not None and r5.ends.tolist() == square_rises(t5, 75.0, 30.0).tolist()),
      (96, 75.0, True))
held = rh_square(t2h)
in_pause = (t2h > 3000) & (t2h < 5400)
held[in_pause] = 0.910                                     # cycling paused 40 min, held dry
rises_held = t2h[1:][(held[:-1] < 0.905) & (held[1:] > 0.905)]
held[in_pause] += rng_o.normal(0, 0.0003, int(in_pause.sum()))
check('a 40 min pause held dry, under 0.3 mV of noise: the cycles either side of it, and the wet half it '
      'cut short, none in it', ocv_rh_cycle_ends(t2h, held).ends.tolist(), rises_held.tolist())
gap_t = np.r_[t2h[:3000], t2h[4040:]]                     # nothing logged from 3,001 s to 4,040 s
check('a stretch missing from the log: the cycles either side of it',
      ocv_rh_cycle_ends(gap_t, rh_square(gap_t)).ends.tolist(), square_rises(gap_t).tolist())
mess_v = rh_square(t2h)
mess_v[::50] = np.nan
order_o = np.random.default_rng(3).permutation(len(t2h))
check('readings out of order, some twice and some blank: read in time order, the same cycles, each at a '
      'reading', ocv_rh_cycle_ends(np.r_[t2h[order_o], t2h[:20]],
                                   np.r_[mess_v[order_o], mess_v[:20]]).ends.tolist(), want_r.tolist())
bad_t = t2h.copy()
bad_t[3000] = 1e9
r_bad = ocv_rh_cycle_ends(bad_t, rh_square(t2h))
check('a stamp far off the rest, too few readings for a grid of them: not read, and said so',
      (r_bad.ends, r_bad.why.startswith('its OCV is logged too sparsely to read: 7,200 readings over')),
      (None, True))
two_h = np.r_[np.arange(1.0, 3601.0), np.arange(34000.0, 37501.0)]
check('a 10.4 h step with 2 h of it logged: too sparse to read, said so',
      ocv_rh_cycle_ends(two_h, rh_square(two_h)).why,
      'its OCV is logged too sparsely to read: 7,101 readings over 10.4 h, mostly 1 s apart')
check('... and one with no reading', ocv_rh_cycle_ends(np.arange(3.0), np.full(3, np.nan)).why,
      'it logged no OCV to read')

# An FCTS folder from station CSVs: 30% RH holds, then 0% RH steps that cycle
# RH, read off their OCV; the second one starts in a dry half.
n_o = '01M12umSynth_2'
fo = tmp / 'fcts_ocv' / '40H_1'
write_four(fo / f'c1_{n_o}_OCV_90Cell-30RH_0.1WH2-0.1WN2.csv', 3600, 10, v=0.11)
write_four(fo / f'c3_{n_o}_OCV_90Cell-30RH_0.1WH2-0.1WA.csv', 7200, 5, v=lambda t: 0.92 - 2e-7 * t)
write_four(fo / f'c4_{n_o}_OCV_90Cell-100RH_0.1WH2-0.1WN2.csv', 3600, 10, v=0.12)
write_four(fo / f'c6_{n_o}_OCV_90Cell-0RH_0.1WH2-0.1DA.csv', 9030, 2, v=lambda t: rh_square(t) - 1e-7 * t)
write_four(fo / f'd3_{n_o}_OCV_90Cell-30RH_0.1WH2-0.1WA.csv', 7200, 5, v=lambda t: 0.915 - 2e-7 * t)
write_four(fo / f'd6_{n_o}_OCV_90Cell-0RH_0.1WH2-0.1DA.csv', 9030, 2,
           v=lambda t: rh_square(t, phase=-50.0) - 1e-7 * t)
tl_o = build_timeline(str(tmp / 'fcts_ocv'))
SWUNG = 'OCV, swinging 10.1 mV every 76 s'
check('an FCTS folder from station CSVs: the 0% RH steps\' cycles read off their OCV, the 30% holds\' '
      'none',
      [(f.step, f.rh_cycles, f.rh_from, f.C_start, f.C_end) for f in tl_o.holds],
      [('c3', 0, '', 0, 0), ('c6', 118, SWUNG, 0, 118), ('d3', 0, '', 118, 118),
       ('d6', 118, SWUNG, 118, 236)])
check('... each cycle ending at its first dry reading on the step\'s own clock; the cycle 76 s, steady',
      (tl_o.holds[1].rh_ends_s.tolist() == square_rises(np.arange(2.0, 9001.0, 2.0)).tolist(),
       tl_o.holds[3].rh_ends_s[:2].tolist(), tl_o.rh_period_s, tl_o.rh_steady),
      (True, [70.0, 146.0], 76.0, True))
c6_o, d6_o = tl_o.holds[1], tl_o.holds[3]
check('... and read off them by hold hours: 13 by 1,000 s into the first, 118 at its end, 26 more by '
      '2,000 s into the second',
      [tl_o.cycles_at(c6_o.H_start_h + 1000 / 3600), tl_o.cycles_at(c6_o.H_end_h),
       tl_o.cycles_at(d6_o.H_start_h + 2000 / 3600), tl_o.cycles_tested], [13, 118, 144, 236])
check('the rule: a rise of the OCV, in a hold that logs no Wet flag',
      rh_cycle_rule(tl_o),
      'a cycle counts once its wet half is over: every rise of the OCV from its low (wet) level to its '
      'high (dry) one, where it swings 2 mV or more at a steady pace for 10 cycles or more, logged every 5 s '
      'or finer, in a hold that logs no Wet flag; cycles, like hours, count in OCV holds only')
tl_fd = build_timeline(str(fd))
check('... and the .fcd run\'s holds named for RH cycling, whose flags fall, not named as uncounted',
      (tl_fd.cycles_tested, rh_cycles_uncounted(tl_fd)), (264, None))
check('... and the .fcd run\'s, counted by its Wet flags, as before',
      rh_cycle_rule(tl_fd),
      "a cycle counts once its wet half is over: every fall from wet to dry of the station's Wet flag (the "
      "one that falls most, the cathode's on a tie), in any hold; cycles, like hours, count in OCV holds only")
late_o = tmp / 'fcts_ocv_late'
write_four(late_o / f'c6_{n_o}_OCV_90Cell-0RH_0.1WH2-0.1DA.csv', 9030, 2, v=rh_square, t0=600.0)
tl_lo = build_timeline(str(late_o))
f_lo = tl_lo.holds[0]
tr_lo = load_hold_trace(tl_lo.segments[0])
k52 = int(np.searchsorted(tr_lo.t_s, 52.0))
check('a log that started late (600 s): its cycles placed by its own span, as its readings are: the '
      'first wet half it logs ends 52 s in',
      (log_zero_s(f_lo), f_lo.rh_ends_s[:2].tolist(), f_lo.rh_cycles,
       float(tr_lo.t_s[k52]), bool(tr_lo.v[k52 - 1] < 0.905 < tr_lo.v[k52])),
      (600.0, [52.0, 128.0], 110, 52.0, True))
# A .fcd hold whose Wet columns are blank is read so too; one whose flag
# reads is not, its OCV swinging or not: the flag says RH held.
t_fb = np.arange(1.0, 7201.0)
fb = tmp / 'fcd_blank_cycling'
write_fcd(fb / 'c5_OCV-RHC_90C-100RH(90C)-0RH_0o1H2-0o1Air_0kPa.fcd', 80000.0 + t_fb,
          {'E_Stack (V)': rh_square(t_fb, 75.0, 30.0, 0.02), 'Wet_Anode': np.nan, 'Wet_Cathode': np.nan},
          desc=TERM)
tl_fb = build_timeline(str(fb))
check('a .fcd hold whose Wet columns are blank: its cycles read off its OCV, and said so, its own '
      'limit not judged without the flag',
      ([(f.rh_cycles, f.rh_from, f.wet_logged) for f in tl_fb.holds], [n for n in tl_fb.notes if 'RH' in n]),
      ([(96, 'OCV, swinging 20 mV every 75 s', False)],
       ['c5: named for RH cycling, but it logs no Wet flag (Wet_Cathode, Wet_Anode): its RH cycles are '
        'read off its OCV, and its own voltage limit is not judged']))
fw = tmp / 'fcd_flag_held'
write_fcd(fw / 'c4_OCV_90C-100RH_0o1H2-0o1Air_0kPa.fcd', 100.0 + t_fb,
          {'E_Stack (V)': rh_square(t_fb, 75.0, 30.0, 0.02), 'Wet_Cathode': 1.0})
tl_fw = build_timeline(str(fw))
check('... one whose Wet flag reads but never falls counts none, though its OCV swings',
      ([f.rh_cycles for f in tl_fw.holds], rh_cycles_absent(tl_fw)),
      ([0], 'none: the Wet flags never fall from wet to dry in a hold (RH held steady)'))
mb = tmp / 'rh_both'
write_four(mb / 'c3_S_OCV_90Cell-0RH_0.1WH2-0.1DA.csv', 7200, 2, v=rh_square)
write_fcd(mb / 'c4_OCV-RHC_90C-100RH(90C)-0RH_0o1H2-0o1Air_0kPa.fcd', 9000.0 + np.arange(15, 7201.0, 15),
          {'E_Stack (V)': 0.95, 'Wet_Cathode': (np.arange(15, 7201.0, 15) % 75 < 30).astype(float)})
tl_mb = build_timeline(str(mb))
check('a station CSV beside a .fcd hold: one counted off its OCV, one by its Wet flag, the rule naming both',
      ([(f.step, f.rh_cycles, f.rh_from) for f in tl_mb.holds], rh_cycle_rule(tl_mb)),
      ([('c3', 94, SWUNG), ('c4', 96, 'Wet flag')],
       "a cycle counts once its wet half is over: every fall from wet to dry of the station's Wet flag "
       "(the one that falls most, the cathode's on a tie), in a hold that logs one; in a hold that logs "
       "none, every rise of the OCV from its low (wet) level to its high (dry) one, where it swings 2 mV "
       "or more at a steady pace for 10 cycles or more, logged every 5 s or finer; cycles, like hours, count "
       "in OCV holds only"))
csv_st = tmp / 'rh_csv_steady'
write_four(csv_st / 'c3_S_OCV_90Cell-30RH_0.1WH2-0.1WA.csv', 7200, 5, v=lambda t: 0.92 - 2e-7 * t)
check('station CSVs whose OCV holds steady, logged every 5 s: why none says RH held steady',
      rh_cycles_absent(build_timeline(str(csv_st))),
      'none: no Wet flag is logged, and the OCV never swings 2 mV or more at a steady pace for 10 cycles (RH '
      'held steady)')
ms = tmp / 'rh_both_steady'
write_four(ms / 'c3_S_OCV_90Cell-30RH_0.1WH2-0.1WA.csv', 7200, 5, v=0.92)
write_fcd(ms / 'c4_OCV_90C-100RH_0o1H2-0o1Air_0kPa.fcd', 9000.0 + np.arange(15, 7201.0, 15),
          {'E_Stack (V)': 0.95, 'Wet_Cathode': 1.0})
check('... and both steady: why none says so for each',
      rh_cycles_absent(build_timeline(str(ms))),
      'none: the Wet flags never fall from wet to dry, and where none is logged, the OCV never swings '
      '2 mV or more at a steady pace for 10 cycles (RH held steady)')
# A hold whose OCV cannot be read: a note says so, and so does the Summary.
ur = tmp / 'rh_unread'
ur.mkdir(parents=True)
(ur / 'c3_S_OCV_90Cell-30RH_0.1WH2-0.1WA.csv').write_text(
    'Time Stamp,Current\n' + ''.join(f'{t_:.1f},0\n' for t_ in np.arange(10.0, 7171.0, 10.0)))
write_four(ur / 'c4_S_OCV_90Cell-0RH_0.1WH2-0.1DA.csv', 7200, 5, v=0.9)
tl_ur = build_timeline(str(ur))
check('a hold whose OCV cannot be read: a note says so, and why no RH cycles names it',
      ([(f.step, f.role, f.rh_unread) for f in tl_ur.holds], [n for n in tl_ur.notes if 'RH cycles' in n],
       rh_cycles_absent(tl_ur)),
      ([('c3', 'hold', True), ('c4', 'hold', False)],
       ["c3: its OCV could not be read for RH cycles (ValueError: no 'Working Electrode Voltage' column), "
        "so they are not counted"],
       'none counted: c3: unreadable (see the notes); the other holds give none'))
write_four(ur / 'c4_S_OCV_90Cell-0RH_0.1WH2-0.1DA.csv', 7200, 2, v=rh_square)
tl_ur2 = build_timeline(str(ur))
check('... beside cycles counted, named as not counted in',
      (tl_ur2.cycles_tested, rh_cycles_uncounted(tl_ur2), rh_cycles_absent(tl_ur2)),
      (94, 'c3: unreadable (see the notes)', None))
# One whose OCV swings but whose cycles cannot be counted -- logged too
# coarsely for its cycle -- is unreadable too, and said so; it did not hold
# RH steady.
coarse = tmp / 'rh_coarse'
write_four(coarse / 'c6_S_OCV_90Cell-0RH_0.1WH2-0.1DA.csv', 37530, 30, v=rh_square)
tl_co = build_timeline(str(coarse))
check('a 0% RH step cycling every 76 s but logged every 30 s: unreadable, the note saying why',
      ([(f.step, f.rh_cycles, f.rh_unread) for f in tl_co.holds], [n_ for n_ in tl_co.notes if 'RH cycles' in n_],
       rh_cycles_absent(tl_co)),
      ([('c6', 0, True)], ['c6: its OCV is logged only every 30 s, too coarsely to read, so its RH cycles '
                           'are not counted'],
       'none counted: c6: unreadable (see the notes)'))
# A hold named for RH cycling whose OCV swings throughout, but too slowly to
# rise ten times, is unreadable, not a hold in which no cycle is seen.
slow8 = tmp / 'rh_slow800'
write_fcd(slow8 / 'c5_OCV-RHC_90C-100RH(90C)-0RH_0o1H2-0o1Air_0kPa.fcd', 80000.0 + t2h,
          {'E_Stack (V)': rh_square(t2h, 800.0, 320.0), 'Wet_Anode': np.nan, 'Wet_Cathode': np.nan},
          desc=TERM)
tl_s8 = build_timeline(str(slow8))
check('a 2 h hold named for RH cycling, its OCV swinging every 800 s, nine times: unreadable, said so',
      ([(f.step, f.rh_cycles, f.rh_unread) for f in tl_s8.holds], [n_ for n_ in tl_s8.notes if 'RH' in n_],
       rh_cycles_absent(tl_s8)),
      ([('c5', 0, True)],
       ['c5: its OCV swings 10 mV at a 800 s period, but rises from its low level to its high one only 9 '
        'times, so its RH cycles are not counted, nor is its own voltage limit judged'],
       'none counted: c5: unreadable (see the notes)'))
# Logged more coarsely than every 30 s, a hold is not read, and says so; a
# hold named for RH cycling whose own limit needs a Wet flag to be placed
# says that limit is not judged either.
coarse60 = tmp / 'rh_coarse60'
t60 = np.arange(60.0, 7201.0, 60.0)
write_fcd(coarse60 / 'c5_OCV-RHC_90C-100RH(90C)-0RH_0o1H2-0o1Air_0kPa.fcd', 80000.0 + t60,
          {'E_Stack (V)': rh_square(t60, 75.0, 30.0, 0.02), 'Wet_Anode': np.nan, 'Wet_Cathode': np.nan},
          desc=TERM)
write_four(coarse60 / 'c6_S_OCV_90Cell-0RH_0.1WH2-0.1DA.csv', 9030, 60, v=rh_square)
write_fcd(coarse60 / 'c3_OCV_90C-0RH_0o1H2-0o1Air_0kPa.fcd', 60000.0 + t60,
          {'E_Stack (V)': rh_square(t60, 75.0, 30.0, 0.02), 'Wet_Anode': np.nan, 'Wet_Cathode': np.nan},
          desc=TERM)
write_fcd(coarse60 / 'c4_OCV-RHC_90C-100RH(90C)-0RH_0o1H2-0o1Air_0kPa.fcd', 70000.0 + t60,
          {'E_Stack (V)': rh_square(t60, 75.0, 30.0, 0.02), 'Wet_Anode': np.nan, 'Wet_Cathode': np.nan})
tl_c60 = build_timeline(str(coarse60))
check('logged every 60 s: .fcd holds with blank flags, with a voltage limit or not and named for RH '
      'cycling or not, and a station CSV, each unreadable, said so, with where its limit is judged',
      ([(f.step, f.rh_cycles, f.rh_unread) for f in tl_c60.holds], [n_ for n_ in tl_c60.notes if 'RH' in n_],
       rh_cycles_absent(tl_c60)),
      ([('c3', 0, True), ('c4', 0, True), ('c5', 0, True), ('c6', 0, True)],
       ['c3: its OCV is logged only every 60 s, too coarsely to read, so its RH cycles are not counted, and '
        'its own voltage limit is judged only where a run ends under it',
        'c4: its OCV is logged only every 60 s, too coarsely to read, so its RH cycles are not counted',
        'c5: its OCV is logged only every 60 s, too coarsely to read, so its RH cycles are not counted, nor '
        'is its own voltage limit judged',
        'c6: its OCV is logged only every 60 s, too coarsely to read, so its RH cycles are not counted'],
       'none counted: c3, c4, c5, c6: unreadable (see the notes)'))
# Cycling in under half of a hold is counted, unless the hold's name gives a
# humidified RH -- a dew point too -- but not one that gives 0% RH, a dry
# gas or no RH at all; the note on one not counted says where its own limit
# is judged.
names = tmp / 'rh_names'
part_v = np.where(t2h < 2880, rh_square(t2h, 75.0, 30.0, 0.02), 0.9)
write_fcd(names / 'c3_OCV_93C-(90C)RH_0o1H2-0o1Air_0kPa.fcd', 30000.0 + t2h,
          {'E_Stack (V)': part_v, 'Wet_Anode': np.nan, 'Wet_Cathode': np.nan}, desc=TERM)
write_fcd(names / 'c4_OCV_25C-Dry_0o1H2-0o1Air_0kPa.fcd', 40000.0 + t2h,
          {'E_Stack (V)': part_v, 'Wet_Anode': np.nan, 'Wet_Cathode': np.nan})
write_fcd(names / 'c5_OCV_83C_0o1H2-0o1Air_0kPa.fcd', 50000.0 + t2h,
          {'E_Stack (V)': part_v, 'Wet_Anode': np.nan, 'Wet_Cathode': np.nan})
tl_nm = build_timeline(str(names))
check('cycling in 40% of a hold: under a dew point, not counted, its limit judged where a run ends; under '
      'Dry, or no RH, counted; each said so',
      ([(f.step, f.rh_label, f.rh_cycles, f.rh_unread) for f in tl_nm.holds],
       [n_ for n_ in tl_nm.notes if 'RH' in n_]),
      ([('c3', 'dew point 90 °C', 0, True), ('c4', '0% RH', 38, False), ('c5', 'RH not given', 38, False)],
       ['c3: its name gives dew point 90 °C, and its OCV swings 20 mV every 75 s in only 40% of it, so the 38 '
        'cycles it shows there are not counted as RH cycles, and its own voltage limit is judged only where a '
        'run ends under it',
        'c4: its OCV swings 2 mV or more in only 40% of it; all 38 RH cycles counted in it fall there',
        'c5: its OCV swings 2 mV or more in only 40% of it; all 38 RH cycles counted in it fall there']))
# Where the cycles read off the OCV need a second look, a note says so: in a
# hold that cycles in under half of its time; in one whose name gives a
# humidified RH; and in a .fcd hold, not named for RH cycling, whose own
# voltage limit its Wet flag would have placed in its halves. Cycling in
# under half of a hold whose name gives a humidified RH is a disturbance, not
# RH cycling: it is not counted, and the hold is unreadable.
looks = tmp / 'rh_looks'
write_four(looks / 'c3_S_OCV_90Cell-30RH_0.1WH2-0.1WA.csv', 36030, 2,
           v=lambda t: np.where(t < 14400, rh_square(t), 0.9))
write_fcd(looks / 'c4_OCV_90C-0RH_0o1H2-0o1Air_0kPa.fcd', 50000.0 + t2h,
          {'E_Stack (V)': rh_square(t2h, 75.0, 30.0, 0.02), 'Wet_Anode': np.nan, 'Wet_Cathode': np.nan},
          desc=TERM)
write_four(looks / 'c5_S_OCV_90Cell-0RH_0.1WH2-0.1DA.csv', 36030, 2,
           v=lambda t: np.where(t < 14400, rh_square(t), 0.9))
write_four(looks / 'c6_S_OCV_90Cell-30RH_0.1WH2-0.1WA.csv', 7200, 2, v=rh_square)
tl_lk = build_timeline(str(looks))
check('... cycling in 4 h of a 10 h hold named 0% RH, or throughout one named 30% RH, and in a .fcd hold '
      'with a voltage limit and blank flags: counted, each noted; in 4 h of one named 30% RH: unreadable, '
      'said so',
      ([(f.step, f.rh_cycles, f.rh_unread, f.loop) for f in tl_lk.holds],
       [n_ for n_ in tl_lk.notes if 'RH' in n_], rh_cycles_uncounted(tl_lk)),
      ([('c3', 0, True, False), ('c4', 96, False, False), ('c5', 189, False, False), ('c6', 94, False, False)],
       ['c3: its name gives 30% RH, and its OCV swings 10.1 mV every 76 s in only 40% of it, so the 189 cycles '
        'it shows there are not counted as RH cycles',
        'c4: it logs no Wet flag (Wet_Cathode, Wet_Anode), so its RH cycles are read off its OCV, and its own '
        'voltage limit is judged only where a run ends under it, not in the wet halves it may trip in',
        'c5: its OCV swings 2 mV or more in only 40% of it; all 189 RH cycles counted in it fall there',
        'c6: its name gives 30% RH, but its OCV swings 10.1 mV every 76 s, as under RH cycling: 94 RH cycles '
        'are counted from it'],
       'c3: unreadable (see the notes)'))
# Readings drawing current are no OCV: they are left out of the read, as the
# OCV rule leaves them out.
loaded = tmp / 'rh_loaded'
loaded.mkdir(parents=True)
t_ld = np.arange(2.0, 7171.0, 2.0)
on_ld = (t_ld >= 3600) & (t_ld < 3720)
v_ld = np.where(on_ld, 0.60, rh_square(t_ld))
(loaded / 'c6_S_OCV_90Cell-0RH_0.1WH2-0.1DA.csv').write_text(
    FOUR + '\n' + ''.join(f'{a_:.1f},1e-05,{b_:.6f},{0.5 if c_ else 0.0:g}\n'
                          for a_, b_, c_ in zip(t_ld, v_ld, on_ld)))
tl_ld = build_timeline(str(loaded))
check('a 2 min load mid-hold: its readings left out of the read, so its cycles are those at open circuit '
      '(the 4 s of a dry half the log resumes in too short to count)',
      (tl_ld.holds[0].rh_ends_s.tolist() == ocv_rh_cycle_ends(t_ld[~on_ld], v_ld[~on_ld]).ends.tolist(),
       tl_ld.holds[0].rh_cycles, ocv_rh_cycle_ends(t_ld, v_ld).ends is None
       or len(ocv_rh_cycle_ends(t_ld, v_ld).ends) != tl_ld.holds[0].rh_cycles),
      (True, 92, True))

# run(): the Summary, the manifest, the OCV plot's bin and the RH cycles plot.
res_o = quiet(run, str(tmp / 'fcts_ocv'), str(tmp / 'out_fcts_ocv'), {'image_format': 'png'})
wb_o = openpyxl.load_workbook(tmp / 'out_fcts_ocv' / 'fc_durability.xlsx', read_only=True)
sum_o = {r_[0]: r_[1] for r_ in wb_o['Summary'].iter_rows(min_row=2, values_only=True)}
man_o = list(wb_o['Manifest'].iter_rows(values_only=True))
wb_o.close()
h_o = man_o[0]
check('run(): the Summary counts them, with the cycle and the rule; the stored row too',
      (sum_o.get('RH cycles tested'), sum_o.get('RH cycle (s)'), sum_o.get('RH cycle rule') == rh_cycle_rule(tl_o),
       'RH cycles not counted in' in sum_o, res_o['summary'][0].get('RH cycles tested')),
      (236, 76, True, False, 236))
check('... the manifest says what each hold\'s cycles came from',
      [(r_[h_o.index('Step')], r_[h_o.index('RH cycles')], r_[h_o.index('RH cycles from')])
       for r_ in man_o[1:] if r_[h_o.index('RH cycles')] is not None],
      [('c6', 118, SWUNG), ('d6', 118, SWUNG)])
check('... the OCV plot binned by whole cycles, and the RH cycles plot drawn',
      (str(sum_o.get('OCV plot bin (min)')).split(';')[0],
       (tmp / 'out_fcts_ocv' / 'fcdur_rh_cycles.png').exists()),
      ('1.267 (automatic: 1 whole RH cycle of 76 s, at least 1 min and no more than about 3,000 bins '
       'over the test)', True))

# Through the portal: stored with its RH cycles and its map, and drawn against
# them as a .fcd test is.
FCTS_S = '261015_SYN-FCTS-OCV'
submit('job-fcts-ocv', tmp / 'fcts_ocv', {'sample_name': FCTS_S, 'image_format': 'png', 'stand': 'FCTS 2'})
key_fo = next(e['bin_id'] for e in store.index['runs'] if e['sample_name'] == FCTS_S)
mx_o, my_o = rh_cycle_map(tl_o)
with contextlib.redirect_stdout(io.StringIO()):
    got_fo = viewstore.plot_series_batch([dict(ask_c, key=key_fo)])
    ocv_fo = next(s['plot'] for s in viewstore.plot_series_batch(
        [{'key': key_fo, 'analysis': 'fc_durability', 'step': '', 'plot_type': 'fcdur_ocv_vs_time'}])['series'])
    src_fo = viewstore.materialize_for_compare(
        [{'key': key_fo, 'plot': ocv_fo, 'label': FCTS_S, 'band': True, 'x': 'cycles'}], tmp / 'cyc_in_fo')
sent_fo = next((ln for s in got_fo['series'] for ax in s['axes'] for ln in ax['lines']
                if ln['label'] == 'RH cycles completed'), {})
staged_fo = plot_compare.load_sidecar(plot_compare.find_sidecar(
    src_fo[0]['output_dir'], src_fo[0]['filename']))['data']['axes'][0] if src_fo else {}
xs_fo = [x for tr in staged_fo.get('lines', []) for x in tr['x'] if x is not None]
check('through the portal: its RH cycles in the index beside the hours, its map sent to the chart as '
      'worked out, and Export staging its OCV against them, from 0 to 236',
      ({k: v for k, v in kv_of(FCTS_S).items() if k.startswith(('Hours', 'RH'))},
       len(sent_fo.get('x') or []) == len(mx_o) and np.allclose(sent_fo['x'], mx_o, atol=1e-4)
       and list(sent_fo['y']) == my_o.tolist(),
       [s['label'] for s in src_fo], staged_fo.get('xlabel'),
       bool(xs_fo) and min(xs_fo) >= 0 and 230 < max(xs_fo) <= 236),
      ({'Hours tested': 9.01667, 'RH cycles tested': 236}, True, [FCTS_S], 'RH cycles', True))

shutil.rmtree(tmp, ignore_errors=True)
print(f'\n{passed} passed, {failed} failed')
sys.exit(1 if failed else 0)
