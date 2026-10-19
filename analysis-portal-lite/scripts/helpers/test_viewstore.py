"""
Offline harness for viewstore.py. JSONBin transport is mocked.

Run:  python3 scripts/helpers/test_viewstore.py
"""

import json
import shutil
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

from scripts.helpers import jsonbin, viewstore, xlsx_export  # noqa: E402
from scripts.helpers.record import encode_sidecars  # noqa: E402

_passed = 0
_failed = 0


def check(name, got, want):
    global _passed, _failed
    if got == want:
        _passed += 1
    else:
        _failed += 1
        print(f"  FAIL {name}\n       got:  {got!r}\n       want: {want!r}")


def check_true(name, cond, detail=''):
    global _passed, _failed
    if cond:
        _passed += 1
    else:
        _failed += 1
        print(f"  FAIL {name} {detail}")


# ─────────────────────────────────────────────────────────────────────
#  Fixtures
# ─────────────────────────────────────────────────────────────────────

def sidecar(plot_type, text=''):
    return {'plot_type': plot_type, 'data': {'axes': [{
        'texts': [{'text': text}] if text else [],
        'axhlines': [], 'axvlines': [],
        'lines': [{'x': [0, 1], 'y': [1, 0], 'label': 'L'}]}]}}


DETAILS = {
    'job-A': {
        'schema': 2, 'job_id': 'job-A', 'sample_name': '260126_FCS6',
        'script': 'FC Polarization Curve', 'timestamp': '2026-05-01T10:00:00Z',
        'metrics': {'polcurve': {
            'polcurve_b14_80C_100RH': {
                'conditions': {'step': 'b14', 'T_C': 80.0},
                'values': {'OCV': {'value': 0.95, 'unit': 'V'}}}}},
        'sidecars': encode_sidecars(
            {'polcurve_b14_80C_100RH': sidecar('polcurve', 'OCV = 0.95 V')}),
    },
    'job-B': {
        'schema': 2, 'job_id': 'job-B', 'sample_name': '260421_GSMA-Qual-1',
        'script': 'Fuel Cell Full Analysis', 'timestamp': '2026-07-22T22:26:04Z',
        'metrics': {
            'polcurve': {'polcurve_b4_80C_95RH': {
                'conditions': {'step': 'b4'}, 'values': {}}},
            # Cleaning metrics are present but its sidecar was excluded at
            # write time, so this plot must report renderable: False.
            'cleaning': {'cleaning_cycles_a4': {
                'conditions': {}, 'values': {'Cycles': 50.0}}},
        },
        'sidecars': encode_sidecars(
            {'polcurve_b4_80C_95RH': sidecar('polcurve')}),
    },
}

INDEX = {'schema': 2, 'runs': [
    {'job_id': 'job-A', 'sample_name': '260126_FCS6',
     'script': 'FC Polarization Curve', 'timestamp': '2026-05-01T10:00:00Z',
     'bin_id': 'BIN-A',
     'Data': [{'Analysis': 'polcurve', 'step': 'b14',
               'Conditions': {'T_C': 80.0}, 'key_values': {'OCV': 0.95}}]},
    {'job_id': 'job-B', 'sample_name': '260421_GSMA-Qual-1',
     'script': 'Fuel Cell Full Analysis', 'timestamp': '2026-07-22T22:26:04Z',
     'bin_id': 'BIN-B',
     'Data': [{'Analysis': 'polcurve', 'step': 'b4', 'Conditions': {}},
              {'Analysis': 'cleaning', 'step': '', 'Conditions': {}}]},
]}

BIN_TO_JOB = {'BIN-A': 'job-A', 'BIN-B': 'job-B'}
_calls = []


def mock_request(url, method='GET', body=None, extra_headers=None):
    _calls.append((method, url))
    for bin_id, job_id in BIN_TO_JOB.items():
        if f'/{bin_id}/' in url:
            return {'record': DETAILS[job_id], 'metadata': {}}
    return {'record': INDEX, 'metadata': {}}


jsonbin._request = mock_request
jsonbin._JSONBIN_BASE = 'https://api.jsonbin.io/v3/b'


def reset():
    viewstore.clear_cache()
    _calls.clear()


# ─────────────────────────────────────────────────────────────────────

print("listing and filters")
reset()
res = viewstore.list_runs()
check('all runs', res['total'], 2)
check('newest first', [r['job_id'] for r in res['runs']], ['job-B', 'job-A'])

check('sample substring', viewstore.list_runs(sample='GSMA')['total'], 1)
check('sample case-insensitive', viewstore.list_runs(sample='gsma')['total'], 1)
check('sample no match', viewstore.list_runs(sample='zzz')['total'], 0)
check('script exact', viewstore.list_runs(script='FC Polarization Curve')['total'], 1)
check('analysis present in Data',
      viewstore.list_runs(analysis='cleaning')['total'], 1)
check('analysis in both', viewstore.list_runs(analysis='polcurve')['total'], 2)
check('since filter', viewstore.list_runs(since='2026-06-01')['total'], 1)
check('until filter', viewstore.list_runs(until='2026-06-01')['total'], 1)
check('range excludes both',
      viewstore.list_runs(since='2026-06-01', until='2026-06-02')['total'], 0)
check('combined filters',
      viewstore.list_runs(sample='FCS6', analysis='polcurve')['total'], 1)

r = viewstore.list_runs(limit=1)
check('limit caps returned', r['returned'], 1)
check('limit preserves total', r['total'], 2)

print("ordering by experiment date")
# The index entries were analysed in one batch but the experiments span months,
# so ordering must follow run_date where it exists.
_orig = INDEX['runs']
INDEX['runs'] = [
    {'job_id': 'old-exp', 'sample_name': '260407_BM1', 'script': 'S',
     'timestamp': '2026-07-23T18:00:00Z', 'run_date': '2026-04-07',
     'bin_id': 'B1', 'Data': []},
    {'job_id': 'new-exp', 'sample_name': '260511_Gen2', 'script': 'S',
     'timestamp': '2026-07-22T10:00:00Z', 'run_date': '2026-05-11',
     'bin_id': 'B2', 'Data': []},
    {'job_id': 'undated', 'sample_name': 'NoPrefix', 'script': 'S',
     'timestamp': '2026-07-23T19:00:00Z', 'bin_id': 'B3', 'Data': []},
]
viewstore.clear_cache()
order = [r['job_id'] for r in viewstore.list_runs()['runs']]
check('newest experiment first, undated falls back to analysis date',
      order, ['undated', 'new-exp', 'old-exp'])
check_true('run_date beats a later analysis timestamp',
           order.index('new-exp') < order.index('old-exp'))
INDEX['runs'] = _orig
viewstore.clear_cache()

print("facets")
f = viewstore.list_runs()['facets']
check('sample facet', f['samples'], ['260126_FCS6', '260421_GSMA-Qual-1'])
check('analysis facet', f['analyses'], ['cleaning', 'polcurve'])
check('script facet', len(f['scripts']), 2)

print("index caching")
reset()
viewstore.list_runs()
n1 = len(_calls)
viewstore.list_runs()
viewstore.list_runs()
check('index fetched once within TTL', len(_calls), n1)
viewstore.list_runs(force=True)
check_true('force refetches', len(_calls) > n1)

print("detail fetch and cache")
reset()
d = viewstore.fetch_detail('job-A')
check('right record', d['job_id'], 'job-A')
n_after_first = len(_calls)
viewstore.fetch_detail('job-A')
check('second fetch served from cache', len(_calls), n_after_first)

stats = viewstore.cache_stats()
check('one cached', stats['detail_cached'], 1)
check('capacity reported', stats['detail_capacity'], viewstore.DETAIL_CACHE_SIZE)

raised = False
try:
    viewstore.fetch_detail('job-missing')
except KeyError:
    raised = True
check_true('unknown job raises KeyError', raised)

# LRU eviction: shrink the cache and confirm the oldest goes first.
orig = viewstore.DETAIL_CACHE_SIZE
viewstore.DETAIL_CACHE_SIZE = 1
reset()
viewstore.fetch_detail('job-A')
viewstore.fetch_detail('job-B')
check('evicted to capacity', viewstore.cache_stats()['detail_cached'], 1)
check_true('newest retained',
           any('BIN-B' in k for k in viewstore.cache_stats()['detail_keys']))
viewstore.DETAIL_CACHE_SIZE = orig

print("keying by bin id, sample and job")
reset()
# bin_id is unambiguous; sample_name and job_id resolve too.
check('resolve by bin id', viewstore.fetch_detail('BIN-A')['job_id'], 'job-A')
check('resolve by job id', viewstore.fetch_detail('job-A')['job_id'], 'job-A')
check('resolve by sample name',
      viewstore.fetch_detail('260126_FCS6')['job_id'], 'job-A')
raised = False
try:
    viewstore.fetch_detail('nothing-like-this')
except KeyError:
    raised = True
check_true('unknown key raises', raised)

# Merging rewrites a bin in place, so the cache must key on the entry's
# timestamp rather than treating detail records as immutable.
reset()
viewstore.fetch_detail('BIN-A')
n_before = len(_calls)
viewstore.fetch_detail('BIN-A')
check('unchanged entry is served from cache', len(_calls), n_before)
INDEX['runs'][0]['timestamp'] = '2026-08-01T00:00:00Z'   # simulate a merge
viewstore.clear_cache()
_calls.clear()
viewstore.fetch_detail('BIN-A')
first = len(_calls)
INDEX['runs'][0]['timestamp'] = '2026-08-02T00:00:00Z'   # merged again
viewstore._index_cache['at'] = 0                          # index TTL expires
viewstore.fetch_detail('BIN-A')
check_true('a merged bin is re-read rather than served stale',
           len(_calls) > first)
INDEX['runs'][0]['timestamp'] = '2026-05-01T10:00:00Z'
reset()

print("plot inventory")
reset()
plots = viewstore.run_plots('job-B')
check('both plots listed', len(plots), 2)
by_name = {p['plot']: p for p in plots}
check_true('stored sidecar is renderable',
           by_name['polcurve_b4_80C_95RH']['renderable'])
check_true('excluded cleaning sidecar is not renderable',
           not by_name['cleaning_cycles_a4']['renderable'])
check('cleaning metrics still present',
      by_name['cleaning_cycles_a4']['values'], {'Cycles': 50.0})
check('analysis tagged', by_name['cleaning_cycles_a4']['analysis'], 'cleaning')

print("materialisation")
reset()
tmp = tempfile.mkdtemp()
try:
    written = viewstore.materialize_sidecars('job-A', Path(tmp))
    check('one sidecar written', written, ['polcurve_b14_80C_100RH'])
    path = Path(tmp) / '_plot_data' / 'polcurve_b14_80C_100RH.json'
    check_true('written in analysis-output layout', path.exists())
    check('content round-trips', json.loads(path.read_text())['plot_type'],
          'polcurve')

    # find_sidecar must locate it — this is the whole point of the layout.
    from scripts.helpers.plot_compare import find_sidecar
    found = find_sidecar(tmp, 'polcurve_b14_80C_100RH.png')
    check_true('find_sidecar locates the materialised file', bool(found),
               f'(got {found!r})')

    # Requesting an unstored plot yields nothing rather than erroring.
    none_written = viewstore.materialize_sidecars(
        'job-B', Path(tmp) / 'b', plots=['cleaning_cycles_a4'])
    check('unstored plot skipped', none_written, [])
finally:
    shutil.rmtree(tmp)

print("compare staging")
reset()
tmp = tempfile.mkdtemp()
try:
    sources = viewstore.materialize_for_compare([
        {'job_id': 'job-A', 'plot': 'polcurve_b14_80C_100RH'},
        {'job_id': 'job-B', 'plot': 'polcurve_b4_80C_95RH'},
        # Not stored — must be dropped rather than producing a broken source.
        {'job_id': 'job-B', 'plot': 'cleaning_cycles_a4'},
    ], Path(tmp))
    check('unstored selection dropped', len(sources), 2)
    check('sample names attached',
          sorted(s['sample_name'] for s in sources),
          ['260126_FCS6', '260421_GSMA-Qual-1'])
    check_true('per-run directories',
              sources[0]['output_dir'] != sources[1]['output_dir'])
    check_true('filenames carry an extension',
               all(s['filename'].endswith('.png') for s in sources))

    from scripts.helpers.plot_compare import find_sidecar
    check_true('every staged source resolves',
               all(find_sidecar(s['output_dir'], s['filename'])
                   for s in sources))

    check('label defaults empty', sources[0]['label'], '')
    labelled = viewstore.materialize_for_compare(
        [{'job_id': 'job-A', 'plot': 'polcurve_b14_80C_100RH',
          'label': 'baseline'}], Path(tmp) / 'l')
    check('label passed through', labelled[0]['label'], 'baseline')
finally:
    shutil.rmtree(tmp)

print("RH cycles axis")
# FC Durability stores, with a test that cycles RH, its hold hours against RH
# cycles completed; Comparisons' Export can stage a durability plot against
# cycles off it. Here a first hold cycles RH for 2 h (50 an hour) and a second
# holds it steady for 1 h.
MAP = {'plot_type': 'fcdur_rh_cycles', 'data': {'axes': [{
    'xlabel': 'Hold hours (h)', 'ylabel': 'RH cycles', 'texts': [], 'axhlines': [],
    'axvlines': [{'x': 2.0}],
    'lines': [{'x': [0.0, 1.0, 2.0, 3.0], 'y': [0, 50, 100, 100],
               'label': 'RH cycles completed'}]}]}}


def ocv_plot():
    return {'plot_type': 'fcdur_ocv_vs_time', 'data': {'axes': [{
        'xlabel': 'Hold time (h)', 'ylabel': 'Hold OCV (V)', 'xlim': [-0.15, 3.15],
        'texts': [{'text': 'OCV loss = 1 %', 'x': 0.02, 'y': 0.9, 'in_axes_fraction': True},
                  {'text': 'death', 'x': 2.75, 'y': 0.7, 'in_axes_fraction': False}],
        'axhlines': [{'y': 0.7}], 'axvlines': [{'x': 2.0}, {'x': 2.75}],
        'lines': [{'x': [0.5, 1.5], 'y': [0.95, 0.94], 'label': '100↔0% RH cycling'},
                  {'x': [2.5], 'y': [0.93], 'label': '30% RH'}],
        'scatters': [{'x': [0.5, 1.5], 'y': [0.9, 0.9], 'label': 'bin low'}]}]}}


check('the map: hours and cycles, in order',
      viewstore.rh_cycle_map({'m': MAP, 'o': ocv_plot()}), ([0.0, 1.0, 2.0, 3.0], [0.0, 50.0, 100.0, 100.0]))
stepped = json.loads(json.dumps(MAP))
stepped['data']['axes'][0]['lines'][0].update(x=[0.0, 1.0, 0.5, 2.0, None], y=[0, 50, 25, 100, 120])
check('... a point stepping back in hours, or missing a number, left out',
      viewstore.rh_cycle_map({'m': stepped}), ([0.0, 1.0, 2.0], [0.0, 50.0, 100.0]))
short = json.loads(json.dumps(MAP))
short['data']['axes'][0]['lines'][0].update(x=[0.0], y=[0])
renamed = json.loads(json.dumps(MAP))
renamed['data']['axes'][0]['lines'][0]['label'] = 'cycles'
check('... none from one point, another line, or a test without the plot',
      [viewstore.rh_cycle_map(s) for s in ({'m': short}, {'m': renamed}, {'o': ocv_plot()}, {})],
      [None, None, None, None])
H, C = viewstore.rh_cycle_map({'m': MAP})
check('cycles at an hour: straight between points, held at the ends, a gap passed through',
      [viewstore.cycles_at(H, C, h) for h in (-1, 0, 0.5, 1.0, 1.25, 2.5, 3.0, 9.0, None)],
      [0.0, 0.0, 25.0, 50.0, 62.5, 100.0, 100.0, 100.0, None])
moved = viewstore.sidecar_on_cycles(ocv_plot(), H, C)['data']['axes'][0]
check('a plot against cycles: lines, points, vertical lines, text placed by data, range, label',
      ([ln['x'] for ln in moved['lines']], moved['scatters'][0]['x'],
       [v['x'] for v in moved['axvlines']], [t['x'] for t in moved['texts']],
       moved['xlim'], moved['xlabel'], moved['ylabel'], moved['axhlines']),
      ([[25.0, 75.0], [100.0]], [25.0, 75.0], [100.0, 100.0], [0.02, 100.0],
       [0.0, 100.0], 'RH cycles', 'Hold OCV (V)', [{'y': 0.7}]))

_orig = INDEX['runs']
INDEX['runs'] = _orig + [
    {'job_id': 'job-C', 'sample_name': '261001_CYC', 'script': 'FC Durability',
     'timestamp': '2026-10-01T10:00:00Z', 'bin_id': 'BIN-C',
     'Data': [{'Analysis': 'fc_durability', 'step': '', 'Conditions': {}}]},
    {'job_id': 'job-D', 'sample_name': '261002_STEADY', 'script': 'FC Durability',
     'timestamp': '2026-10-02T10:00:00Z', 'bin_id': 'BIN-D',
     'Data': [{'Analysis': 'fc_durability', 'step': '', 'Conditions': {}}]}]
BIN_TO_JOB.update({'BIN-C': 'job-C', 'BIN-D': 'job-D'})
long_map = json.loads(json.dumps(MAP))
long_map['data']['axes'][0]['lines'][0].update(x=[k / 100 for k in range(901)],
                                               y=[min(k, 450) for k in range(901)])
for job, sample, plots in (('job-C', '261001_CYC', {'ocv': ocv_plot(), 'cyc': MAP}),
                           ('job-D', '261002_STEADY', {'ocv': ocv_plot()}),
                           ('job-L', '261003_LONG', {'cyc': long_map})):
    DETAILS[job] = {
        'schema': 2, 'job_id': job, 'sample_name': sample, 'script': 'FC Durability',
        'metrics': {'fc_durability': {f'{sample}-FCDurability-{k}': {
            'conditions': {'step': ''}, 'values': {}} for k in plots}},
        'sidecars': encode_sidecars({f'{sample}-FCDurability-{k}': v for k, v in plots.items()})}
INDEX['runs'].append({'job_id': 'job-L', 'sample_name': '261003_LONG', 'script': 'FC Durability',
                      'timestamp': '2026-10-03T10:00:00Z', 'bin_id': 'BIN-L',
                      'Data': [{'Analysis': 'fc_durability', 'step': '', 'Conditions': {}}]})
BIN_TO_JOB['BIN-L'] = 'job-L'
reset()
tmp = tempfile.mkdtemp()
try:
    ask = {'analysis': 'fc_durability', 'step': '', 'plot_type': 'fcdur_rh_cycles', 'full': True}
    got = viewstore.plot_series_batch([dict(ask, key='BIN-L'), dict(ask, key='BIN-D')])
    line = got['series'][0]['axes'][0]['lines'][0]
    check('the map sent whole to the chart, 901 points not thinned; a test without one missing',
          (len(line['x']), line.get('decimated', False), [m['key'] for m in got['missing']]),
          (901, False, ['BIN-D']))
    OCV_C, OCV_D = '261001_CYC-FCDurability-ocv', '261002_STEADY-FCDurability-ocv'
    sources = viewstore.materialize_for_compare([
        {'key': 'BIN-C', 'plot': OCV_C, 'label': 'cycled', 'x': 'cycles'},
        {'key': 'BIN-C', 'plot': OCV_C, 'label': 'again', 'x': 'cycles'},
        {'key': 'BIN-D', 'plot': OCV_D, 'label': 'steady', 'x': 'cycles'},
        {'key': 'BIN-A', 'plot': 'polcurve_b14_80C_100RH', 'label': 'polcurve', 'x': 'cycles'},
    ], Path(tmp))
    check('staged against cycles: the cycling test (once for both), not the steady one; '
          'not a durability plot, as stored',
          [(s['label'], s['plot_type']) for s in sources],
          [('cycled', 'fcdur_ocv_vs_time'), ('again', 'fcdur_ocv_vs_time'), ('polcurve', 'polcurve')])
    from scripts.helpers.plot_compare import find_sidecar, load_sidecar
    staged = [load_sidecar(find_sidecar(s['output_dir'], s['filename']))['data']['axes'][0]
              for s in sources]
    check('... its x read off its own map, converted once',
          ([ln['x'] for ln in staged[0]['lines']], [v['x'] for v in staged[0]['axvlines']],
           staged[0]['xlabel'], staged[1]['lines'][0]['x']),
          ([[25.0, 75.0], [100.0]], [100.0, 100.0], 'RH cycles', [25.0, 75.0]))
    check('... the polcurve staged as stored', (staged[2]['lines'][0]['x'], staged[2].get('xlabel')),
          ([0, 1], None))
    plain = viewstore.materialize_for_compare([
        {'key': 'BIN-C', 'plot': OCV_C}, {'key': 'BIN-D', 'plot': OCV_D, 'x': 'hours'}], Path(tmp) / 'h')
    check('... and on hold hours, every test as stored',
          [load_sidecar(find_sidecar(s['output_dir'], s['filename']))['data']['axes'][0]['lines'][0]['x']
           for s in plain], [[0.5, 1.5], [0.5, 1.5]])
    both = viewstore.materialize_for_compare([
        {'key': 'BIN-C', 'plot': OCV_C, 'x': 'cycles'}, {'key': 'BIN-C', 'plot': OCV_C},
        {'key': 'BIN-C', 'plot': '261001_CYC-FCDurability-cyc', 'x': 'cycles'}], Path(tmp) / 'b')
    check('... the same plot asked for on both axes staged both ways; a plot not against hold '
          'hours (the map itself) as stored',
          [(lambda a: (a['lines'][0]['x'][:2], a.get('xlabel')))(
              load_sidecar(find_sidecar(s['output_dir'], s['filename']))['data']['axes'][0]) for s in both],
          [([25.0, 75.0], 'RH cycles'), ([0.5, 1.5], 'Hold time (h)'), ([0.0, 1.0], 'Hold hours (h)')])
finally:
    shutil.rmtree(tmp)
    INDEX['runs'] = _orig
    for b, j in (('BIN-C', 'job-C'), ('BIN-D', 'job-D'), ('BIN-L', 'job-L')):
        BIN_TO_JOB.pop(b)
        DETAILS.pop(j)
    reset()

print("custom plot: points, labels and the workbook")
# The Comparisons custom plot reads what other charts leave out: the points an
# FC Durability plot marks values with (R_short on its own twin axis, a hold's
# few OCV decay windows), and, to list what a sample stores, labels alone.
XOVER = {'plot_type': 'fcdur_xover_vs_time', 'data': {'axes': [
    {'xlabel': 'Hold hours (h)', 'ylabel': 'H₂ crossover (mA/cm²)', 'texts': [], 'axhlines': [],
     'lines': [{'x': [0.0, 1.0, 2.0], 'y': [2.9, 3.0, 3.1], 'label': 'Crossover, 100% RH, window mean'},
               {'x': [0.0, 1.0, 2.0], 'y': [2.7, None, 2.8], 'label': 'Crossover, 50% RH, window mean'}],
     'scatters': [{'x': [0.0], 'y': [2.95], 'label': 'Crossover, 100% RH, earlier repeats'}]},
    {'ylabel': 'R_short (Ω·cm²), ◆; DOE target ≥ 1,000 (dash-dot)', 'is_twin': True,
     'scatters': [{'x': [1.0, 2.0], 'y': [9000.0, 4000.0], 'label': 'R_short'}]}]}}
OCVD = {'plot_type': 'fcdur_ocv_vs_time', 'data': {'axes': [
    {'ylabel': 'Hold OCV (V)', 'texts': [], 'axhlines': [],
     'lines': [{'x': [0.5, 1.5], 'y': [0.95, 0.94], 'label': '100↔0% RH cycling'}],
     'scatters': [{'x': [0.5, 1.5], 'y': [0.9, 0.9], 'label': 'bin low'},
                  {'x': [0.5, 1.5], 'y': [0.99, 0.98], 'label': 'bin high'},
                  {'x': [0.2, 1.9], 'y': [0.96, 0.93], 'label': 'settled start and end (30-min medians)'}]},
    {'xlabel': 'Hold hours (h)', 'ylabel': 'OCV decay (μV/h)\n10 h windows',
     'lines': [{'x': [0.5, 1.0, 1.5], 'y': [300.0, 320.0, 340.0], 'label': None}],
     'scatters': [{'x': [2.5], 'y': [500.0], 'label': None}]}]}}
POL = {'plot_type': 'polcurve', 'data': {'axes': [
    {'xlabel': 'Current density (A/cm²)', 'ylabel': 'Voltage (V)', 'title': 'Polarization Curve',
     'lines': [{'x': [0.0, 0.5, 1.0], 'y': [0.95, 0.75, 0.6], 'label': 'Cell voltage'},
               {'x': [0.0, 0.5, 1.0], 'y': [0.95, 0.78, 0.66], 'label': 'iR-free voltage'}]},
    {'ylabel': 'Power density (mW/cm²)', 'is_twin': True,
     'lines': [{'x': [0.0, 0.5, 1.0], 'y': [0.0, 375.0, 600.0], 'label': 'Power density'}]}]}}
# A short first hold, its decay windows few enough to be stored as points,
# then a long one whose windows are a line: the points come first in x.
OCVQ = {'plot_type': 'fcdur_ocv_vs_time', 'data': {'axes': [
    {'ylabel': 'Hold OCV (V)', 'lines': [{'x': [0.5, 8.0], 'y': [0.95, 0.9], 'label': '30% RH'}]},
    {'xlabel': 'Hold hours (h)', 'ylabel': 'OCV decay (μV/h)\n10 h windows',
     'lines': [{'x': [5.0, 6.0, 7.0], 'y': [300.0, 320.0, 340.0], 'label': None}],
     'scatters': [{'x': [2.5], 'y': [900.0], 'label': None}]}]}}
# Holds stored out of x order, each band's edges listed apart from its hold:
# the middle hold has a gap in its values and in its low edge, the last
# hold's edges are on an x of their own, so it has no band; a hold listed
# after them has a low edge and no high, and another edges with no number.
OCVB = {'plot_type': 'fcdur_ocv_vs_time', 'data': {'axes': [
    {'ylabel': 'Hold OCV (V)', 'lines': [
        {'x': [5.0, 6.0, 7.0], 'y': [0.93, None, 0.91], 'label': '30% RH'},
        {'x': [0.5, 1.0], 'y': [0.95, 0.94], 'label': '30% RH'},
        {'x': [9.0, 10.0], 'y': [0.9, 0.89], 'label': '30% RH'},
        {'x': [5.0, 6.0, 7.0], 'y': [0.9, 0.88, None], 'label': 'bin low'},
        {'x': [0.5, 1.0], 'y': [0.92, 0.93], 'label': 'bin low'},
        {'x': [0.5, 1.0], 'y': [0.97, 0.96], 'label': 'bin high'},
        {'x': [5.0, 6.0, 7.0], 'y': [0.95, 0.94, 0.93], 'label': 'bin high'},
        {'x': [9.0, 10.5], 'y': [0.88, 0.87], 'label': 'bin low'},
        {'x': [9.0, 10.5], 'y': [0.92, 0.91], 'label': 'bin high'},
        {'x': [12.0, 13.0], 'y': [0.88, 0.87], 'label': '30% RH'},
        {'x': [12.0, 13.0], 'y': [0.85, 0.84], 'label': 'bin low'},
        {'x': [14.0, 15.0], 'y': [0.86, 0.85], 'label': '30% RH'},
        {'x': [14.0, 15.0], 'y': [None, None], 'label': 'bin low'},
        {'x': [14.0, 15.0], 'y': [None, None], 'label': 'bin high'}]},
    {'xlabel': 'Hold hours (h)', 'ylabel': 'OCV decay (μV/h)\n10 h windows',
     'lines': [{'x': [5.0, 6.0, 7.0], 'y': [300.0, 320.0, 340.0], 'label': None}]}]}}
_orig = INDEX['runs']
INDEX['runs'] = _orig + [
    {'job_id': 'job-X', 'sample_name': '261005_XO', 'script': 'FC Durability',
     'timestamp': '2026-10-05T10:00:00Z', 'bin_id': 'BIN-X',
     'Data': [{'Analysis': 'fc_durability', 'step': '', 'Conditions': {}}]},
    {'job_id': 'job-S', 'sample_name': '261006_STEADY', 'script': 'FC Durability',
     'timestamp': '2026-10-06T10:00:00Z', 'bin_id': 'BIN-S',
     'Data': [{'Analysis': 'fc_durability', 'step': '', 'Conditions': {}}]},
    {'job_id': 'job-P', 'sample_name': '261007_POL', 'script': 'FC Polarization Curve',
     'timestamp': '2026-10-07T10:00:00Z', 'bin_id': 'BIN-P',
     'Data': [{'Analysis': 'polcurve', 'step': 'b4', 'Conditions': {}}]},
    {'job_id': 'job-Q', 'sample_name': '261008_ORDER', 'script': 'FC Durability',
     'timestamp': '2026-10-08T10:00:00Z', 'bin_id': 'BIN-Q',
     'Data': [{'Analysis': 'fc_durability', 'step': '', 'Conditions': {}}]},
    {'job_id': 'job-B', 'sample_name': '261009_BAND', 'script': 'FC Durability',
     'timestamp': '2026-10-09T10:00:00Z', 'bin_id': 'BIN-B',
     'Data': [{'Analysis': 'fc_durability', 'step': '', 'Conditions': {}}]}]
for job, sample, bucket, plots in (
        ('job-X', '261005_XO', 'fc_durability', {'xo': XOVER, 'ocv': OCVD, 'cyc': MAP}),
        ('job-S', '261006_STEADY', 'fc_durability', {'xo': XOVER}),
        ('job-P', '261007_POL', 'polcurve', {'b4': POL}),
        ('job-Q', '261008_ORDER', 'fc_durability', {'ocv': OCVQ}),
        ('job-B', '261009_BAND', 'fc_durability', {'ocv': OCVB})):
    BIN_TO_JOB['BIN-' + job[-1]] = job
    DETAILS[job] = {
        'schema': 2, 'job_id': job, 'sample_name': sample, 'script': 'S',
        'metrics': {bucket: {f'{sample}-{k}': {
            'conditions': {'step': 'b4' if bucket == 'polcurve' else ''}, 'values': {}}
            for k in plots}},
        'sidecars': encode_sidecars({f'{sample}-{k}': v for k, v in plots.items()})}
reset()
try:
    FD = {'analysis': 'fc_durability', 'step': ''}
    plain = viewstore.plot_series('BIN-X', plot_type='fcdur_xover_vs_time', **FD)
    check('asked for as before, the R_short axis (points only) is left out, and no axis has points',
          ([a['ylabel'] for a in plain['axes']], any('points' in a for a in plain['axes']),
           [ln['label'] for ln in plain['axes'][0]['lines']]),
          (['H₂ crossover (mA/cm²)'], False,
           ['Crossover, 100% RH, window mean', 'Crossover, 50% RH, window mean']))
    pts = viewstore.plot_series('BIN-X', plot_type='fcdur_xover_vs_time', points=True, **FD)
    check('with points: each axis its points, and the R_short axis with only its points',
          [(a['ylabel'][:7], [ln['label'] for ln in a['lines']],
            [(p['label'], p['x'], p['y']) for p in a['points']]) for a in pts['axes']],
          [('H₂ cros', ['Crossover, 100% RH, window mean', 'Crossover, 50% RH, window mean'],
            [('Crossover, 100% RH, earlier repeats', [0.0], [2.95])]),
           ('R_short', [], [('R_short', [1.0, 2.0], [9000.0, 4000.0])])])
    ocv = viewstore.plot_series('BIN-X', plot_type='fcdur_ocv_vs_time', points=True, **FD)
    check('... a band\'s edges stay lines, not points; unlabelled decay windows come as lines '
          'and points with an empty label',
          [([ln['label'] for ln in a['lines']], [p['label'] for p in a['points']]) for a in ocv['axes']],
          [(['100↔0% RH cycling', 'bin low', 'bin high'], ['settled start and end (30-min medians)']),
           ([''], [''])])
    lab = viewstore.plot_series('BIN-X', plot_type='fcdur_xover_vs_time', points=True,
                                labels=True, **FD)
    check('labels alone: every line and point its label and stored count, no numbers',
          [(a['lines'], a['points']) for a in lab['axes']],
          [([{'label': 'Crossover, 100% RH, window mean', 'n': 3},
             {'label': 'Crossover, 50% RH, window mean', 'n': 3}],
            [{'label': 'Crossover, 100% RH, earlier repeats', 'n': 1}]),
           ([], [{'label': 'R_short', 'n': 2}])])
    batch = viewstore.plot_series_batch([dict(FD, key='BIN-X', plot_type='fcdur_xover_vs_time',
                                              points=True, labels=True),
                                         dict(FD, key='BIN-X', plot_type='fcdur_xover_vs_time',
                                              points='yes')])
    check('the batch passes points and labels on, only as true',
          ([len(a['points']) for a in batch['series'][0]['axes']], 'n' in batch['series'][0]['axes'][0]['lines'][0],
           len(batch['series'][1]['axes']), 'points' in batch['series'][1]['axes'][0]),
          ([1, 1], True, 1, False))
    long_pts = json.loads(json.dumps(XOVER))
    long_pts['data']['axes'][1]['scatters'][0].update(x=list(range(1000)), y=[float(v) for v in range(1000)])
    DETAILS['job-S']['sidecars'] = encode_sidecars({'261006_STEADY-xo': long_pts})
    reset()
    thin = viewstore.plot_series('BIN-S', plot_type='fcdur_xover_vs_time', points=True, **FD)['axes'][1]['points'][0]
    whole = viewstore.plot_series('BIN-S', plot_type='fcdur_xover_vs_time', points=True, stored=True,
                                  **FD)['axes'][1]['points'][0]
    check('points thinned as lines are; read for an export, every stored point',
          (thin.get('decimated'), len(thin['x']) <= viewstore.MAX_SERIES_POINTS + 2,
           whole.get('decimated', False), len(whole['x'])), (True, True, False, 1000))
    DETAILS['job-S']['sidecars'] = encode_sidecars({'261006_STEADY-xo': XOVER})
    reset()

    def curve(key, src, pick, **kw):
        return dict({'key': key, 'name': kw.pop('name', 'shown'), 'stored': 'stored', 'quantity': 'q',
                     'unit': 'u', 'axis': 'Left', 'series': 's', 'cond': 'c', 'src': src, 'pick': pick}, **kw)
    XO = dict(FD, plot_type='fcdur_xover_vs_time', points=True)
    OC = dict(FD, plot_type='fcdur_ocv_vs_time', points=True)
    PO = {'analysis': 'polcurve', 'step': 'b4'}
    got = viewstore.custom_plot_curves([
        curve('BIN-X', XO, {'axis': 1, 'points': [0], 'points_labels': ['R_short']}),
        curve('BIN-X', XO, {'axis': 0, 'lines': [1], 'lines_labels': ['Crossover, 50% RH, window mean']}),
        curve('BIN-X', OC, {'axis': 1, 'lines': [0], 'lines_labels': [''], 'points': [0],
                            'points_labels': ['']}),
        curve('BIN-P', PO, {'axis': 0, 'lines': [0], 'lines_labels': ['Cell voltage'], 'power': True}),
        curve('BIN-P', PO, {'axis': 0, 'lines': [1], 'lines_labels': ['Cell voltage']}),
        curve('BIN-P', PO, {'axis': 5, 'lines': [0], 'lines_labels': ['Cell voltage']}),
        curve('BIN-A', {'analysis': 'polcurve', 'step': 'zz'}, {'axis': 0, 'lines': [0]}),
    ])
    check('the workbook\'s curves: points and lines where the chart found them, a point without '
          'both numbers left out, the decay windows as lines and points, a power curve j·V in mW/cm²',
          [c['parts'] for c in got[:4]],
          [[[[1.0, 9000.0], [2.0, 4000.0]]], [[[0.0, 2.7], [2.0, 2.8]]],
           [[[0.5, 300.0], [1.0, 320.0], [1.5, 340.0]], [[2.5, 500.0]]],
           [[[0.0, 0.0], [0.5, 0.75 * 0.5 * 1000.0], [1.0, 600.0]]]])
    check('... one whose label has moved, or whose axis is gone, or whose plot is not stored, '
          'has no points and says why',
          [(c['parts'], c['reason']) for c in got[4:]],
          [([], 'not found in its stored plot (stored again since?)'),
           ([], 'not found in its stored plot (stored again since?)'),
           ([], 'no stored plot data')])
    check('... its descriptive fields as sent', (got[0]['key'], got[0]['name'], got[0]['stored'],
                                                got[0]['axis']), ('BIN-X', 'shown', 'stored', 'Left'))
    cyc = viewstore.custom_plot_curves([
        curve('BIN-X', XO, {'axis': 0, 'lines': [0], 'lines_labels': ['Crossover, 100% RH, window mean']}),
        curve('BIN-S', XO, {'axis': 0, 'lines': [0], 'lines_labels': ['Crossover, 100% RH, window mean']}),
    ], 'cyc')
    check('against RH cycles: hours read off the test\'s own map; a test without one says so',
          [(c['parts'], c['reason']) for c in cyc],
          [([[[0.0, 2.9], [50.0, 3.0], [100.0, 3.1]]], ''), ([], 'no RH cycles counted')])

    from openpyxl import load_workbook
    import io as _io
    wbk = load_workbook(_io.BytesIO(xlsx_export.build_custom_workbook(
        {'xtitle': 'Hold time (h)', 'layout': 'overlaid', 'exported': '15 Oct 2026',
         'axes': [{'side': 'Left axis', 'text': 'R_short (Ω·cm²)'}]}, got[:2] + got[4:5])))
    check('workbook sheets', wbk.sheetnames, ['Custom plot', 'Curves', 'Data'])
    rows_c = [[c.value for c in r] for r in wbk['Curves'].iter_rows()]
    check('Curves: one row per curve as the legend named it, with its stored name; '
          'one with no points says why',
          [r[:3] + r[8:] for r in rows_c],
          [['Curve', 'Sample', 'Stored name', 'Points', 'Note'], [1, 'shown', 'stored', 2, None],
           [2, 'shown', 'stored', 2, None],
           [3, 'shown', 'stored', 0, 'not exported: not found in its stored plot (stored again since?)']])
    data = [[c.value for c in r] for r in wbk['Data'].iter_rows()]
    check('Data: a block per curve with points, x and y columns with their titles',
          data[:4], [['1. shown · q · s', None, None, '2. shown · q · s', None],
                     ['Hold time (h)', 'q (u)', None, 'Hold time (h)', 'q (u)'],
                     [1, 9000, None, 0, 2.7], [2, 4000, None, 2, 2.8]])
    info = {r[0].value: r[1].value for r in wbk['Custom plot'].iter_rows() if r[0].value}
    check('the plot described', (info.get('x'), info.get('Left axis'), info.get('Curves')),
          ('Hold time (h)', 'R_short (Ω·cm²)', '3, 2 with points'))
    check_true('... its note names no band where no curve has one (review, 16 Oct)',
               not any('bin' in str(r[0].value) for r in wbk['Custom plot'].iter_rows() if r[0].value),
               [r[0].value for r in wbk['Custom plot'].iter_rows()])

    # Parts in x order: a short hold's decay windows (points) before a long
    # one's (a line), though a pick lists lines first.
    OQ = dict(FD, plot_type='fcdur_ocv_vs_time', points=True)
    order = viewstore.custom_plot_curves([curve('BIN-Q', OQ, {
        'axis': 1, 'lines': [0], 'lines_labels': [''], 'points': [0], 'points_labels': ['']})])
    check('a curve\'s parts in x order, whatever kind each is stored as',
          order[0]['parts'], [[[2.5, 900.0]], [[5.0, 300.0], [6.0, 320.0], [7.0, 340.0]]])
    wbq = load_workbook(_io.BytesIO(xlsx_export.build_custom_workbook(
        {'xtitle': 'Hold time (h)'}, [dict(order[0], name='=1+1', stored='=HYPERLINK("x")')])))
    dq = [[c for c in r] for r in wbq['Data'].iter_rows()]
    cq = [[c for c in r] for r in wbq['Curves'].iter_rows()]
    check('... each part numbered in its y column\'s title, x beside y',
          [c.value for c in dq[1]], ['Hold time (h)', 'q (u), part 1 of 2', 'Hold time (h)', 'q (u), part 2 of 2'])
    check('names that begin with = are written as text, not formulas',
          [(c.value, c.data_type) for c in cq[1][1:3]] + [(dq[0][0].value, dq[0][0].data_type)],
          [('=1+1', 's'), ('=HYPERLINK("x")', 's'), ('1. =1+1 · q · s', 's')])

    # A Hold OCV curve whose high–low band the chart shades (16 Oct, 17:36):
    # each part's bin low and bin high, paired as the chart pairs them.
    OB = dict(FD, plot_type='fcdur_ocv_vs_time', points=True)
    holds = {'axis': 0, 'lines': [0, 1, 2], 'lines_labels': ['30% RH'] * 3}
    bnd = viewstore.custom_plot_curves([
        curve('BIN-B', OB, holds, band=True),
        curve('BIN-B', OB, holds, band=False),
        curve('BIN-B', OB, holds),
        curve('BIN-Q', OB, {'axis': 0, 'lines': [0], 'lines_labels': ['30% RH']}, band=True),
        curve('BIN-B', OB, {'axis': 1, 'lines': [0], 'lines_labels': ['']}, band=True)])
    check('a band shaded: each hold\'s bin low and high paired by its own x, each pair once, in x order '
          'with its hold; a value with no number left out with its band, and a point with an edge that is '
          'not a number left blank, as the chart leaves it unshaded; a hold whose edges are on another x has none',
          (bnd[0]['parts'], bnd[0]['bands'], bnd[0]['band']),
          ([[[0.5, 0.95], [1.0, 0.94]], [[5.0, 0.93], [7.0, 0.91]], [[9.0, 0.9], [10.0, 0.89]]],
           [[[0.92, 0.97], [0.93, 0.96]], [[0.9, 0.95], [None, None]], None], 'shaded'))
    check('... turned off on the chart, or a curve with none, the same parts and no band; '
          'shaded with none stored, or on a line with no edges, says so',
          [(c['parts'] == bnd[0]['parts'], c['bands'], c['band']) for c in bnd[1:3]]
          + [(c['bands'], c['band']) for c in bnd[3:]],
          [(True, [None, None, None], 'not shaded'), (True, [None, None, None], ''),
           ([None], 'none stored'), ([None], 'none stored')])
    twice = viewstore.custom_plot_curves([
        curve('BIN-B', OB, {'axis': 0, 'lines': [1, 1], 'lines_labels': ['30% RH'] * 2}, band=True)])
    check('... a pair of edges goes to one trace only (a hold picked twice has its band once)',
          twice[0]['bands'], [[[0.92, 0.97], [0.93, 0.96]], None])
    half = viewstore.custom_plot_curves([
        curve('BIN-B', OB, {'axis': 0, 'lines': [9], 'lines_labels': ['30% RH']}, band=True)])
    check('... and a low edge with no high is no band', (half[0]['parts'], half[0]['bands'], half[0]['band']),
          ([[[12.0, 0.88], [13.0, 0.87]]], [None], 'none stored'))
    blank = viewstore.custom_plot_curves([
        curve('BIN-B', OB, {'axis': 0, 'lines': [11], 'lines_labels': ['30% RH']}, band=True),
        curve('BIN-X', OC, {'axis': 0, 'lines': [0], 'lines_labels': ['100↔0% RH cycling'], 'points': [0],
                            'points_labels': ['settled start and end (30-min medians)']}, band=True)])
    check('... nor are edges with no number in them (review, 16 Oct)', (blank[0]['bands'], blank[0]['band']),
          ([None], 'none stored'))
    check('... and a curve\'s points have no band, its line beside them in x order keeping its own',
          (blank[1]['parts'], blank[1]['bands'], blank[1]['band']),
          ([[[0.2, 0.96], [1.9, 0.93]], [[0.5, 0.95], [1.5, 0.94]]], [None, [[0.9, 0.99], [0.9, 0.98]]], 'shaded'))
    cyb = viewstore.custom_plot_curves([
        curve('BIN-X', OC, {'axis': 0, 'lines': [0], 'lines_labels': ['100↔0% RH cycling']}, band=True)], 'cyc')
    check('... against RH cycles, its x read off the map and its band as it was',
          (cyb[0]['parts'], cyb[0]['bands'], cyb[0]['band']),
          ([[[25.0, 0.95], [75.0, 0.94]]], [[[0.9, 0.99], [0.9, 0.98]]], 'shaded'))
    wbb = load_workbook(_io.BytesIO(xlsx_export.build_custom_workbook(
        {'xtitle': 'Hold time (h)'},
        [dict(bnd[0], quantity='Hold OCV', unit='V'), dict(bnd[1], quantity='Hold OCV', unit='V'),
         got[0], dict(bnd[3], quantity='Hold OCV', unit='V')])))
    rows_b = [[c.value for c in r] for r in wbb['Curves'].iter_rows()]
    check('Curves, a curve with a band among them: a High–low band column says of each whether its '
          'bins\' lows and highs are in Data',
          [r[8:] for r in rows_b],
          [['Points', 'High–low band', 'Note'], [6, 'shaded: bin low and high in Data', None],
           [6, 'not shaded on the chart: left out', None], [2, None, None], [2, 'none stored now', None]])
    db = [[c.value for c in r] for r in wbb['Data'].iter_rows()]
    check('Data: beside each banded hold\'s values its bin low and bin high, titled with the unit and the '
          'part; a hold with no band, x and y alone',
          (db[1][:10], [r[:10] for r in db[2:4]]),
          (['Hold time (h)', 'Hold OCV (V), part 1 of 3', 'bin low (V), part 1 of 3', 'bin high (V), part 1 of 3',
            'Hold time (h)', 'Hold OCV (V), part 2 of 3', 'bin low (V), part 2 of 3', 'bin high (V), part 2 of 3',
            'Hold time (h)', 'Hold OCV (V), part 3 of 3'],
           [[0.5, 0.95, 0.92, 0.97, 5, 0.93, 0.9, 0.95, 9, 0.9],
            [1, 0.94, 0.93, 0.96, 7, 0.91, None, None, 10, 0.89]]))
    check('... the band turned off: x and y alone', db[1][11:17],
          ['Hold time (h)', 'Hold OCV (V), part 1 of 3', 'Hold time (h)', 'Hold OCV (V), part 2 of 3',
           'Hold time (h)', 'Hold OCV (V), part 3 of 3'])
    info_b = [r[0].value for r in wbb['Custom plot'].iter_rows() if r[0].value]
    check_true('the note says a shaded band\'s lows and highs are in the data',
               any('each bin\'s low and high' in str(v) for v in info_b), info_b)
    wbo = load_workbook(_io.BytesIO(xlsx_export.build_custom_workbook(
        {'xtitle': 'Hold time (h)'}, [dict(bnd[1], quantity='Hold OCV', unit='V')])))
    info_o = [r[0].value for r in wbo['Custom plot'].iter_rows() if r[0].value]
    check_true('... and says nothing of bands where none is shaded (second review, 16 Oct)',
               not any('bin' in str(v) for v in info_o)
               and [c.value for c in wbo['Curves'][1]][8:11] == ['Points', 'High–low band', 'Note'], info_o)

    # Reads: a record read and decoded once per batch or export, a plot once
    # per export, and a key that cannot be read failing once.
    _decode, _series = viewstore.decode_sidecars, viewstore.plot_series
    n_dec, n_ser = [], []
    viewstore.decode_sidecars = lambda d: (n_dec.append(1), _decode(d))[1]
    viewstore.plot_series = lambda *a, **k: (n_ser.append(1), _series(*a, **k))[1]
    try:
        reset()
        b4 = viewstore.plot_series_batch([
            dict(FD, key='BIN-X', plot_type='fcdur_xover_vs_time', points=True, labels=True),
            dict(FD, key='BIN-X', plot_type='fcdur_ocv_vs_time', points=True, labels=True),
            dict(FD, key='BIN-X', plot_type='fcdur_xover_vs_time', points=True),
            dict(FD, key='BIN-S', plot_type='fcdur_xover_vs_time', points=True)])
        check('a batch decodes each record once, however many of its plots it asks for',
              (len(b4['series']), len(n_dec)), (4, 2))
        reset()
        bad = viewstore.plot_series_batch([dict(FD, key='NOPE', plot_type='fcdur_xover_vs_time'),
                                           dict(FD, key='NOPE', plot_type='fcdur_ocv_vs_time')])
        check('... and a key it cannot read fails once: the index read again once, both missing',
              ([m['reason'][:8] for m in bad['missing']], sum(1 for _m, u in _calls if '/BIN-' not in u)),
              (['KeyError', 'KeyError'], 2))
        n_dec.clear()
        n_ser.clear()
        reset()
        twice = viewstore.custom_plot_curves([
            curve('BIN-X', XO, {'axis': 0, 'lines': [0], 'lines_labels': ['Crossover, 100% RH, window mean']}),
            curve('BIN-X', XO, {'axis': 0, 'lines': [1], 'lines_labels': ['Crossover, 50% RH, window mean']}),
            curve('BIN-X', XO, {'axis': 1, 'points': [0], 'points_labels': ['R_short']}),
            curve('BIN-S', XO, {'axis': 0, 'lines': [0], 'lines_labels': ['Crossover, 100% RH, window mean']}),
            curve('NOPE', XO, {'axis': 0, 'lines': [0], 'lines_labels': ['x']}),
            curve('NOPE', XO, {'axis': 0, 'lines': [1], 'lines_labels': ['y']})], 'cyc')
        check('an export reads each plot once and decodes each record once, its map included',
              (len(n_ser), len(n_dec), [bool(c['parts']) for c in twice]),
              (3, 2, [True, True, True, False, False, False]))
        check('... a key it cannot read says so for each of its curves',
              [c['reason'][:20] for c in twice[4:]], ['could not be read (K'] * 2)
    finally:
        viewstore.decode_sidecars, viewstore.plot_series = _decode, _series
finally:
    INDEX['runs'] = _orig
    for b in ('BIN-X', 'BIN-S', 'BIN-P', 'BIN-Q', 'BIN-B'):
        DETAILS.pop(BIN_TO_JOB.pop(b))
    reset()

print(f"\n{_passed} passed, {_failed} failed")
sys.exit(1 if _failed else 0)
